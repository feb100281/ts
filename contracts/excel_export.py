# contracts/excel_export.py
"""
Выгрузка списка договоров в Excel.

Оформление — фирменное: тёмно-зелёная шапка, Roboto Light 10,
зебра, закреплённая строка заголовков, автофильтр.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# ПАЛИТРА
# ============================================================

DARK_GREEN = "2F6656"
LIGHT_GREEN = "E7F1ED"
STRIPE = "F6F6F6"
BORDER_GRAY = "D9D9D9"
WHITE = "FFFFFF"
TEXT = "050505"
MUTED = "7A7A7A"

FONT_NAME = "Roboto Light"


# ============================================================
# ЭЛЕМЕНТЫ ОФОРМЛЕНИЯ
# ============================================================

_thin = Side(style="thin", color=BORDER_GRAY)
BORDER_THIN = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

FONT_TITLE = Font(name=FONT_NAME, size=13, bold=True, color=DARK_GREEN)
FONT_SUBTITLE = Font(name=FONT_NAME, size=9, color=MUTED)
FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
FONT_NORMAL = Font(name=FONT_NAME, size=10, color=TEXT)
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color=TEXT)

FILL_HEADER = PatternFill("solid", fgColor=DARK_GREEN)
FILL_TOTAL = PatternFill("solid", fgColor=LIGHT_GREEN)
FILL_STRIPE = PatternFill("solid", fgColor=STRIPE)
FILL_NONE = PatternFill(fill_type=None)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)

FMT_DATE = "dd.mm.yyyy"
FMT_DATETIME = "dd.mm.yyyy hh:mm"


# ============================================================
# КОЛОНКИ
# ============================================================

#: (заголовок, ширина, как достать значение из договора)
COLUMNS = (
    (
        "Контрагент",
        44,
        lambda contract: (
            contract.cp.name if contract.cp_id else None
        ),
    ),
    (
        "Тип документа",
        30,
        lambda contract: (
            contract.title.title if contract.title_id else None
        ),
    ),
    (
        "Номер договора",
        24,
        lambda contract: contract.number or "без номера",
    ),
    (
        "Дата договора",
        16,
        lambda contract: contract.date,
    ),
)


def _clean(value):
    """Приводит значение к типу, который понимает openpyxl."""
    if value is None:
        return None

    if isinstance(value, datetime):
        # Excel не хранит часовой пояс, openpyxl на нём падает.
        if value.tzinfo is not None:
            value = value.astimezone().replace(tzinfo=None)
        return value

    return value


# ============================================================
# СБОРКА КНИГИ
# ============================================================

def make_contracts_excel(contracts, subtitle: str = "") -> bytes:
    """
    Собирает книгу со списком договоров.

    contracts — queryset или список объектов Contracts.
    subtitle  — строка под заголовком: сколько записей,
                какой фильтр был применён.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Договоры"

    sheet.sheet_view.showGridLines = False

    width = len(COLUMNS)

    # --------------------------------------------------------
    # Шапка листа
    # --------------------------------------------------------
    sheet.cell(row=1, column=1, value="Договоры").font = FONT_TITLE
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=width,
    )
    sheet.row_dimensions[1].height = 22

    stamp = f"Выгружено {datetime.now():%d.%m.%Y %H:%M}"

    sheet.cell(
        row=2,
        column=1,
        value=f"{subtitle} · {stamp}" if subtitle else stamp,
    ).font = FONT_SUBTITLE

    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=width,
    )

    # --------------------------------------------------------
    # Заголовки
    # --------------------------------------------------------
    header_row = 4

    for index, (title, column_width, _) in enumerate(
        COLUMNS,
        start=1,
    ):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER_THIN
        cell.alignment = ALIGN_CENTER

        sheet.column_dimensions[
            get_column_letter(index)
        ].width = column_width

    sheet.row_dimensions[header_row].height = 26

    # --------------------------------------------------------
    # Строки
    # --------------------------------------------------------
    row_index = header_row

    for offset, contract in enumerate(contracts):
        row_index = header_row + 1 + offset

        for column_index, (_, _, getter) in enumerate(
            COLUMNS,
            start=1,
        ):
            value = _clean(getter(contract))

            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

            cell.border = BORDER_THIN
            cell.font = FONT_NORMAL
            cell.fill = FILL_STRIPE if offset % 2 else FILL_NONE
            cell.alignment = ALIGN_LEFT

            if isinstance(value, (date, datetime)):
                cell.number_format = FMT_DATE
                cell.alignment = ALIGN_CENTER

    count = row_index - header_row

    if not count:
        cell = sheet.cell(
            row=header_row + 1,
            column=1,
            value="Не выбрано ни одного договора",
        )
        cell.font = FONT_SUBTITLE
        sheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=width,
        )
    else:
        # ----------------------------------------------------
        # Итог: только количество. Складывать тут нечего,
        # но знать, сколько строк в файле, полезно.
        # ----------------------------------------------------
        total_row = row_index + 1

        for column_index in range(1, width + 1):
            cell = sheet.cell(row=total_row, column=column_index)
            cell.border = BORDER_THIN
            cell.fill = FILL_TOTAL
            cell.font = FONT_BOLD
            cell.alignment = ALIGN_LEFT

        sheet.cell(
            row=total_row,
            column=1,
            value=f"ИТОГО договоров: {count}",
        )

        sheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(width)}{row_index}"
        )

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer.read()
