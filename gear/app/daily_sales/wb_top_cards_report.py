# gear/app/daily_sales/wb_top_cards_report.py
"""
Отдельная "быстрая" выгрузка по двум самым дорогим разделам --
"3.1 Логистика" и "3.4. Штрафы": аналитическая PDF-записка (по
категориям, брендам и карточкам) и полная построчная детализация в
Excel. Это ОТДЕЛЬНАЯ кнопка в том же меню "Экспорт" -- не трогает и
не заменяет уже готовый "Анализ расходов WB" (wb_expenses_excel.py,
его больше не редактируем).

Данные берём тем же fetch_wb_expenses_detail() из wb_expenses_excel.py
(тот же запрос, та же категоризация под wb_costs.sql) -- одна точка
правды, чтобы суммы здесь и в основном отчёте не могли разойтись.
Категорию (WB "предмет") добираем отдельным лёгким запросом к
inventories.wb_product -- тем же полем subject_name и тем же
фолбэком "Не указана", что уже использует ai_analysis/data.py для
разрезов по категориям. Если этот отдельный запрос не отработает --
это не должно ронять весь отчёт, категория просто уйдёт в
"Не указана" (см. try/except в _fetch_target_df).

Период -- тот же date-picker дашборда (filters.date_picker_id), не
выбран -- значит весь период (см. HISTORY_START в wb_expenses_excel.py).

Два разных пункта меню под два формата:
  -- PDF-записка -- по каждому разделу: короткий текстовый вывод,
     разбивка по категориям и брендам (с полосками -- наглядно,
     кто сколько тянет), разбивка по статьям (с тем, что конкретно
     формирует каждую статью -- карточка, бренд или категория),
     топ-N карточек;
  -- Excel-расшифровка -- построчная детализация rrd_id/nm_id по
     обоим разделам (с категорией отдельной колонкой), без урезки
     топом -- можно самой построить любой свод.

"Кол-во"/"операций" -- число операций (уникальных rrd_id), не штук
товара: в sales.sales_long нет надёжного поля с количеством товара
на уровне одной строки расходов.

"Продано / вернули" здесь намеренно НЕ считаем: расход по возврату
регистрируется у WB датой самой операции (btn/sop_name), а не датой
исходной продажи. За произвольный период в выгрузку попадают только
сами операции возврата -- продажи, которые их породили, могли быть
раньше или позже выбранного периода. Сравнение с продажами за тот же
период дало бы некорректный "процент возврата". Показываем только
факт: сколько операций и на какую сумму -- без сопоставления с
продажами.
"""

from __future__ import annotations

import html as html_lib
from datetime import date, datetime
from io import BytesIO

import pandas as pd
import dash_mantine_components as dmc
from dash import Input, Output, State, dcc, no_update

from openpyxl import Workbook

from . import excel_report_style as style
from .wb_expenses_excel import (
    HISTORY_START,
    ACCOUNT_LABELS,
    fetch_wb_expenses_detail,
    export_menu_item,
    report_menu_group,
)

#: Сколько карточек показывать в топе PDF-записки по каждому разделу.
TOP_N = 15

#: Сколько категорий/брендов показывать в полосках по каждому разделу.
BREAKDOWN_TOP_N = 8

#: Разделы, для которых собираем расшифровку -- те же коды account,
#: что и в WB_EXPENSE_ACCOUNTS (wb_expenses_excel.py).
TARGET_ACCOUNTS = ("3.4. Штрафы", "3.1 Логистика")

#: Короткие имена листов Excel для этих разделов (лимит Excel -- 31
#: символ, и хочется покороче, чем "3.4. Штрафы"/"3.1 Логистика").
SHEET_NAMES = {
    "3.4. Штрафы": "Штрафы",
    "3.1 Логистика": "Логистика",
}


# ================================================================ ID's
TOP_CARDS_PDF_ITEM_ID = "wb-top-cards-pdf-item"
TOP_CARDS_EXCEL_ITEM_ID = "wb-top-cards-excel-item"
TOP_CARDS_PDF_DOWNLOAD_ID = "wb-top-cards-pdf-download"
TOP_CARDS_EXCEL_DOWNLOAD_ID = "wb-top-cards-excel-download"
TOP_CARDS_STATUS_ID = "wb-top-cards-status"
TOP_CARDS_CLICKS_ID = "wb-top-cards-clicks"

TOP_CARDS_MENU_KINDS = (
    ("top_cards_pdf", TOP_CARDS_PDF_ITEM_ID),
    ("top_cards_excel", TOP_CARDS_EXCEL_ITEM_ID),
)


# ================================================================ данные
def _fetch_categories() -> pd.DataFrame:
    """nm_id -> категория (WB "предмет", subject_name).

    Тот же паттерн, что и в ai_analysis/data.py для разрезов по
    категориям: card_id -> subject_name, тот же фолбэк на
    "Не указана". Отдельный лёгкий запрос, а не правка
    fetch_wb_expenses_detail() -- тот отчёт трогать нельзя.
    """
    from conns import get_duckdb_conn_with_opt

    with get_duckdb_conn_with_opt() as con:
        return con.execute(
            """
            SELECT
                card_id AS nm_id,
                COALESCE(NULLIF(TRIM(subject_name), ''), 'Не указана') AS category
            FROM inventories.wb_product
            """
        ).df()


def _fetch_target_df(start_date: str, end_date: str) -> pd.DataFrame:
    df = fetch_wb_expenses_detail(start_date, end_date)
    target = df[df["account"].isin(TARGET_ACCOUNTS)].copy() if not df.empty else df.copy()

    if target.empty:
        # На пустом df ещё нет колонки "category" -- добавляем явно,
        # чтобы дальше по коду (топ карточек, разбивки) можно было
        # рассчитывать на её присутствие независимо от того, есть
        # данные или нет.
        target["category"] = pd.Series(dtype=object)
        return target

    try:
        # merge, а не колонка-заглушка заранее: если завести колонку
        # "category" в target ДО merge, pandas при слиянии по nm_id
        # переименует обе одноимённые колонки в "category_x"/
        # "category_y" вместо одной "category" -- дальше по коду это
        # просто упадёт с KeyError.
        cats = _fetch_categories()
        target = target.merge(cats, on="nm_id", how="left")
    except Exception:
        # Категория -- дополнительный разрез, не основной отчёт: если
        # этот отдельный запрос почему-то не отработал, не роняем всю
        # выгрузку -- просто помечаем как неизвестную категорию.
        target["category"] = None

    target["category"] = target["category"].fillna("Не указана")
    return target


def _sort_by_abs(g: pd.DataFrame, top_n: int) -> pd.DataFrame:
    return g.reindex(g["amount_rub"].abs().sort_values(ascending=False).index).head(top_n)


def _top_cards(df: pd.DataFrame, account: str, top_n: int = TOP_N) -> pd.DataFrame:
    """Топ-N карточек по |сумме| для одного раздела.

    Группируем СТРОГО по nm_id (dropna=False -- см. историю с
    "Без уточнения" у cost_item в wb_expenses_excel.py, pandas
    groupby по умолчанию молча теряет строки с пустым ключом).
    title/brand/category берём через first(): группировка по ним
    вместе с nm_id развалила бы одну и ту же карточку на несколько
    строк при малейшем разночтении в этих полях -- сумма карточки
    оказалась бы в топе заниженной, а сама карточка -- задвоенной.
    """
    sub = df.loc[df["account"] == account]
    if sub.empty:
        return pd.DataFrame(columns=["nm_id", "title", "brand", "category", "ops", "amount_rub"])

    g = (
        sub.groupby("nm_id", dropna=False)
        .agg(
            title=("title", "first"),
            brand=("brand", "first"),
            category=("category", "first"),
            ops=("rrd_id", "nunique"),
            amount_rub=("amount_rub", "sum"),
        )
        .reset_index()
    )
    return _sort_by_abs(g, top_n)


def _breakdown_by(df: pd.DataFrame, account: str, key_col: str, top_n: int = BREAKDOWN_TOP_N) -> pd.DataFrame:
    """Топ-N значений key_col (категория/бренд) по |сумме|."""
    sub = df.loc[df["account"] == account]
    if sub.empty:
        return pd.DataFrame(columns=[key_col, "ops", "amount_rub"])

    g = (
        sub.groupby(key_col, dropna=False)
        .agg(ops=("rrd_id", "nunique"), amount_rub=("amount_rub", "sum"))
        .reset_index()
    )
    return _sort_by_abs(g, top_n)


def _account_totals(df: pd.DataFrame, account: str):
    sub = df.loc[df["account"] == account]
    total = float(sub["amount_rub"].sum()) if not sub.empty else 0.0
    ops = int(sub["rrd_id"].nunique()) if not sub.empty else 0
    cards = int(sub["nm_id"].nunique()) if not sub.empty else 0
    return total, ops, cards


def _dominant_share(sub: pd.DataFrame, key_col: str):
    """Значение key_col с наибольшей |суммой| внутри sub и его доля
    (в %) от суммарного |amount_rub| в sub. (None, 0.0), если считать
    не из чего."""
    if sub.empty:
        return None, 0.0
    g = sub.groupby(key_col, dropna=False)["amount_rub"].sum()
    total_abs = sub["amount_rub"].abs().sum()
    if g.empty or not total_abs:
        return None, 0.0
    top_key = g.abs().idxmax()
    share = abs(g.loc[top_key]) / total_abs * 100
    return top_key, share


#: Сколько карточек показывать списком внутри каждой статьи в
#: таблице "По статьям" -- раньше показывали только топ-1, стало
#: мало: одна карточка не объясняет, куда делись остальные деньги
#: статьи, если явного лидера нет.
STATYA_CARDS_N = 5


def _statya_top_cards(sub: pd.DataFrame, top_n: int = STATYA_CARDS_N) -> pd.DataFrame:
    """Топ-N карточек (по |сумме|) внутри уже отфильтрованного по
    статье sub, с долей (%) каждой от суммы статьи."""
    if sub.empty:
        return pd.DataFrame(columns=["nm_id", "title", "amount_rub", "share"])

    g = (
        sub.groupby("nm_id", dropna=False)
        .agg(title=("title", "first"), amount_rub=("amount_rub", "sum"))
        .reset_index()
    )
    total_abs = sub["amount_rub"].abs().sum()
    g["share"] = (g["amount_rub"].abs() / total_abs * 100) if total_abs else 0.0
    return _sort_by_abs(g, top_n)


def _statya_driver(df: pd.DataFrame, account: str, cost_item):
    """Что конкретно формирует статью: одна карточка, один бренд,
    одна категория -- или просто много похожих операций без явного
    лидера.

    Возвращает (kind, text, top_cards):
      -- kind -- "card" / "brand" / "category" / "spread" / "none";
      -- text -- человеческое объяснение (одна фраза);
      -- top_cards -- топ-STATYA_CARDS_N карточек внутри этой статьи
         (см. _statya_top_cards), чтобы в таблице всегда был виден
         список конкретных товаров, а не только "бренд такой-то".

    Смотрим на карточку первой -- это самый конкретный, самый
    "проверяемый" уровень (можно открыть эту карточку и посмотреть
    руками), и только если она не тянет статью в одиночку --
    поднимаемся до бренда, а затем до категории.
    """
    sub = df.loc[(df["account"] == account) & (df["cost_item"] == cost_item)]
    if sub.empty:
        return "none", "данных нет", pd.DataFrame(columns=["nm_id", "title", "amount_rub", "share"])

    top_cards = _statya_top_cards(sub)
    card_key, card_share = _dominant_share(sub, "nm_id")

    if card_key is not None and card_share >= 40:
        title = sub.loc[sub["nm_id"] == card_key, "title"].iloc[0]
        return "card", f"почти целиком карточка «{_esc(title)}» ({card_share:.0f}% статьи)", top_cards

    brand_key, brand_share = _dominant_share(sub, "brand")
    if brand_key is not None and brand_share >= 40 and str(brand_key) not in ("Не указан", "None", "", "nan"):
        return "brand", f"в основном бренд «{_esc(brand_key)}» ({brand_share:.0f}% статьи)", top_cards

    cat_key, cat_share = _dominant_share(sub, "category")
    if cat_key is not None and cat_share >= 40:
        return "category", f"больше всего у категории «{_esc(cat_key)}» ({cat_share:.0f}% статьи)", top_cards

    if card_key is not None:
        title = sub.loc[sub["nm_id"] == card_key, "title"].iloc[0]
        return "spread", f"разбросано по многим карточкам, заметнее других «{_esc(title)}» ({card_share:.0f}%)", top_cards

    return "none", "явного лидера нет, операции разного рода", top_cards


def _statya_driver_text(df: pd.DataFrame, account: str, cost_item) -> str:
    """Только фраза-объяснение (используется в короткой записке
    наверху раздела, где карточка внутри статьи не нужна)."""
    _, text, _ = _statya_driver(df, account, cost_item)
    return text


def _money(v) -> str:
    v = round(float(v))
    sign = "-" if v < 0 else ""
    return f"{sign}{abs(v):,.0f}".replace(",", " ") + " ₽"


def _period_label(start_date: str, end_date: str) -> str:
    start_label = pd.to_datetime(start_date).strftime("%d.%m.%Y")
    end_label = pd.to_datetime(end_date).strftime("%d.%m.%Y")
    return f"{start_label} – {end_label}" if start_date != end_date else start_label


# ================================================================ PDF-записка
def _esc(text) -> str:
    return html_lib.escape("" if text is None else str(text))


_PDF_CSS = f"""
    @page {{ size: A4; margin: 16mm 14mm; }}
    * {{ box-sizing: border-box; }}
    body {{ font-family: "Helvetica Neue", Arial, sans-serif; color: #1f2a24;
            margin: 0; font-size: 10px; }}
    h1 {{ font-size: 18px; color: #{style.NAVY_3}; margin: 0 0 4px; }}
    .period {{ color: #667; font-size: 9.5px; margin: 0 0 16px; line-height: 1.5; }}
    /* .block НАМЕРЕННО без page-break-inside: avoid -- целиком раздел
       (KPI + текст + категории/бренды + статьи + топ-карточек) выше
       одной страницы A4, и "avoid" на таком большом контейнере не
       спасает от разрыва посреди раздела, а только гонит браузер
       начинать его с чистой страницы -- отсюда была пустая страница
       под заголовком. "avoid" вместо этого точечно на мелких блоках
       ниже (KPI-строка, текст-вывод, строка таблицы статей, блок
       топ-карточек с заголовком) -- там он реально работает. */
    .block {{ margin-bottom: 20px; }}
    /* Второй и следующий разделы (Логистика после Штрафов) -- всегда
       с новой страницы; первый раздел (сразу после заголовка) не
       трогаем -- иначе снова получим пустую первую страницу. */
    .block + .block {{ page-break-before: always; break-before: page; }}
    h2 {{ font-size: 14px; color: #{style.NAVY_3}; border-bottom: 2px solid #{style.NAVY};
          padding-bottom: 4px; margin: 0 0 8px; }}
    .kpi-row {{ display: flex; gap: 8px; margin-bottom: 8px;
                page-break-inside: avoid; break-inside: avoid; }}
    .kpi {{ flex: 1; background: #{style.SURFACE_3}; border-radius: 6px; padding: 7px 9px; }}
    .kpi-label {{ font-size: 8.5px; color: #667; text-transform: uppercase; }}
    .kpi-value {{ font-size: 14px; font-weight: 700; color: #{style.NAVY_3}; }}
    .narrative {{ background: #{style.SURFACE_5}; border-left: 3px solid #{style.NAVY};
                  padding: 7px 10px; margin: 0 0 10px; font-size: 9.5px; line-height: 1.5;
                  page-break-inside: avoid; break-inside: avoid; }}
    .subhead {{ font-size: 10.5px; font-weight: 700; color: #{style.NAVY_3};
                margin: 0 0 5px; }}
    .cols {{ display: flex; gap: 14px; margin-bottom: 10px; }}
    .col {{ flex: 1; min-width: 0; }}
    .bar-row {{ display: flex; align-items: center; gap: 6px; margin-bottom: 4px; }}
    .bar-label {{ width: 34%; font-size: 8.7px; overflow: hidden; text-overflow: ellipsis;
                  white-space: nowrap; }}
    .bar-track {{ flex: 1; background: #{style.SURFACE_3}; border-radius: 3px; height: 8px;
                  overflow: hidden; }}
    .bar-fill {{ height: 100%; background: #{style.NAVY}; border-radius: 3px; }}
    .bar-fill.neg {{ background: #B0745A; }}
    .bar-value {{ width: 30%; text-align: right; font-size: 8.5px; color: #445; white-space: nowrap; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 9.5px; table-layout: fixed; }}
    th {{ background: #{style.NAVY}; color: #fff; text-align: left; padding: 5px 6px; }}
    td {{ padding: 4px 6px; border-bottom: 1px solid #{style.LINE}; overflow: hidden;
          text-overflow: ellipsis; white-space: nowrap; }}
    td.num {{ text-align: right; white-space: nowrap; }}
    col.c-num {{ width: 5%; }}
    col.c-name {{ width: 48%; }}
    col.c-nmid {{ width: 15%; }}
    col.c-ops {{ width: 12%; }}
    col.c-sum {{ width: 20%; }}
    tr:nth-child(even) td {{ background: #{style.SURFACE_2}; }}
    .statya-table {{ margin-bottom: 10px; }}
    .statya-table td {{ white-space: normal; overflow: visible; text-overflow: clip;
                         line-height: 1.4; vertical-align: top; }}
    .statya-table tr {{ page-break-inside: avoid; break-inside: avoid; }}
    .statya-card-line {{ margin-top: 4px; color: #445; font-size: 8.7px; font-weight: 700; }}
    .statya-card-list {{ margin: 2px 0 0; padding-left: 14px; color: #556; font-size: 8.5px;
                          line-height: 1.45; }}
    .statya-card-list li {{ margin-bottom: 1px; }}
    col.sc-name {{ width: 22%; }}
    col.sc-ops {{ width: 9%; }}
    col.sc-sum {{ width: 17%; }}
    col.sc-driver {{ width: 52%; }}
    .empty {{ text-align: center; color: #888; padding: 8px; font-size: 9px; }}
    /* Топ-N карточек: заголовок и таблица держатся вместе -- если не
       помещаются на текущей странице, целиком уходят на следующую,
       а не разрываются между заголовком и таблицей. */
    .top-cards-block {{ page-break-inside: avoid; break-inside: avoid; margin-top: 4px; }}
    .footer {{ margin-top: 10px; font-size: 8px; color: #999; }}
    /* Мелкая светло-серая подпись внизу каждого раздела -- период и
       НДС, чтобы не потерять контекст, если раздел ушёл на отдельную
       страницу (в отличие от .footer это повторяется в каждом
       разделе, а не один раз в конце документа). */
    .section-note {{ margin-top: 10px; font-size: 8.5px; color: #999; }}
"""


def _bar_list_html(rows: pd.DataFrame, label_col: str) -> str:
    if rows.empty:
        return '<div class="empty">Данных нет</div>'

    max_abs = rows["amount_rub"].abs().max() or 1.0
    items = []
    for r in rows.itertuples(index=False):
        value = getattr(r, "amount_rub")
        pct = min(100.0, abs(value) / max_abs * 100) if max_abs else 0.0
        neg_class = " neg" if value < 0 else ""
        label = getattr(r, label_col)
        ops = getattr(r, "ops")
        items.append(f"""
        <div class="bar-row">
            <div class="bar-label" title="{_esc(label)}">{_esc(label)}</div>
            <div class="bar-track"><div class="bar-fill{neg_class}" style="width:{pct:.1f}%"></div></div>
            <div class="bar-value">{_money(value)} · {int(ops)} оп.</div>
        </div>
        """)
    return "".join(items)


def _statya_table_html(df: pd.DataFrame, account: str, top_n: int = BREAKDOWN_TOP_N) -> str:
    """Таблица "по статьям": сумма/операции по каждой статье, плюс
    что конкретно за ней стоит.

    Под фразой-объяснением (карточка/бренд/категория -- см.
    _statya_driver) всегда даём список топ-STATYA_CARDS_N карточек
    внутри этой статьи с суммой и долей каждой: "в основном бренд
    «H&M»" само по себе не говорит, какие именно товары этого бренда
    чаще всего брали в брак и в каком соотношении -- список отвечает
    на это прямо, даже когда лидер и так одна карточка (тогда видно,
    что осталось на остальные).
    """
    rows = _breakdown_by(df, account, "cost_item", top_n=top_n)
    if rows.empty:
        return '<div class="empty">Данных нет</div>'

    trs = []
    for r in rows.itertuples(index=False):
        cost_item = getattr(r, "cost_item")
        kind, text, top_cards = _statya_driver(df, account, cost_item)

        cell = text
        if not top_cards.empty:
            items = "".join(
                f"<li>«{_esc(c.title)}» (nm_id {int(c.nm_id)}) -- "
                f"{_money(c.amount_rub)} ({c.share:.0f}%)</li>"
                for c in top_cards.itertuples(index=False)
            )
            cell += (
                f'<div class="statya-card-line">топ-{len(top_cards)} товаров статьи:</div>'
                f'<ol class="statya-card-list">{items}</ol>'
            )

        trs.append(f"""
        <tr>
            <td>{_esc(cost_item)}</td>
            <td class="num">{int(getattr(r, "ops"))}</td>
            <td class="num">{_money(getattr(r, "amount_rub"))}</td>
            <td>{cell}</td>
        </tr>
        """)

    return f"""
    <table class="statya-table">
        <colgroup>
            <col class="sc-name"><col class="sc-ops"><col class="sc-sum"><col class="sc-driver">
        </colgroup>
        <thead>
            <tr><th>Статья</th><th>Опер.</th><th>Сумма, ₽</th><th>Бренд / карточка</th></tr>
        </thead>
        <tbody>{"".join(trs)}</tbody>
    </table>
    """


def _narrative_html(df: pd.DataFrame, account: str, total: float, ops: int) -> str:
    """Кратко объясняет итог по разделу и показывает основные суммы."""

    def operations_text(n: int) -> str:
        number = f"{n:,}".replace(",", " ")
        if 11 <= n % 100 <= 14:
            ending = "операций"
        elif n % 10 == 1:
            ending = "операция"
        elif 2 <= n % 10 <= 4:
            ending = "операции"
        else:
            ending = "операций"
        return f"{number} {ending}"

    if ops == 0:
        return '<p class="narrative">Операций за период нет.</p>'

    sentences = [
        f"Итог по разделу за период {_money(total)} "
        f"по {operations_text(ops)}."
    ]

    statyas = _breakdown_by(df, account, "cost_item", top_n=1)
    if not statyas.empty:
        s = statyas.iloc[0]
        driver = _statya_driver_text(df, account, s["cost_item"])
        driver = driver.strip().rstrip(".").removeprefix("тут ")

        sentence = (
            f"Самая крупная статья — «{_esc(s['cost_item'])}»: "
            f"{_money(s['amount_rub'])}"
        )
        if driver:
            sentence += f"; {driver}"
        sentences.append(sentence + ".")

    brands = _breakdown_by(df, account, "brand", top_n=1)
    if not brands.empty:
        b = brands.iloc[0]
        sentences.append(
            f"Среди брендов наибольшая сумма приходится на "
            f"«{_esc(b['brand'])}» — {_money(b['amount_rub'])}."
        )

    cats = _breakdown_by(df, account, "category", top_n=1)
    if not cats.empty:
        c = cats.iloc[0]
        share = abs(c["amount_rub"]) / abs(total) * 100 if total else 0.0
        sentences.append(
            f"Среди товарных категорий наибольшая сумма — у "
            f"«{_esc(c['category'])}»: {_money(c['amount_rub'])} "
            f"({share:.0f}% от итога раздела)."
        )

    cards = _top_cards(df, account, top_n=1)
    if not cards.empty:
        k = cards.iloc[0]
        sentences.append(
            f"Дороже всего обошлась карточка «{_esc(k['title'])}» "
            f"(nm_id {int(k['nm_id'])}) — "
            f"{_money(k['amount_rub'])} за {operations_text(int(k['ops']))}."
        )

    return '<p class="narrative">' + " ".join(sentences) + "</p>"


def _account_section_html(account: str, df: pd.DataFrame, period_label: str) -> str:
    label = ACCOUNT_LABELS.get(account, account)
    total, ops, cards_n = _account_totals(df, account)
    narrative = _narrative_html(df, account, total, ops)

    cats_html = _bar_list_html(_breakdown_by(df, account, "category"), "category")
    brands_html = _bar_list_html(_breakdown_by(df, account, "brand"), "brand")
    statya_html = _statya_table_html(df, account)

    top = _top_cards(df, account)
    if top.empty:
        rows_html = '<tr><td colspan="5" class="empty">Операций за период нет</td></tr>'
    else:
        rows_html = "".join(
            f"""
            <tr>
                <td class="num">{i + 1}</td>
                <td>{_esc(r.title)}</td>
                <td class="num">{int(r.nm_id) if pd.notna(r.nm_id) else "--"}</td>
                <td class="num">{int(r.ops)}</td>
                <td class="num">{_money(r.amount_rub)}</td>
            </tr>
            """
            for i, r in enumerate(top.itertuples(index=False))
        )

    return f"""
    <section class="block">
        <h2>{_esc(label)}</h2>
        <div class="kpi-row">
            <div class="kpi">
                <div class="kpi-label">Сумма за период</div>
                <div class="kpi-value">{_money(total)}</div>
            </div>
            <div class="kpi">
                <div class="kpi-label">Операций</div>
                <div class="kpi-value">{ops}</div>
            </div>
            <div class="kpi">
                <div class="kpi-label">Затронуто карточек</div>
                <div class="kpi-value">{cards_n}</div>
            </div>
        </div>
        {narrative}
        <div class="cols">
            <div class="col">
                <div class="subhead">По категориям</div>
                {cats_html}
            </div>
            <div class="col">
                <div class="subhead">По брендам</div>
                {brands_html}
            </div>
        </div>
        <div class="subhead">По статьям</div>
        {statya_html}
        <div class="top-cards-block">
            <div class="subhead">Топ-{TOP_N} карточек: {_esc(label)}</div>
            <table>
                <colgroup>
                    <col class="c-num"><col class="c-name"><col class="c-nmid">
                    <col class="c-ops"><col class="c-sum">
                </colgroup>
                <thead>
                    <tr><th>#</th><th>Наименование</th><th>nm_id</th>
                        <th>Опер.</th><th>Сумма, ₽</th></tr>
                </thead>
                <tbody>{rows_html}</tbody>
            </table>
        </div>
        <div class="section-note">{_esc(period_label)} · без НДС</div>
    </section>
    """


def build_top_cards_html(df: pd.DataFrame, period_label: str) -> str:
    body = "".join(_account_section_html(account, df, period_label) for account in TARGET_ACCOUNTS)

    return f"""<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>Штрафы и Логистика -- записка</title>
<style>{_PDF_CSS}</style>
</head>
<body>
<h1>Штрафы и Логистика: по категориям, брендам и карточкам</h1>
<p class="period">
    {_esc(period_label)} · без НДС · топ-{BREAKDOWN_TOP_N} категорий/брендов/статей и
    топ-{TOP_N} карточек по разделу.
</p>
{body}
<div class="footer">
    Сформировано {datetime.now():%d.%m.%Y %H:%M}
</div>
</body>
</html>
"""


def build_top_cards_pdf(df: pd.DataFrame, period_label: str) -> bytes:
    try:
        from weasyprint import HTML
    except ImportError as exc:
        raise RuntimeError(
            "Для PDF установите WeasyPrint: pip install weasyprint"
        ) from exc

    buffer = BytesIO()
    HTML(string=build_top_cards_html(df, period_label)).write_pdf(buffer)
    return buffer.getvalue()


def build_top_cards_pdf_bytes(start_date: str, end_date: str) -> bytes:
    df = _fetch_target_df(start_date, end_date)
    return build_top_cards_pdf(df, _period_label(start_date, end_date))


# ================================================================ Excel-расшифровка
DETAIL_HEADERS = [
    "Дата", "rrd_id", "nm_id", "Наименование карточки", "Бренд", "Категория",
    "Раздел", "Статья", "Комментарий WB (sop_name)", "Под-операция (btn)",
    "Операция", "Ставка НДС, %", "Сумма без НДС, ₽",
]


def _write_account_detail_sheet(wb, sheet_name, title, subtitle, params, df, account):
    ws = wb.create_sheet(sheet_name)
    sub = df.loc[df["account"] == account].sort_values(["date_from", "rrd_id"])

    col_end = len(DETAIL_HEADERS)
    row = style.write_sheet_header(ws, title, subtitle, params, col_end, landscape=True)
    header_row = row
    style.write_table_header(ws, header_row, 1, DETAIL_HEADERS)

    row = header_row + 1
    first_data_row = row
    numeric_cols = (12, 13)
    formats = {12: style.FMT_PCT, 13: style.FMT_MONEY_DEC}

    for _, r in sub.iterrows():
        ws.cell(row=row, column=1, value=r["date_from"])
        ws.cell(row=row, column=1).number_format = style.FMT_DATE
        ws.cell(row=row, column=2, value=str(r["rrd_id"]))
        ws.cell(row=row, column=3, value=int(r["nm_id"]) if pd.notna(r["nm_id"]) else None)
        ws.cell(row=row, column=4, value=r["title"])
        ws.cell(row=row, column=5, value=r["brand"])
        ws.cell(row=row, column=6, value=r["category"])
        ws.cell(row=row, column=7, value=r["account_label"])
        ws.cell(row=row, column=8, value=r["cost_item"])
        ws.cell(row=row, column=9, value=r["sop_name"])
        ws.cell(row=row, column=10, value=r["btn"])
        ws.cell(row=row, column=11,
                value="Дебет" if r["oper"] == "dt" else "Кредит (возврат)")
        ws.cell(row=row, column=12,
                value=float(r["vat_rate"]) if pd.notna(r["vat_rate"]) else 20.0)
        ws.cell(row=row, column=13, value=round(float(r["amount_rub"]), 2))

        style.style_data_row(ws, row, 1, col_end,
                              numeric_cols=numeric_cols, formats=formats)
        row += 1

    last_row = max(row - 1, first_data_row)
    if sub.empty:
        ws.cell(row=row, column=1, value="Нет операций за выбранный период")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_end)
        last_row = row

    style.apply_zebra(ws, first_data_row, last_row, 1, col_end)
    style.apply_column_dividers(ws, header_row, last_row, 1, col_end)
    style.apply_column_widths(
        ws,
        {
            "Дата": 12, "rrd_id": 14, "nm_id": 12,
            "Наименование карточки": 42, "Бренд": 16, "Категория": 22,
            "Раздел": 26, "Статья": 26,
            "Комментарий WB (sop_name)": 36, "Под-операция (btn)": 36,
            "Операция": 16, "Ставка НДС, %": 12, "Сумма без НДС, ₽": 16,
        },
        header_row,
    )
    style.freeze_table(ws, header_row, first_col=4)
    if not sub.empty:
        style.enable_autofilter(ws, header_row, 1, col_end, last_row)

    return ws


def _write_toc(wb, df, period_label, sheet_entries):
    ws = wb.create_sheet(style.TOC_SHEET_NAME)
    col_end = 6
    style.sheet_base_setup(ws, landscape=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_end)
    c = ws.cell(row=1, column=1, value="Топ карточек: Штрафы и Логистика")
    c.font = style.FONT_SHEET_TITLE
    c.alignment = style.ALIGN_LEFT
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_end)
    c = ws.cell(row=2, column=1, value=period_label)
    c.font = style.FONT_SHEET_SUBTITLE
    c.alignment = style.ALIGN_LEFT

    cards = []
    for account in TARGET_ACCOUNTS:
        total, ops, cards_n = _account_totals(df, account)
        cards.append((
            ACCOUNT_LABELS.get(account, account),
            _money(total),
            f"{ops} операций · {cards_n} карточек",
        ))

    row = 4
    style.write_kpi_cards(ws, row, cards, col_start=1, card_width=3, gap_after=0)
    row += 4

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_end)
    ws.cell(row=row, column=1, value="Листы").font = style.FONT_TABLE_HEADER
    row += 1
    row = style.write_toc_links(ws, row, sheet_entries, col_label=1,
                                 col_desc=2, desc_span=col_end - 1)
    row += 1

    style.write_footer_note(
        ws, row, col_end,
        f"Полная детализация по разделам 'Штрафы' и 'Логистика' -- та же "
        f"категоризация, что и в отчёте 'Анализ расходов WB'. "
        f"Сформировано {datetime.now():%d.%m.%Y %H:%M}.",
    )

    style.apply_column_widths(ws, {}, header_row=row, default_max=30)
    for letter in ("A", "B", "C", "D", "E", "F"):
        if ws.column_dimensions[letter].width is None or ws.column_dimensions[letter].width < 14:
            ws.column_dimensions[letter].width = 18

    return ws


def build_top_cards_excel(start_date: str, end_date: str) -> bytes:
    df = _fetch_target_df(start_date, end_date)
    return _assemble_excel(df, start_date, end_date)


def _assemble_excel(df: pd.DataFrame, start_date: str, end_date: str) -> bytes:
    period_label = _period_label(start_date, end_date)
    params = f"Период: {period_label} · без НДС · сформировано {datetime.now():%d.%m.%Y %H:%M}"

    wb = Workbook()
    wb.remove(wb.active)

    for account in TARGET_ACCOUNTS:
        label = ACCOUNT_LABELS.get(account, account)
        _write_account_detail_sheet(
            wb, SHEET_NAMES[account],
            f"Расходы WB: {label}",
            "Построчная детализация rrd_id/nm_id за выбранный период",
            params,
            df, account,
        )

    sheet_entries = [
        (SHEET_NAMES[account], f"Построчная детализация: {ACCOUNT_LABELS.get(account, account)}")
        for account in TARGET_ACCOUNTS
    ]
    _write_toc(wb, df, params, sheet_entries)

    style.finalize_workbook(
        wb,
        order=[style.TOC_SHEET_NAME] + [SHEET_NAMES[a] for a in TARGET_ACCOUNTS],
    )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


# ================================================================ UI: пункты меню "Экспорт"
def top_cards_menu_items():
    """Вызывается из ui.py -- отчёт "Штрафы / Логистика" как один
    report_menu_group с двумя форматами внутри (та же вложенная
    подача, что и у "Анализ расходов WB" в export_menu_dropdown_items)."""
    return [
        report_menu_group(
            "Штрафы / Логистика",
            [
                export_menu_item(
                    "Записка, PDF",
                    f"По категориям, брендам и статьям, топ-{TOP_N} карточек",
                    TOP_CARDS_PDF_ITEM_ID,
                ),
                export_menu_item(
                    "Расшифровка, Excel",
                    "Построчная детализация rrd_id/nm_id по обоим разделам",
                    TOP_CARDS_EXCEL_ITEM_ID,
                ),
            ],
        ),
    ]


# ================================================================ callback
def register_top_cards_callbacks(app, filters):
    """
    Тот же паттерн, что и у register_wb_expenses_excel_callbacks:
    один callback на оба пункта меню, счётчик нажатий в Store вместо
    ctx.triggered_id (под django-plotly-dash не работает), период --
    тот же date-picker дашборда (период не выбран -- весь период).
    """

    @app.callback(
        Output(TOP_CARDS_PDF_DOWNLOAD_ID, "data"),
        Output(TOP_CARDS_EXCEL_DOWNLOAD_ID, "data"),
        Output(TOP_CARDS_STATUS_ID, "children"),
        Output(TOP_CARDS_CLICKS_ID, "data"),

        Input(TOP_CARDS_PDF_ITEM_ID, "n_clicks"),
        Input(TOP_CARDS_EXCEL_ITEM_ID, "n_clicks"),

        State(filters.date_picker_id, "value"),
        State(TOP_CARDS_CLICKS_ID, "data"),

        prevent_initial_call=True,
    )
    def export_top_cards(pdf_clicks, excel_clicks, date_range, seen_clicks):
        counts = {
            "top_cards_pdf": int(pdf_clicks or 0),
            "top_cards_excel": int(excel_clicks or 0),
        }
        seen = seen_clicks or {}

        changed = [
            kind for kind, _ in TOP_CARDS_MENU_KINDS
            if counts.get(kind, 0) != int(seen.get(kind) or 0)
        ]
        if not changed:
            return no_update, no_update, no_update, counts
        kind = changed[0]

        if date_range and len(date_range) == 2:
            picked_start, picked_end = date_range
            if picked_start and picked_end:
                start_date, end_date = picked_start, picked_end
            elif picked_start or picked_end:
                return no_update, no_update, no_update, counts
            else:
                start_date = HISTORY_START.isoformat()
                end_date = date.today().isoformat()
        else:
            start_date = HISTORY_START.isoformat()
            end_date = date.today().isoformat()

        period_note = (
            f"{start_date} – {end_date}" if start_date != end_date
            else start_date
        )

        try:
            if kind == "top_cards_pdf":
                content = build_top_cards_pdf_bytes(start_date, end_date)
                filename = f"wb_top_cards_{start_date}_{end_date}.pdf"
                status = dmc.Alert(
                    title="Записка готова", color="teal", withCloseButton=True,
                    children=dmc.Text(
                        f"{filename} · {len(content) / 1_000_000:.1f} МБ · "
                        f"период {period_note}",
                        size="sm",
                    ),
                )
                return dcc.send_bytes(content, filename=filename), no_update, status, counts

            content = build_top_cards_excel(start_date, end_date)
            filename = f"wb_top_cards_detail_{start_date}_{end_date}.xlsx"
            status = dmc.Alert(
                title="Файл готов", color="teal", withCloseButton=True,
                children=dmc.Text(
                    f"{filename} · {len(content) / 1_000_000:.1f} МБ · "
                    f"период {period_note}",
                    size="sm",
                ),
            )
            return no_update, dcc.send_bytes(content, filename=filename), status, counts

        except Exception as exc:
            return (
                no_update, no_update,
                dmc.Alert(
                    title="Не удалось собрать файл", color="red", withCloseButton=True,
                    children=dmc.Text(
                        f"{type(exc).__name__}: {exc} (период запроса: {period_note})",
                        size="sm",
                    ),
                ),
                counts,
            )
