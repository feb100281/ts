# gear/app/daily_sales/cards_compliance.py
# =============================================================================
#  Проверка карточек WB: GTIN, ТН ВЭД, декларации/сертификаты.
#
#  В выгрузку попадают только размеры с остатком на выбранную дату
#  (склады WB + FBS + в пути), как в выгрузке остатков.
#  Уровень строки: артикул WB + размер (штрихкоды размера — в одной ячейке).
# =============================================================================

from io import BytesIO
from datetime import date

import numpy as np
import pandas as pd
from dash import dcc, Input, Output, State, no_update
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

from conns import get_duckdb_conn_with_opt
from .excel_report_style import (
    TOC_SHEET_NAME, FONT, TEXT, EXPENSE, FREE, FREE_BG, OCCUPIED_BG, WARN_BG,
    FMT_QTY, FMT_DATE,
    write_sheet_header, write_table_header, style_data_row, apply_zebra,
    apply_column_widths, freeze_table, enable_autofilter, write_kpi_cards,
    write_toc_links, write_footer_note, finalize_workbook,
)
from .ui import (
    STOCKS_DATE_PICKER_ID,
    CARDS_CHECK_BTN_ID,
    CARDS_CHECK_DOWNLOAD_ID,
    CARDS_CHECK_LOADING_ID,
)


CARDS_CHECK_SQL = r"""
WITH wb AS (
    SELECT nm_id, chrt_id,
           SUM(COALESCE(quantity, 0)) AS wb_qty,
           SUM(COALESCE(in_way_to_client, 0)
               + COALESCE(in_way_from_client, 0)) AS transit_qty
    FROM stocks.unpacked_stocks
    WHERE date_from::DATE = $report_date::DATE
    GROUP BY 1, 2
),
fbs AS (
    SELECT nm_id, chrt_id, SUM(COALESCE(quantity, 0)) AS fbs_qty
    FROM stocks.unpacked_fbs_stocks
    WHERE date_from::DATE = $report_date::DATE
      AND nm_id IS NOT NULL
    GROUP BY 1, 2
),
stock AS (
    SELECT
        COALESCE(wb.nm_id, fbs.nm_id)::BIGINT AS nm_id,
        COALESCE(wb.chrt_id, fbs.chrt_id)::BIGINT AS chrt_id,
        COALESCE(wb.wb_qty, 0) AS wb_qty,
        COALESCE(fbs.fbs_qty, 0) AS fbs_qty,
        COALESCE(wb.transit_qty, 0) AS transit_qty
    FROM wb
    FULL OUTER JOIN fbs
        ON wb.nm_id = fbs.nm_id AND wb.chrt_id = fbs.chrt_id
),
stock_pos AS (
    SELECT *, wb_qty + fbs_qty + transit_qty AS total_qty
    FROM stock
    WHERE wb_qty + fbs_qty + transit_qty > 0
),
skus AS (
    SELECT DISTINCT
        c.nm_id::BIGINT AS nm_id,
        c.chrt_id::BIGINT AS chrt_id,
        c.vendor_code, c.brand, c.subject_name, c.title, c.tech_size,
        COALESCE(c.kiz_marked, FALSE) AS kiz_marked,
        TRIM(c.sku) AS barcode,
        TRIM(COALESCE(c.tnved, '')) AS tnved_raw,
        TRIM(c.declaration_number) AS doc_number,
        try_strptime(c.cert_end_date, '%d.%m.%Y')::DATE AS doc_end_date
    FROM cards.unpacked_cards c
),
sku_check AS (
    SELECT
        *,
        CASE
            WHEN COALESCE(barcode, '') = '' THEN 'нет штрихкода'
            WHEN NOT regexp_matches(barcode, '^[0-9]+$') THEN 'есть не только цифры'
            WHEN length(barcode) NOT IN (8, 12, 13, 14)
                THEN 'длина ' || length(barcode) || ' знаков'
            WHEN length(barcode) = 13 AND starts_with(barcode, '2')
                THEN 'внутренний код (на «2»), не GTIN'
            WHEN (10 - list_sum(list_transform(
                    range(1, length(barcode)),
                    i -> CAST(substr(barcode, length(barcode) - i, 1) AS INT)
                         * CASE WHEN i % 2 = 1 THEN 3 ELSE 1 END)) % 10) % 10
                 <> CAST(right(barcode, 1) AS INT)
                THEN 'неверная контрольная цифра'
        END AS gtin_issue_one
    FROM skus
),
-- одна строка на размер: остаток размера не дублируется по штрихкодам
sizes AS (
    SELECT
        nm_id, chrt_id,
        any_value(vendor_code) AS vendor_code,
        any_value(brand) AS brand,
        any_value(subject_name) AS subject_name,
        any_value(title) AS title,
        any_value(tech_size) AS tech_size,
        bool_or(kiz_marked) AS kiz_marked,
        string_agg(DISTINCT barcode, ', ') AS barcode,
        string_agg(
            CASE WHEN gtin_issue_one IS NOT NULL
                 THEN COALESCE(NULLIF(barcode, ''), '—') || ': ' || gtin_issue_one END,
            '; ') AS gtin_issue,
        any_value(tnved_raw) AS tnved_raw,
        any_value(doc_number) AS doc_number,
        max(doc_end_date) AS doc_end_date
    FROM sku_check
    GROUP BY nm_id, chrt_id
),
joined AS (
    SELECT s.*, z.* EXCLUDE (nm_id, chrt_id),
           regexp_replace(COALESCE(z.tnved_raw, ''), '[^0-9]', '', 'g') AS tnved
    FROM stock_pos s
    LEFT JOIN sizes z ON z.nm_id = s.nm_id AND z.chrt_id = s.chrt_id
)
SELECT
    * EXCLUDE (gtin_issue, tnved),
    CASE WHEN vendor_code IS NULL THEN 'карточка не найдена' ELSE gtin_issue END
        AS gtin_issue,
    CASE
        WHEN vendor_code IS NULL THEN 'карточка не найдена'
        WHEN tnved = '' THEN 'не заполнен'
        WHEN length(tnved) <> 10 THEN length(tnved) || ' цифр вместо 10'
    END AS tnved_issue,
    CASE
        WHEN vendor_code IS NULL THEN 'карточка не найдена'
        WHEN COALESCE(doc_number, '') = '' THEN 'нет номера документа'
        WHEN NOT regexp_matches(upper(doc_number), 'ЕАЭС|EAЭС|EAEU|RU')
            THEN 'формат номера не похож на ЕАЭС / RU'
        WHEN doc_end_date IS NULL THEN 'не указан срок действия'
        WHEN doc_end_date < $report_date::DATE
            THEN 'срок истёк ' || strftime(doc_end_date, '%d.%m.%Y')
    END AS doc_issue
FROM joined
ORDER BY nm_id, tech_size
"""


def get_cards_check_data(report_date) -> pd.DataFrame:
    """Размеры с остатком на дату и результаты проверок."""
    with get_duckdb_conn_with_opt() as con:
        df = con.execute(
            CARDS_CHECK_SQL,
            {"report_date": pd.to_datetime(report_date).date()},
        ).df()

    df["issues"] = (
        df[["gtin_issue", "tnved_issue", "doc_issue"]].notna().sum(axis=1)
    )
    return df


# ------------------------------------------------------------------ Excel

COLUMNS = [
    ("nm_id", "Артикул WB", None),
    ("vendor_code", "Артикул продавца", None),
    ("brand", "Бренд", None),
    ("subject_name", "Предмет", None),
    ("title", "Наименование", None),
    ("tech_size", "Размер", None),
    ("total_qty", "Остаток, шт", FMT_QTY),
    ("wb_qty", "Склады WB", FMT_QTY),
    ("fbs_qty", "FBS", FMT_QTY),
    ("transit_qty", "В пути", FMT_QTY),
    ("kiz_marked", "Маркировка ЧЗ", None),
    ("barcode", "Штрихкоды", None),
    ("gtin_issue", "GTIN: проблема", None),
    ("tnved_raw", "ТН ВЭД", None),
    ("tnved_issue", "ТН ВЭД: проблема", None),
    ("doc_number", "Декларация / сертификат", None),
    ("doc_end_date", "Действует до", FMT_DATE),
    ("doc_issue", "Документ: проблема", None),
    ("issues", "Проблем", FMT_QTY),
]
ISSUE_KEYS = {"gtin_issue", "tnved_issue", "doc_issue"}
WIDTHS = {"Наименование": 42, "Декларация / сертификат": 30,
          "GTIN: проблема": 30, "ТН ВЭД: проблема": 20,
          "Документ: проблема": 30, "Штрихкоды": 24}

SHEET_ISSUES = "Карточки с проблемами"
SHEET_ALL = "Все с остатком"


def _cell_value(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if value is pd.NaT:
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, (bool, np.bool_)):
        return "да" if value else "нет"
    return value


def _fit_width(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _write_table(ws, df, title, subtitle, params):
    ncols = len(COLUMNS)
    hdr = write_sheet_header(ws, title, subtitle, params, ncols)
    write_table_header(ws, hdr, 1, [c[1] for c in COLUMNS])

    numeric = {i + 1 for i, c in enumerate(COLUMNS) if c[2]}
    formats = {i + 1: c[2] for i, c in enumerate(COLUMNS) if c[2]}
    bad_fill = PatternFill("solid", fgColor=OCCUPIED_BG)
    bad_font = Font(name=FONT, size=10, color=EXPENSE)

    row = hdr + 1
    for rec in df.to_dict("records"):
        for j, (key, _label, _fmt) in enumerate(COLUMNS, start=1):
            ws.cell(row=row, column=j, value=_cell_value(rec.get(key)))
        style_data_row(ws, row, 1, ncols, numeric_cols=numeric, formats=formats)
        for j, (key, _label, _fmt) in enumerate(COLUMNS, start=1):
            if key in ISSUE_KEYS and rec.get(key):
                cell = ws.cell(row=row, column=j)
                cell.fill = bad_fill
                cell.font = bad_font
        row += 1

    last = max(row - 1, hdr + 1)
    apply_zebra(ws, hdr + 1, last, 1, ncols)
    apply_column_widths(ws, WIDTHS, hdr)
    freeze_table(ws, hdr, first_col=3)
    enable_autofilter(ws, hdr, 1, ncols, last)
    _fit_width(ws)


def make_cards_check_excel(df, report_date) -> bytes:
    report_date = pd.to_datetime(report_date).date()
    d = report_date.strftime("%d.%m.%Y")
    issues = df[df["issues"] > 0].sort_values(
        ["issues", "total_qty"], ascending=[False, False])

    wb = Workbook()
    ws = wb.active
    ws.title = TOC_SHEET_NAME
    write_sheet_header(
        ws, "ПРОВЕРКА КАРТОЧЕК WB",
        "GTIN, ТН ВЭД, декларации и сертификаты · только товары с остатком",
        "Остатки на %s · склады WB + FBS + в пути" % d, 10, landscape=False)
    # на оглавлении кнопка возврата не нужна
    ws.unmerge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    for col in (1, 2):
        cell = ws.cell(row=1, column=col)
        cell.value, cell.hyperlink = None, None
        cell.fill, cell.border = PatternFill(fill_type=None), Border()

    cards_total = df["nm_id"].nunique()
    cards_bad = issues["nm_id"].nunique()
    qty_bad = issues["total_qty"].sum()
    write_kpi_cards(ws, 7, [
        ("КАРТОЧЕК С ОСТАТКОМ", f"{cards_total:,}".replace(",", " "),
         f"остаток {int(df['total_qty'].sum()):,} шт".replace(",", " ")),
        ("С ПРОБЛЕМАМИ", f"{cards_bad:,}".replace(",", " "),
         f"остаток {int(qty_bad):,} шт".replace(",", " ")),
        ("GTIN", str(df.loc[df["gtin_issue"].notna(), "nm_id"].nunique()), "карточек"),
        ("ТН ВЭД", str(df.loc[df["tnved_issue"].notna(), "nm_id"].nunique()), "карточек"),
        ("ДЕКЛАРАЦИИ / СЕРТИФИКАТЫ",
         str(df.loc[df["doc_issue"].notna(), "nm_id"].nunique()), "карточек"),
    ], col_start=1, card_width=2)

    row = write_toc_links(ws, 12, [
        (SHEET_ISSUES, "Размеры, где есть хотя бы одна проблема"),
        (SHEET_ALL, "Все размеры с остатком на дату — для сверки"),
    ], col_label=1, col_desc=2, desc_span=9)

    notes = [
        "GTIN — штрихкод из 8/12/13/14 цифр с верной контрольной цифрой (GS1). "
        "Код из 13 цифр на «2» — внутренний (в т. ч. сгенерированный WB), не GTIN. "
        "Для товаров с маркировкой ЧЗ GTIN должен быть зарегистрирован в Честном знаке.",
        "ТН ВЭД — ровно 10 цифр.",
        "Документ — номер декларации/сертификата заполнен, похож на ЕАЭС/RU, "
        "срок действия не истёк на дату остатков.",
        "Наличие GTIN в Честном знаке и документа в реестре ФСА по этим данным "
        "не проверяется — только заполненность и формат.",
    ]
    row += 1
    for text in notes:
        write_footer_note(ws, row, 10, text)
        ws.cell(row=row, column=1).alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 28
        row += 1
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["A"].width = 26
    _fit_width(ws)

    params = "Остатки на %s · уровень: артикул WB + размер" % d
    _write_table(wb.create_sheet(SHEET_ISSUES), issues,
                 "КАРТОЧКИ С ПРОБЛЕМАМИ",
                 "Сначала — больше проблем и больше остаток", params)
    _write_table(wb.create_sheet(SHEET_ALL),
                 df.sort_values(["nm_id", "tech_size"]),
                 "ВСЕ КАРТОЧКИ С ОСТАТКОМ", "Для сверки", params)

    finalize_workbook(wb, [TOC_SHEET_NAME, SHEET_ISSUES, SHEET_ALL])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------- callback

def register_cards_check_callbacks(app):
    @app.callback(
        Output(CARDS_CHECK_DOWNLOAD_ID, "data"),
        Output(CARDS_CHECK_LOADING_ID, "children"),
        Input(CARDS_CHECK_BTN_ID, "n_clicks"),
        State(STOCKS_DATE_PICKER_ID, "value"),
        prevent_initial_call=True,
    )
    def export_cards_check(n_clicks, report_date):
        if not n_clicks:
            return no_update, no_update
        report_date = pd.to_datetime(report_date).date() if report_date else (
            date.today() - pd.Timedelta(days=1))
        df = get_cards_check_data(report_date)
        content = make_cards_check_excel(df, report_date)
        return (
            dcc.send_bytes(content, filename="cards_check_%s.xlsx"
                           % report_date.strftime("%Y-%m-%d")),
            "",
        )
