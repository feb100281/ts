# gear/app/daily_sales/excel_report_style.py
"""
Фирменный стиль отчётов в Excel — ТРЕНДСЕТТЕР, внутренний стандарт.

Реализует часть стандарта, которая не требует нативных сводных
Excel (openpyxl не умеет создавать их с нуля — нужен отдельный
файл-скелет, которого в проекте пока нет): палитру, шрифты,
числовые форматы, анатомию листа, зебру, разделители колонок,
уровни вложенности, оглавление с KPI-карточками и навигацией.

Один раз собранный модуль — переиспользуется любой новой
Excel-выгрузкой дашбордов, не только "Анализ расходов WB".
"""

from __future__ import annotations

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties


# ================================================================ палитра
NAVY = "2F6656"        # основной: шапки, плашки разделов
NAVY_2 = "3D7A67"      # светлее: чередование годов
NAVY_3 = "1F5E4E"      # темнее: итоги за период, ссылки
ACCENT = "1F5E4E"      # линия под заголовком листа

TEXT = "1F1F1F"        # основной текст
TEXT_2 = "4A4A4A"      # второстепенный текст
MUTED = "8A8A8A"       # коды статей, подписи, сноски

PAGE = "F7F7F7"        # фон страницы
SURFACE = "FFFFFF"     # фон карточек и таблиц
SURFACE_2 = "F3F8F6"   # чётные строки
SURFACE_3 = "EDF5F1"   # подраздел / колонки "Итого"
SURFACE_4 = "E7F1ED"   # раздел, строки итогов, подзаголовки, кнопка возврата
SURFACE_5 = "FAFCFB"   # почти белый с зелёным подтоном
TOTAL_ROW = "E7F1ED"
ZEBRA_ROW = "F7F7F7"

LINE = "D9D9D9"
LINE_STRONG = "BFBFBF"

INCOME = "3C4043"      # доходы и количества
EXPENSE = "7B4437"     # расходы — приглушённый коричневый, читается
                       # как минус, но не кричит красным

FREE, FREE_BG = "2F6656", "E7F1ED"
INFO, INFO_BG = "2F75B5", "EAF2FB"
WARN, WARN_BG = "9A6100", "FDF3DE"
OCCUPIED_BG = "F6E9E4"

# ================================================================ шрифт
# Helvetica Light — управленческий пакет; на сводных и дашбордных
# выгрузках стандарт явно требует Roboto (есть не у всех машин
# Helvetica Light, Excel тогда молча подставляет свой шрифт).
FONT = "Roboto"

# ================================================================ форматы чисел
FMT_MONEY = '#,##0;(#,##0);"—"'
FMT_MONEY_DEC = '#,##0.00;(#,##0.00);"—"'
FMT_PCT = '#,##0.0" %";(#,##0.0" %");"—"'
FMT_QTY = '#,##0;(#,##0);"—"'
FMT_PRICE = '#,##0.00;(#,##0.00);"—"'
FMT_DATE = "dd.mm.yyyy"
FMT_DATETIME = "dd.mm.yyyy hh:mm"

# ================================================================ навигация
TOC_SHEET_NAME = "Оглавление"

# Уровни вложенности иерархических таблиц: заливка + жирность.
# Ниже — работает зебра, без явной заливки.
LEVEL_STYLES = [
    (SURFACE_4, True),   # 0 — раздел
    (SURFACE_3, True),   # 1 — подраздел
    (SURFACE_2, False),  # 2 — статья
]


# ================================================================ базовые объекты стиля
def _side(color):
    return Side(style="thin", color=color)


BORDER_THIN = Border(
    left=_side(LINE), right=_side(LINE),
    top=_side(LINE), bottom=_side(LINE),
)

FILL_HEADER = PatternFill("solid", fgColor=NAVY)
FILL_ZEBRA = PatternFill("solid", fgColor=ZEBRA_ROW)
FILL_TOTAL = PatternFill("solid", fgColor=TOTAL_ROW)
FILL_TOC_BTN = PatternFill("solid", fgColor=SURFACE_4)
FILL_NONE = PatternFill(fill_type=None)

FONT_SHEET_TITLE = Font(name=FONT, size=14, bold=True, color=TEXT)
FONT_SHEET_SUBTITLE = Font(name=FONT, size=9, color=TEXT_2)
FONT_SHEET_PARAMS = Font(name=FONT, size=9, color=MUTED)
FONT_TABLE_HEADER = Font(name=FONT, size=10, bold=True, color=SURFACE)
FONT_DATA = Font(name=FONT, size=10, color=TEXT)
FONT_TOTAL = Font(name=FONT, size=10, bold=True, color=TEXT)
FONT_FOOTNOTE = Font(name=FONT, size=8, color=MUTED)
FONT_TOC_LINK = Font(name=FONT, size=10, bold=True, color=NAVY_3, underline=None)
FONT_KPI_LABEL = Font(name=FONT, size=8, bold=True, color=TEXT_2)
FONT_KPI_VALUE = Font(name=FONT, size=16, bold=True, color=NAVY_3)
FONT_KPI_SUB = Font(name=FONT, size=8, color=MUTED)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


# ================================================================ зебра и разделители
def apply_zebra(ws, first_row, last_row, first_col, last_col):
    """
    Красит только чётные строки диапазона данных и только те
    ячейки, у которых заливки ещё нет — итоги, шапка и колонки
    "Итого" уже покрашены по смыслу, зебра их не трогает.
    """
    for row in range(first_row, last_row + 1):
        if row % 2:
            continue

        for col in range(first_col, last_col + 1):
            cell = ws.cell(row=row, column=col)

            if cell.fill.fill_type:
                continue

            cell.fill = FILL_ZEBRA


def apply_column_dividers(ws, first_row, last_row, first_col, last_col):
    """Тонкие вертикальные линии между колонками, поверх остальных границ."""
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col):
            cell = ws.cell(row=row, column=col)
            old = cell.border

            cell.border = Border(
                left=old.left,
                right=_side(LINE),
                top=old.top,
                bottom=old.bottom,
            )


# ================================================================ анатомия листа
def sheet_base_setup(ws, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetupPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True


def back_to_toc(ws, row=1, col_start=1, col_end=2):
    """Светло-зелёная плашка с зелёным текстом-ссылкой на оглавление."""
    ws.merge_cells(start_row=row, start_column=col_start,
                    end_row=row, end_column=col_end)

    c = ws.cell(row=row, column=col_start, value="←  Оглавление")
    c.hyperlink = f"#'{TOC_SHEET_NAME}'!A1"
    c.font = FONT_TOC_LINK
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_TOC_BTN
        cell.border = Border(bottom=_side(LINE))

    ws.row_dimensions[row].height = 16


def write_sheet_header(ws, title, subtitle, params, col_end, landscape=True):
    """
    Строки 1—5: кнопка "Оглавление", заголовок, подзаголовок,
    параметры, фирменная зелёная линия. Возвращает номер первой
    свободной строки (6) — с неё уже идёт шапка таблицы.
    """
    sheet_base_setup(ws, landscape=landscape)
    back_to_toc(ws, row=1, col_start=1, col_end=min(2, col_end))

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_end)
    c = ws.cell(row=2, column=1, value=title)
    c.font = FONT_SHEET_TITLE
    c.alignment = ALIGN_LEFT
    ws.row_dimensions[2].height = 24

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_end)
    c = ws.cell(row=3, column=1, value=subtitle)
    c.font = FONT_SHEET_SUBTITLE
    c.alignment = ALIGN_LEFT
    ws.row_dimensions[3].height = 14

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=col_end)
    c = ws.cell(row=4, column=1, value=params)
    c.font = FONT_SHEET_PARAMS
    c.alignment = ALIGN_LEFT
    ws.row_dimensions[4].height = 14

    for col in range(1, col_end + 1):
        cell = ws.cell(row=5, column=col)
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[5].height = 6

    return 6


# ================================================================ таблица
def write_table_header(ws, row, col_start, headers):
    """Шапка таблицы: полужирный белый текст на заливке NAVY."""
    for offset, label in enumerate(headers):
        col = col_start + offset
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = FILL_HEADER
        cell.font = FONT_TABLE_HEADER
        cell.border = BORDER_THIN
        cell.alignment = ALIGN_CENTER

    ws.row_dimensions[row].height = 26


def style_data_row(ws, row, col_start, col_end, numeric_cols=(), formats=None):
    """
    Базовый стиль строки данных: обычный шрифт, тонкая граница,
    числа вправо с нужным форматом, текст влево.
    """
    formats = formats or {}

    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = FONT_DATA
        cell.border = BORDER_THIN

        if col in numeric_cols:
            cell.alignment = ALIGN_RIGHT
            if col in formats:
                cell.number_format = formats[col]
        else:
            cell.alignment = ALIGN_LEFT


def style_level_row(ws, row, col_start, col_end, level, bold_value_cols=()):
    """
    Строка иерархии (раздел / подраздел / статья) — заливка и
    жирность по уровню, ниже статьи заливки нет, работает зебра.
    """
    if level < len(LEVEL_STYLES):
        fill_color, bold = LEVEL_STYLES[level]
    else:
        fill_color, bold = None, False

    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER_THIN

        if fill_color:
            cell.fill = PatternFill("solid", fgColor=fill_color)

        if bold or col in bold_value_cols:
            cell.font = Font(name=FONT, size=10, bold=True, color=TEXT)


def style_total_row(ws, row, col_start, col_end):
    """Строка ИТОГО: заливка E7F1ED, полужирный, сверху medium, снизу double."""
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = FILL_TOTAL
        cell.font = FONT_TOTAL
        cell.border = Border(
            left=_side(LINE), right=_side(LINE),
            top=Side(style="medium", color=NAVY),
            bottom=Side(style="double", color=NAVY),
        )


def set_number_format(ws, row, col, fmt):
    ws.cell(row=row, column=col).number_format = fmt


def apply_column_widths(ws, widths_by_header, header_row, default_max=38):
    """
    Фиксированная ширина для повторяющихся заголовков, для
    остальных — по самому длинному значению в колонке с потолком.
    """
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        header = ws.cell(row=header_row, column=col).value

        if header in widths_by_header:
            ws.column_dimensions[letter].width = widths_by_header[header]
            continue

        max_len = len(str(header)) if header else 0
        for row in range(header_row + 1, ws.max_row + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None:
                max_len = max(max_len, len(str(value)))

        ws.column_dimensions[letter].width = min(max_len + 2, default_max)


def freeze_table(ws, header_row, first_col=2):
    """
    Заморозка: шапка таблицы и колонки левее first_col.

    Координата собирается как строка (а не через ws.cell(...)):
    если якорная строка данных вдруг целиком объединена (например,
    "нет операций за период" на пустой выгрузке), ws.cell() в этом
    месте вернёт MergedCell, а openpyxl не умеет разбирать такую
    ячейку как координату freeze_panes -- падает с
    "'MergedCell' object is not iterable".
    """
    letter = get_column_letter(first_col)
    ws.freeze_panes = f"{letter}{header_row + 1}"


def enable_autofilter(ws, header_row, col_start, col_end, last_row):
    start_letter = get_column_letter(col_start)
    end_letter = get_column_letter(col_end)
    ws.auto_filter.ref = f"{start_letter}{header_row}:{end_letter}{last_row}"


# ================================================================ оглавление
def write_kpi_cards(ws, row, cards, col_start=1, card_width=2, gap_after=0):
    """
    KPI-карточки в ряд. cards — список (label, value, sub).
    Первая карточка выделяется заливкой SURFACE_4.
    """
    col = col_start

    for i, (label, value, sub) in enumerate(cards):
        c0, c1 = col, col + card_width - 1

        for r, font, text, align in (
            (row, FONT_KPI_LABEL, label, ALIGN_LEFT_WRAP),
            (row + 1, FONT_KPI_VALUE, value, ALIGN_LEFT),
            (row + 2, FONT_KPI_SUB, sub, ALIGN_LEFT),
        ):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
            cell = ws.cell(row=r, column=c0, value=text)
            cell.font = font
            cell.alignment = Alignment(
                horizontal=align.horizontal, vertical="center", indent=1,
                wrap_text=(r == row),
            )

        for r in range(row, row + 3):
            for cc in range(c0, c1 + 1):
                cell = ws.cell(row=r, column=cc)
                if i == 0:
                    cell.fill = PatternFill("solid", fgColor=SURFACE_4)
                cell.border = Border(
                    top=Side(style="thin", color=NAVY) if r == row else None,
                    left=_side(LINE) if cc == c0 else None,
                    right=_side(LINE) if cc == c1 else None,
                    bottom=_side(LINE) if r == row + 2 else None,
                )

        ws.row_dimensions[row].height = 16
        ws.row_dimensions[row + 1].height = 26
        ws.row_dimensions[row + 2].height = 14

        col = c1 + 1 + gap_after

    return col


def write_toc_links(ws, row, entries, col_label=1, col_desc=2, desc_span=3):
    """
    entries — список (sheet_name, description). Возвращает строку
    сразу после последней ссылки.
    """
    for name, description in entries:
        c = ws.cell(row=row, column=col_label, value=f"›  {name}")
        c.hyperlink = f"#'{name}'!A1"
        c.font = FONT_TOC_LINK
        c.fill = PatternFill("solid", fgColor=SURFACE_4)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c.border = Border(bottom=_side(LINE))

        ws.merge_cells(
            start_row=row, start_column=col_desc,
            end_row=row, end_column=col_desc + desc_span - 1,
        )
        d = ws.cell(row=row, column=col_desc, value=description)
        d.font = FONT_SHEET_SUBTITLE
        d.alignment = Alignment(horizontal="left", vertical="center",
                                 wrap_text=True, indent=1)
        d.border = Border(bottom=_side(LINE))

        ws.row_dimensions[row].height = 28
        row += 1

    return row


def write_footer_note(ws, row, col_end, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_end)
    c = ws.cell(row=row, column=1, value=text)
    c.font = FONT_FOOTNOTE
    c.alignment = ALIGN_LEFT


# ================================================================ финализация книги
def finalize_workbook(wb, order):
    """
    Задаёт порядок листов явно, снимает "выделенность" со всех
    листов (иначе Excel откроет книгу в режиме группового
    редактирования и правка на одном листе уйдёт сразу на все
    выделенные), делает активным первый лист.
    """
    ordered = [wb[name] for name in order if name in wb.sheetnames]
    ordered += [ws for ws in wb.worksheets if ws not in ordered]
    wb._sheets = ordered

    for sheet in wb.worksheets:
        for view in sheet.views.sheetView:
            view.tabSelected = False

    wb.active = 0
