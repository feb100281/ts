# gear/app/daily_sales/fbs_orders/excel.py
"""
Excel-отчёт по заказам FBS.

Оформление — в фирменном стиле отчётов проекта: тёмно-зелёная
шапка, Roboto Light 10, зебра, отрицательные значения в скобках.
Каждый лист самодостаточен: заголовок, на что смотреть,
на какой момент данные.
"""

from __future__ import annotations

from datetime import datetime, time
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import (
    CARGO_TYPE_NAMES,
    SLA_LIMIT_HOURS,
    SLA_WARNING_HOURS,
    SUPPLIER_STATUS_NAMES,
    TOTAL_ROW_LABEL,
    WB_STATUS_NAMES,
)


# ============================================================
# ПАЛИТРА
# ============================================================

DARK_GREEN = "2F6656"
LIGHT_GREEN = "E7F1ED"
STRIPE = "F6F6F6"
BORDER_GRAY = "D9D9D9"
WHITE = "FFFFFF"
TEXT = "050505"
MUTED = "7A7A7A"

WARN_BG = "FFF4E6"
WARN_TEXT = "B35309"

ALERT_BG = "FDECEC"
ALERT_TEXT = "A33A3A"

FONT_NAME = "Roboto Light"

# Палитра графиков — та же, что на дашборде.
CHART_BLUE = "2A78D6"
CHART_ORANGE = "EB6834"
CHART_CRITICAL = "D03B3B"
CHART_SERIOUS = "EC835A"
CHART_GRID = "E1E0D9"


# ============================================================
# ФОРМАТЫ ЧИСЕЛ
# ============================================================

FMT_INT = '#,##0;(#,##0);"—"'
FMT_MONEY = '#,##0.00;(#,##0.00);"—"'
FMT_HOURS = '#,##0.0;(#,##0.0);"—"'
FMT_PCT = '#,##0.00"%";(#,##0.00"%");"—"'
FMT_DATE = "dd.mm.yyyy"
FMT_DATETIME = "dd.mm.yyyy hh:mm"


# ============================================================
# ЭЛЕМЕНТЫ ОФОРМЛЕНИЯ
# ============================================================

_thin = Side(style="thin", color=BORDER_GRAY)
BORDER_THIN = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

FONT_TITLE = Font(name=FONT_NAME, size=13, bold=True, color=DARK_GREEN)
FONT_SUBTITLE = Font(name=FONT_NAME, size=9, color=MUTED)
FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
FONT_NORMAL = Font(name=FONT_NAME, size=10, color=TEXT)
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True, color=TEXT)
FONT_ALERT = Font(name=FONT_NAME, size=10, bold=True, color=ALERT_TEXT)
FONT_WARN = Font(name=FONT_NAME, size=10, color=WARN_TEXT)

FILL_HEADER = PatternFill("solid", fgColor=DARK_GREEN)
FILL_TOTAL = PatternFill("solid", fgColor=LIGHT_GREEN)
FILL_STRIPE = PatternFill("solid", fgColor=STRIPE)
FILL_ALERT = PatternFill("solid", fgColor=ALERT_BG)
FILL_WARN = PatternFill("solid", fgColor=WARN_BG)
FILL_NONE = PatternFill(fill_type=None)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_HEADER = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)


# ============================================================
# ПОДГОТОВКА ЗНАЧЕНИЙ
# ============================================================

def _clean(value):
    """
    Приводит значение к типу, который openpyxl умеет писать.

    Отдельно снимается часовой пояс: WB отдаёт даты со
    смещением, а Excel формат с таймзоной не поддерживает
    и openpyxl на такой дате падает.
    """
    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()

    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.astimezone().replace(tzinfo=None)
        return value

    if isinstance(value, time):
        if value.tzinfo is not None:
            value = value.replace(tzinfo=None)
        return value

    if hasattr(value, "item"):
        try:
            return value.item()
        except (ValueError, AttributeError):
            pass

    return value


def _supplier_status_name(value):
    if value in (None, ""):
        return "—"
    return SUPPLIER_STATUS_NAMES.get(value, value)


def _wb_status_name(value):
    if value in (None, ""):
        return "—"
    return WB_STATUS_NAMES.get(value, value)


def _cargo_type_name(value):
    value = _clean(value)
    if value is None:
        return "—"
    try:
        return CARGO_TYPE_NAMES.get(int(value), str(value))
    except (TypeError, ValueError):
        return str(value)


# ============================================================
# ОПИСАНИЕ КОЛОНКИ
# ============================================================

class Column:
    """
    Одна колонка таблицы в отчёте.

    field   — имя колонки в DataFrame;
    title   — заголовок в Excel;
    fmt     — числовой формат;
    width   — ширина;
    total   — нужно ли считать итог ('sum' / 'avg' / None);
    convert — функция преобразования значения.
    """

    def __init__(
        self,
        field,
        title,
        fmt=None,
        width=16,
        total=None,
        convert=None,
        weight_field=None,
    ):
        self.field = field
        self.title = title
        self.fmt = fmt
        self.width = width
        self.total = total
        self.convert = convert
        self.weight_field = weight_field

    def value(self, row):
        raw = _clean(row.get(self.field))

        if self.convert is not None:
            return self.convert(raw)

        return raw


# ============================================================
# ЗАПИСЬ ТАБЛИЦЫ
# ============================================================

def _write_sheet_header(ws, title, subtitle, width):
    ws.sheet_view.showGridLines = False

    ws.cell(row=1, column=1, value=title).font = FONT_TITLE
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(width, 1),
    )
    ws.row_dimensions[1].height = 22

    ws.cell(row=2, column=1, value=subtitle).font = FONT_SUBTITLE
    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=max(width, 1),
    )
    ws.row_dimensions[2].height = 14

    return 4


def _write_table(
    ws,
    df,
    columns,
    start_row,
    total_label_field=None,
    highlight=None,
    autofilter=True,
    freeze=True,
):
    """
    Пишет таблицу и возвращает номер строки после неё.

    highlight — функция от строки, возвращающая 'alert',
    'warn' или None. Ею подсвечиваются просроченные заказы.
    """
    header_row = start_row

    for index, column in enumerate(columns, start=1):
        cell = ws.cell(
            row=header_row,
            column=index,
            value=column.title,
        )
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER_THIN
        cell.alignment = ALIGN_HEADER

        ws.column_dimensions[
            get_column_letter(index)
        ].width = column.width

    ws.row_dimensions[header_row].height = 28

    records = (
        df.to_dict(orient="records")
        if df is not None and not df.empty
        else []
    )

    row_index = header_row

    for offset, record in enumerate(records):
        row_index = header_row + 1 + offset

        mark = highlight(record) if highlight else None

        for column_index, column in enumerate(columns, start=1):
            cell = ws.cell(
                row=row_index,
                column=column_index,
                value=column.value(record),
            )

            cell.border = BORDER_THIN

            if mark == "alert":
                cell.fill = FILL_ALERT
                cell.font = FONT_ALERT
            elif mark == "warn":
                cell.fill = FILL_WARN
                cell.font = FONT_WARN
            else:
                cell.fill = (
                    FILL_STRIPE if offset % 2 else FILL_NONE
                )
                cell.font = FONT_NORMAL

            if column.fmt:
                cell.number_format = column.fmt
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT

    if not records:
        row_index = header_row + 1
        cell = ws.cell(
            row=row_index,
            column=1,
            value="Нет данных по выбранным фильтрам",
        )
        cell.font = FONT_SUBTITLE
        ws.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=len(columns),
        )

    # ------------------------------------------------------------------
    # Итоговая строка
    # ------------------------------------------------------------------
    has_totals = any(column.total for column in columns)

    if has_totals and records:
        total_row = row_index + 1

        for column_index, column in enumerate(columns, start=1):
            cell = ws.cell(row=total_row, column=column_index)
            cell.border = BORDER_THIN
            cell.fill = FILL_TOTAL
            cell.font = FONT_BOLD

            if column.field == total_label_field or (
                total_label_field is None and column_index == 1
            ):
                cell.value = TOTAL_ROW_LABEL
                cell.alignment = ALIGN_LEFT
                continue

            if not column.total:
                cell.alignment = ALIGN_LEFT
                continue

            values = [
                v
                for v in (
                    _clean(record.get(column.field))
                    for record in records
                )
                if isinstance(v, (int, float))
            ]

            if not values:
                cell.alignment = ALIGN_RIGHT
                continue

            if column.total == "wavg":
                # Среднее, взвешенное по количеству заказов.
                # Простое среднее по строкам занижает или
                # завышает итог, когда корзины разного размера.
                pairs = [
                    (
                        _clean(record.get(column.field)),
                        _clean(record.get(column.weight_field)),
                    )
                    for record in records
                ]

                pairs = [
                    (v, w)
                    for v, w in pairs
                    if isinstance(v, (int, float))
                    and isinstance(w, (int, float))
                    and w
                ]

                weight_total = sum(w for _, w in pairs)

                if weight_total:
                    cell.value = round(
                        sum(v * w for v, w in pairs) / weight_total,
                        2,
                    )
                else:
                    cell.alignment = ALIGN_RIGHT
                    continue

            elif column.total == "avg":
                cell.value = round(sum(values) / len(values), 2)

            else:
                cell.value = round(sum(values), 2)

            cell.number_format = column.fmt or FMT_INT
            cell.alignment = ALIGN_RIGHT

        row_index = total_row

    if autofilter and records:
        ws.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(len(columns))}{header_row + len(records)}"
        )

    if freeze:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    return row_index + 2


# ============================================================
# ИТОГ НАД ТАБЛИЦЕЙ
# ============================================================

def _write_totals_banner(ws, row, parts, width):
    """
    Пишет строку итогов над таблицей.

    Нужна, чтобы главные цифры листа были видны сразу, без
    прокрутки до конца таблицы.
    """
    text = "ИТОГО:   " + "     ·     ".join(parts)

    cell = ws.cell(row=row, column=1, value=text)
    cell.font = FONT_BOLD
    cell.fill = FILL_TOTAL
    cell.border = BORDER_THIN
    cell.alignment = ALIGN_LEFT

    for column_index in range(2, max(width, 1) + 1):
        filler = ws.cell(row=row, column=column_index)
        filler.fill = FILL_TOTAL
        filler.border = BORDER_THIN

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=max(width, 1),
    )

    ws.row_dimensions[row].height = 20

    return row + 2


def _totals_parts(df, specs):
    """
    Готовит подписи для плашки итогов.

    specs — список кортежей (поле, подпись, формат),
    где формат — 'int' или 'money'.
    """
    parts = []

    if df is None or df.empty:
        return ["нет данных по выбранным фильтрам"]

    for field, label, kind in specs:
        if field not in df.columns:
            continue

        values = pd.to_numeric(df[field], errors="coerce")

        if values.notna().sum() == 0:
            continue

        total = float(values.sum())

        if kind == "money":
            text = f"{total:,.0f} ₽".replace(",", " ")
        else:
            text = f"{total:,.0f}".replace(",", " ")

        parts.append(f"{label} {text}")

    return parts or ["нет числовых данных"]


# ============================================================
# ГРАФИКИ
# ============================================================

def _style_axes(chart):
    """Приглушённая сетка и оси, как на дашборде."""
    grid = ChartLines()
    grid.spPr = GraphicalProperties(
        ln=LineProperties(solidFill=CHART_GRID, w=9525)
    )

    chart.y_axis.majorGridlines = grid
    chart.x_axis.majorGridlines = None

    for axis in (chart.x_axis, chart.y_axis):
        axis.majorTickMark = "none"
        axis.minorTickMark = "none"
        axis.spPr = GraphicalProperties(
            ln=LineProperties(solidFill=CHART_GRID, w=9525)
        )

    chart.x_axis.delete = False
    chart.y_axis.delete = False


def _add_bucket_chart(ws, anchor_cell, header_row, row_count, title):
    """Столбики: количество заказов по корзинам времени."""
    if row_count <= 0:
        return

    chart = BarChart()
    chart.type = "bar"
    chart.style = None
    chart.title = title
    chart.height = 8.5
    chart.width = 20
    chart.gapWidth = 60
    chart.legend = None

    data = Reference(
        ws,
        min_col=2,
        min_row=header_row,
        max_row=header_row + row_count,
    )
    categories = Reference(
        ws,
        min_col=1,
        min_row=header_row + 1,
        max_row=header_row + row_count,
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    series = chart.series[0]
    series.graphicalProperties = GraphicalProperties(solidFill=CHART_BLUE)
    series.graphicalProperties.line.noFill = True

    _style_axes(chart)

    ws.add_chart(chart, anchor_cell)


def _add_ranked_chart(
    ws,
    anchor_cell,
    header_row,
    row_count,
    value_col,
    category_col,
    title,
):
    """Столбики по произвольной паре «категория — значение»."""
    if row_count <= 0:
        return

    chart = BarChart()
    chart.type = "bar"
    chart.style = None
    chart.title = title
    chart.height = 9.5
    chart.width = 20
    chart.gapWidth = 60
    chart.legend = None

    data = Reference(
        ws,
        min_col=value_col,
        min_row=header_row,
        max_row=header_row + row_count,
    )
    categories = Reference(
        ws,
        min_col=category_col,
        min_row=header_row + 1,
        max_row=header_row + row_count,
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    series = chart.series[0]
    series.graphicalProperties = GraphicalProperties(solidFill=CHART_BLUE)
    series.graphicalProperties.line.noFill = True

    _style_axes(chart)

    ws.add_chart(chart, anchor_cell)


def _add_daily_chart(ws, anchor_cell, header_row, row_count, title):
    """Линии: заказы и отмены по дням."""
    if row_count <= 0:
        return

    chart = LineChart()
    chart.title = title
    chart.style = None
    chart.height = 8.5
    chart.width = 24

    for column_index, color in (
        (2, CHART_BLUE),
        (4, CHART_ORANGE),
    ):
        reference = Reference(
            ws,
            min_col=column_index,
            min_row=header_row,
            max_row=header_row + row_count,
        )
        chart.add_data(reference, titles_from_data=True)

    categories = Reference(
        ws,
        min_col=1,
        min_row=header_row + 1,
        max_row=header_row + row_count,
    )
    chart.set_categories(categories)

    for series, color in zip(
        chart.series,
        (CHART_BLUE, CHART_ORANGE),
    ):
        series.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=color, w=19050)
        )
        series.smooth = False
        series.marker = None

    _style_axes(chart)

    ws.add_chart(chart, anchor_cell)


# ============================================================
# ЛИСТ «СВОДКА»
# ============================================================

def _write_summary(ws, payload, meta):
    kpi = payload["kpi"]
    as_of = payload["as_of"]

    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60

    row = _write_sheet_header(
        ws,
        "Анализ заказов FBS",
        (
            "Данные на "
            f"{as_of:%d.%m.%Y %H:%M}"
            f" · источник: {payload.get('source', '—')}"
        ),
        width=3,
    )

    # ------------------------------------------------------------------
    # Параметры отчёта
    # ------------------------------------------------------------------
    parameters = [
        ("Период заказов", meta.get("period", "не выбран — вся история")),
        ("Бренд", meta.get("brand", "все")),
        ("Категория", meta.get("category", "все")),
        ("Пол", meta.get("gender", "любой")),
        ("Норматив сборки, часов", SLA_LIMIT_HOURS),
        ("Порог внимания, часов", SLA_WARNING_HOURS),
        ("Отчёт сформирован", datetime.now()),
    ]

    cell = ws.cell(row=row, column=1, value="ПАРАМЕТРЫ ОТЧЁТА")
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.border = BORDER_THIN
    cell.alignment = ALIGN_LEFT

    for column_index in (2, 3):
        c = ws.cell(row=row, column=column_index)
        c.fill = FILL_HEADER
        c.border = BORDER_THIN

    row += 1

    for index, (label, value) in enumerate(parameters):
        label_cell = ws.cell(row=row, column=1, value=label)
        value_cell = ws.cell(row=row, column=2, value=_clean(value))

        for cell in (label_cell, value_cell):
            cell.border = BORDER_THIN
            cell.fill = FILL_STRIPE if index % 2 else FILL_NONE

        label_cell.font = FONT_NORMAL
        label_cell.alignment = ALIGN_LEFT

        value_cell.font = FONT_BOLD
        value_cell.alignment = ALIGN_LEFT

        if isinstance(value, datetime):
            value_cell.number_format = FMT_DATETIME

        row += 1

    row += 1

    # ------------------------------------------------------------------
    # Показатели
    # ------------------------------------------------------------------
    metrics = [
        (
            "Заказов всего",
            kpi.get("orders_total"),
            FMT_INT,
            "Все заказы текущего среза с учётом фильтров",
        ),
        (
            "На сборке сейчас",
            kpi.get("orders_in_work"),
            FMT_INT,
            "Продавец подтвердил, WB ждёт сборку",
        ),
        (
            f"Из них просрочено (>{SLA_LIMIT_HOURS} ч)",
            kpi.get("orders_overdue"),
            FMT_INT,
            "Требуют реакции в первую очередь",
        ),
        (
            "Закрыто поставкой",
            kpi.get("orders_closed"),
            FMT_INT,
            "Заказ вошёл в закрытую поставку",
        ),
        (
            "Отменено",
            kpi.get("orders_cancelled"),
            FMT_INT,
            "Отмены покупателем и отказы",
        ),
        (
            "Поставок задействовано",
            kpi.get("supplies_total"),
            FMT_INT,
            "Уникальные supply_id в срезе",
        ),
        (
            "Артикулов в заказах",
            kpi.get("nm_count"),
            FMT_INT,
            "Уникальные nm_id",
        ),
        (
            "Сумма заказов, ₽",
            kpi.get("amount"),
            FMT_MONEY,
            "По финальной цене заказа (final_price_rub)",
        ),
        (
            "Средний возраст на сборке, ч",
            kpi.get("avg_age_in_work"),
            FMT_HOURS,
            "От создания заказа до момента выгрузки",
        ),
        (
            "Максимальный возраст на сборке, ч",
            kpi.get("max_age_in_work"),
            FMT_HOURS,
            "Самый долгий несобранный заказ",
        ),
        (
            "Средний срок до закрытия поставки, ч",
            kpi.get("avg_hours_to_close"),
            FMT_HOURS,
            "По закрытым заказам",
        ),
    ]

    cell = ws.cell(row=row, column=1, value="ПОКАЗАТЕЛЬ")
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.border = BORDER_THIN
    cell.alignment = ALIGN_LEFT

    cell = ws.cell(row=row, column=2, value="ЗНАЧЕНИЕ")
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.border = BORDER_THIN
    cell.alignment = ALIGN_HEADER

    cell = ws.cell(row=row, column=3, value="КАК СЧИТАЕТСЯ")
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.border = BORDER_THIN
    cell.alignment = ALIGN_LEFT

    row += 1

    for index, (label, value, fmt, note) in enumerate(metrics):
        is_alert = (
            label.startswith("Из них просрочено")
            and isinstance(_clean(value), (int, float))
            and _clean(value) > 0
        )

        label_cell = ws.cell(row=row, column=1, value=label)
        value_cell = ws.cell(row=row, column=2, value=_clean(value))
        note_cell = ws.cell(row=row, column=3, value=note)

        for cell in (label_cell, value_cell, note_cell):
            cell.border = BORDER_THIN

            if is_alert:
                cell.fill = FILL_ALERT
            else:
                cell.fill = FILL_STRIPE if index % 2 else FILL_NONE

        label_cell.font = FONT_ALERT if is_alert else FONT_NORMAL
        label_cell.alignment = ALIGN_LEFT

        value_cell.font = FONT_ALERT if is_alert else FONT_BOLD
        value_cell.number_format = fmt
        value_cell.alignment = ALIGN_RIGHT

        note_cell.font = FONT_SUBTITLE
        note_cell.alignment = ALIGN_LEFT

        row += 1

    ws.freeze_panes = "A5"


# ============================================================
# ОПИСАНИЕ ЛИСТОВ
# ============================================================

def _bucket_columns(with_state=False):
    columns = [
        Column("time_group", "Время в работе", width=18),
        Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    ]

    if with_state:
        columns += [
            Column("closed_orders", "Закрыто", FMT_INT, 12, total="sum"),
            Column("open_orders", "Открыто", FMT_INT, 12, total="sum"),
        ]

    columns += [
        Column("share_pct", "Доля", FMT_PCT, 12, total="sum"),
        Column(
            "avg_hours",
            "Средние часы",
            FMT_HOURS,
            15,
            total="wavg",
            weight_field="orders",
        ),
        Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
    ]

    return columns


OVERDUE_COLUMNS = [
    Column("order_id", "Заказ", width=16),
    Column("created_at", "Создан", FMT_DATETIME, 18),
    Column("age_hours", "Часов в работе", FMT_HOURS, 16),
    Column("article", "Артикул", width=20),
    Column("title", "Наименование", width=38),
    Column("brand", "Бренд", width=18),
    Column("subject_name", "Категория", width=22),
    Column("office_name", "Пункт выдачи", width=22),
    Column("warehouse_id", "Склад", width=14),
    Column("supply_name", "Поставка", width=26),
    Column("amount", "Сумма, ₽", FMT_MONEY, 16, total="sum"),
]

LOGISTICS_COLUMNS = [
    Column("warehouse", "Склад отправления", width=20),
    Column("office_name", "Пункт выдачи", width=24),
    Column("destination_office", "Склад WB назначения", width=26),
    Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    Column("orders_in_work", "На сборке", FMT_INT, 13, total="sum"),
    Column("orders_overdue", "Просрочено", FMT_INT, 13, total="sum"),
    Column("supplies", "Поставок", FMT_INT, 12, total="sum"),
    Column(
        "avg_hours",
        "Средние часы",
        FMT_HOURS,
        15,
        total="wavg",
        weight_field="orders",
    ),
    Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
]

SUPPLIES_COLUMNS = [
    Column("supply_id", "Поставка", width=20),
    Column("supply_name", "Название", width=34),
    Column("created_at", "Создана", FMT_DATETIME, 18),
    Column("closed_at", "Закрыта", FMT_DATETIME, 18),
    Column("scan_dt", "Отсканирована", FMT_DATETIME, 18),
    Column(
        "cargo_type",
        "Тип груза",
        width=20,
        convert=_cargo_type_name,
    ),
    Column("is_b2b", "B2B", width=10),
    Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    Column(
        "cancelled_orders",
        "Отменено",
        FMT_INT,
        12,
        total="sum",
    ),
    Column(
        "avg_hours_to_close",
        "Часов до закрытия",
        FMT_HOURS,
        18,
        total="wavg",
        weight_field="orders",
    ),
    Column(
        "avg_hours_to_scan",
        "Часов до скана",
        FMT_HOURS,
        17,
        total="wavg",
        weight_field="orders",
    ),
    Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
]

PRODUCTS_COLUMNS = [
    Column("nm_id", "NM ID", width=14),
    Column("article", "Артикул", width=20),
    Column("title", "Наименование", width=38),
    Column("brand", "Бренд", width=18),
    Column("subject_name", "Категория", width=22),
    Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    Column("orders_in_work", "На сборке", FMT_INT, 13, total="sum"),
    Column("cancelled_orders", "Отменено", FMT_INT, 12, total="sum"),
    Column("zero_orders", "Нулевых", FMT_INT, 12, total="sum"),
    Column(
        "avg_hours",
        "Средние часы",
        FMT_HOURS,
        15,
        total="wavg",
        weight_field="orders",
    ),
    Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
]

STATUS_COLUMNS = [
    Column(
        "supplier_status",
        "Статус продавца",
        width=24,
        convert=_supplier_status_name,
    ),
    Column(
        "wb_status",
        "Статус WB",
        width=26,
        convert=_wb_status_name,
    ),
    Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    Column(
        "avg_hours",
        "Средние часы",
        FMT_HOURS,
        15,
        total="wavg",
        weight_field="orders",
    ),
    Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
]

DAILY_COLUMNS = [
    Column("order_date", "Дата заказа", FMT_DATE, 14),
    Column("orders", "Заказов", FMT_INT, 12, total="sum"),
    Column("closed_orders", "Закрыто", FMT_INT, 12, total="sum"),
    Column("cancelled_orders", "Отменено", FMT_INT, 12, total="sum"),
    Column(
        "avg_hours",
        "Средние часы",
        FMT_HOURS,
        15,
        total="wavg",
        weight_field="orders",
    ),
    Column("amount", "Сумма, ₽", FMT_MONEY, 18, total="sum"),
]

RAW_COLUMNS = [
    Column("order_id", "Заказ", width=16),
    Column("created_at", "Создан", FMT_DATETIME, 18),
    Column(
        "supplier_status",
        "Статус продавца",
        width=20,
        convert=_supplier_status_name,
    ),
    Column(
        "wb_status",
        "Статус WB",
        width=22,
        convert=_wb_status_name,
    ),
    Column("order_state", "Состояние", width=14),
    Column("age_hours", "Часов в работе", FMT_HOURS, 16),
    Column("time_group", "Корзина времени", width=18),
    Column("nm_id", "NM ID", width=14),
    Column("chrt_id", "CHRT ID", width=14),
    Column("article", "Артикул", width=20),
    Column("title", "Наименование", width=38),
    Column("brand", "Бренд", width=18),
    Column("subject_name", "Категория", width=22),
    Column("gender", "Пол", width=12),
    Column("warehouse_id", "Склад", width=14),
    Column("office_name", "Пункт выдачи", width=22),
    Column("destination_office", "Склад WB назначения", width=26),
    Column("supply_id", "Поставка", width=20),
    Column("supply_name", "Название поставки", width=30),
    Column("supply_created_at", "Поставка создана", FMT_DATETIME, 18),
    Column("supply_closed_at", "Поставка закрыта", FMT_DATETIME, 18),
    Column("supply_scan_dt", "Скан поставки", FMT_DATETIME, 18),
    Column("hours_to_supply_close", "Часов до закрытия", FMT_HOURS, 18),
    Column("hours_to_supply_scan", "Часов до скана", FMT_HOURS, 17),
    Column("is_zero_order", "Нулевой заказ", width=14),
    Column("price_rub", "Цена, ₽", FMT_MONEY, 14),
    Column("sale_price_rub", "Цена со скидкой, ₽", FMT_MONEY, 18),
    Column("final_price_rub", "Итоговая цена, ₽", FMT_MONEY, 18),
    Column(
        "order_amount",
        "Сумма заказа, ₽",
        FMT_MONEY,
        18,
        total="sum",
    ),
]


def _overdue_highlight(record):
    hours = _clean(record.get("age_hours"))

    if not isinstance(hours, (int, float)):
        return None

    if hours >= SLA_LIMIT_HOURS:
        return "alert"

    if hours >= SLA_WARNING_HOURS:
        return "warn"

    return None


def _logistics_highlight(record):
    overdue = _clean(record.get("orders_overdue"))

    if isinstance(overdue, (int, float)) and overdue > 0:
        return "warn"

    return None


# ============================================================
# СБОРКА КНИГИ
# ============================================================

def make_fbs_excel(payload: dict, meta: dict, raw_df=None) -> bytes:
    """
    Собирает книгу отчёта и возвращает её содержимое.

    payload — результат collect_fbs_analysis;
    meta    — человекочитаемые параметры отчёта для шапки;
    raw_df  — построчная выгрузка заказов (может отсутствовать).
    """
    as_of = payload["as_of"]
    stamp = f"Данные на {as_of:%d.%m.%Y %H:%M}"
    period = meta.get("period", "вся история")

    workbook = Workbook()

    # ---------------------------------------------------------
    # Сводка
    # ---------------------------------------------------------
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    _write_summary(summary_sheet, payload, meta)

    # ---------------------------------------------------------
    # Контроль сборки
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Контроль сборки")

    row = _write_sheet_header(
        sheet,
        "Контроль сборки заказов",
        (
            f"{stamp} · период: {period} · "
            f"норматив {SLA_LIMIT_HOURS} часов"
        ),
        width=len(_bucket_columns()),
    )

    buckets_in_work = payload["buckets_in_work"]

    row = _write_totals_banner(
        sheet,
        row,
        _totals_parts(
            buckets_in_work,
            [
                ("orders", "заказов на сборке", "int"),
                ("amount", "на сумму", "money"),
            ],
        ),
        width=len(_bucket_columns()),
    )

    cell = sheet.cell(
        row=row,
        column=1,
        value="СЕЙЧАС НА СБОРКЕ",
    )
    cell.font = FONT_TITLE
    row += 1

    bucket_header_row = row
    bucket_rows = 0 if buckets_in_work is None else len(buckets_in_work)

    row = _write_table(
        sheet,
        buckets_in_work,
        _bucket_columns(),
        row,
        total_label_field="time_group",
        freeze=False,
    )

    _add_bucket_chart(
        sheet,
        f"H{bucket_header_row}",
        bucket_header_row,
        bucket_rows,
        "Заказы по времени в работе",
    )

    cell = sheet.cell(
        row=row,
        column=1,
        value="ВСЕ АКТИВНЫЕ ЗАКАЗЫ (включая закрытые поставкой)",
    )
    cell.font = FONT_TITLE
    row += 1

    _write_table(
        sheet,
        payload["buckets_all"],
        _bucket_columns(with_state=True),
        row,
        total_label_field="time_group",
        autofilter=False,
        freeze=False,
    )

    # ---------------------------------------------------------
    # Просроченные заказы
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Просрочено")

    row = _write_sheet_header(
        sheet,
        "Заказы сверх норматива сборки",
        (
            f"{stamp} · заказы на сборке старше "
            f"{SLA_LIMIT_HOURS} часов, от самого старого"
        ),
        width=len(OVERDUE_COLUMNS),
    )

    overdue = payload["overdue"]
    overdue_count = 0 if overdue is None else len(overdue)

    row = _write_totals_banner(
        sheet,
        row,
        [f"заказов {overdue_count:,}".replace(",", " ")]
        + _totals_parts(
            overdue,
            [("amount", "на сумму", "money")],
        ),
        width=len(OVERDUE_COLUMNS),
    )

    _write_table(
        sheet,
        payload["overdue"],
        OVERDUE_COLUMNS,
        row,
        total_label_field="order_id",
        highlight=_overdue_highlight,
    )

    # ---------------------------------------------------------
    # Логистика
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Логистика")

    row = _write_sheet_header(
        sheet,
        "Откуда и куда едут заказы",
        (
            f"{stamp} · склад отправления и пункт выдачи WB, "
            "строки с просрочкой подсвечены"
        ),
        width=len(LOGISTICS_COLUMNS),
    )

    row = _write_totals_banner(
        sheet,
        row,
        _totals_parts(
            payload["logistics"],
            [
                ("orders", "заказов", "int"),
                ("orders_overdue", "просрочено", "int"),
                ("supplies", "поставок", "int"),
                ("amount", "на сумму", "money"),
            ],
        ),
        width=len(LOGISTICS_COLUMNS),
    )

    logistics_header_row = row
    logistics_rows = (
        0 if payload["logistics"] is None else len(payload["logistics"])
    )

    _write_table(
        sheet,
        payload["logistics"],
        LOGISTICS_COLUMNS,
        row,
        total_label_field="warehouse",
        highlight=_logistics_highlight,
    )

    # Столбики строятся по колонке «Заказов» (D) и подписываются
    # пунктом выдачи (B): склад один и тот же, различает направление
    # именно пункт.
    _add_ranked_chart(
        sheet,
        f"K{logistics_header_row}",
        logistics_header_row,
        min(logistics_rows, 15),
        value_col=4,
        category_col=2,
        title="Заказы по направлениям",
    )

    # ---------------------------------------------------------
    # Поставки
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Поставки")

    row = _write_sheet_header(
        sheet,
        "Поставки: сборка и отгрузка",
        f"{stamp} · от создания до скана на складе WB",
        width=len(SUPPLIES_COLUMNS),
    )

    row = _write_totals_banner(
        sheet,
        row,
        _totals_parts(
            payload["supplies"],
            [
                ("orders", "заказов", "int"),
                ("cancelled_orders", "отменено", "int"),
                ("amount", "на сумму", "money"),
            ],
        ),
        width=len(SUPPLIES_COLUMNS),
    )

    _write_table(
        sheet,
        payload["supplies"],
        SUPPLIES_COLUMNS,
        row,
        total_label_field="supply_id",
    )

    # ---------------------------------------------------------
    # Товары
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Товары")

    row = _write_sheet_header(
        sheet,
        "Что заказывают",
        f"{stamp} · период: {period}",
        width=len(PRODUCTS_COLUMNS),
    )

    row = _write_totals_banner(
        sheet,
        row,
        _totals_parts(
            payload["products"],
            [
                ("orders", "заказов", "int"),
                ("cancelled_orders", "отменено", "int"),
                ("amount", "на сумму", "money"),
            ],
        ),
        width=len(PRODUCTS_COLUMNS),
    )

    _write_table(
        sheet,
        payload["products"],
        PRODUCTS_COLUMNS,
        row,
        total_label_field="nm_id",
    )

    # ---------------------------------------------------------
    # Статусы
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("Статусы")

    row = _write_sheet_header(
        sheet,
        "Заказы по статусам",
        f"{stamp} · пары «статус продавца — статус WB»",
        width=len(STATUS_COLUMNS),
    )

    _write_table(
        sheet,
        payload["statuses"],
        STATUS_COLUMNS,
        row,
        total_label_field="supplier_status",
    )

    # ---------------------------------------------------------
    # По дням
    # ---------------------------------------------------------
    sheet = workbook.create_sheet("По дням")

    row = _write_sheet_header(
        sheet,
        "Заказы по дням",
        f"{stamp} · по дате создания заказа",
        width=len(DAILY_COLUMNS),
    )

    daily = payload["daily"]

    row = _write_totals_banner(
        sheet,
        row,
        _totals_parts(
            daily,
            [
                ("orders", "заказов", "int"),
                ("closed_orders", "закрыто", "int"),
                ("cancelled_orders", "отменено", "int"),
                ("amount", "на сумму", "money"),
            ],
        ),
        width=len(DAILY_COLUMNS),
    )

    daily_header_row = row
    daily_rows = 0 if daily is None else len(daily)

    _write_table(
        sheet,
        daily,
        DAILY_COLUMNS,
        row,
        total_label_field="order_date",
    )

    _add_daily_chart(
        sheet,
        f"H{daily_header_row}",
        daily_header_row,
        daily_rows,
        "Заказы и отмены по дням",
    )

    # ---------------------------------------------------------
    # Построчно
    # ---------------------------------------------------------
    if raw_df is not None:
        sheet = workbook.create_sheet("Заказы")

        row = _write_sheet_header(
            sheet,
            "Заказы построчно",
            f"{stamp} · исходные строки для своих сводных",
            width=len(RAW_COLUMNS),
        )

        _write_table(
            sheet,
            raw_df,
            RAW_COLUMNS,
            row,
            total_label_field="order_id",
            highlight=_overdue_highlight,
        )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return buffer.read()
