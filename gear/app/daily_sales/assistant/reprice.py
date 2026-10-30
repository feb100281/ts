# gear/app/daily_sales/assistant/reprice.py
"""Разбор переоценок за выбранные дни: что подешевело, как быстро распродалось,
сколько денег недополучили по сравнению с прежней ценой.

Запуск на сервере:
    python manage.py shell -c "from datetime import date;
    from gear.app.daily_sales.assistant import reprice;
    print(reprice.run(date(2026, 10, 22), date(2026, 10, 23)))"

Цена — наша цена до СПП с НДС (продажи до СПП / шт) из отчёта WB; цена
покупателя — после СПП. «До» — 14 дней перед переоценкой, «после» — с первого
дня переоценки по последний загруженный день. Маржа — по методике дашборда
продаж (без корректировки НДС).
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pandas as pd

BEFORE_DAYS = 14
MIN_CUT = -10.0           # снижение цены, %, начиная с которого артикул попадает в разбор
FAST_DAYS = 7             # «распродали за неделю и быстрее» — отдельная пометка


def _stock_on(d, tag: str = "0") -> pd.DataFrame:
    """Остаток на дату: склады WB + FBS (stock), в пути к покупателю (to),
    в пути от покупателя (frm)."""
    from .drop import _query
    cols = ["nm", f"stock{tag}", f"to{tag}", f"frm{tag}"]
    if d is None:
        return pd.DataFrame(columns=cols)
    df = _query("""
        WITH x AS (
            SELECT nm_id, SUM(COALESCE(quantity, 0)) AS q,
                   SUM(COALESCE(in_way_to_client, 0)) AS t,
                   SUM(COALESCE(in_way_from_client, 0)) AS f
            FROM stocks.unpacked_stocks WHERE date_from::DATE = $d GROUP BY 1
            UNION ALL
            SELECT nm_id, SUM(COALESCE(quantity, 0)), 0, 0 FROM stocks.unpacked_fbs_stocks
            WHERE date_from::DATE = $d GROUP BY 1)
        SELECT nm_id AS nm, SUM(q) AS s, SUM(t) AS t, SUM(f) AS f FROM x GROUP BY 1""",
               {"d": pd.Timestamp(d).date()})
    if df.empty:
        return pd.DataFrame(columns=cols)
    df.columns = cols
    df["nm"] = df["nm"].astype("int64")
    return df


def _arrivals() -> pd.DataFrame:
    """Первое появление товара: первый день с остатком на складах и первая/последняя
    дата УПД поставщика. Два запроса: без УПД (нет связи с Postgres) остаётся склад."""
    from .drop import _query
    st = _query("""
        SELECT nm_id AS nm, MIN(date_from::DATE) AS first_stock FROM (
            SELECT nm_id, date_from FROM stocks.unpacked_stocks WHERE COALESCE(quantity, 0) > 0
            UNION ALL
            SELECT nm_id, date_from FROM stocks.unpacked_fbs_stocks WHERE COALESCE(quantity, 0) > 0
        ) GROUP BY 1""", {})
    upd = _query("""
        SELECT l.nm_id AS nm, MIN(d.date) AS first_upd, MAX(d.date) AS last_upd
        FROM pg.public.upd_income_lines l
        JOIN pg.public.cards_upddocument d ON d.id = l.upd_document_id
        WHERE l.nm_id IS NOT NULL GROUP BY 1""", {})
    out = st if not st.empty else pd.DataFrame(columns=["nm", "first_stock"])
    if not upd.empty:
        out = out.merge(upd, on="nm", how="outer")
    for c in ("first_stock", "first_upd", "last_upd"):
        if c not in out:
            out[c] = pd.NaT
        out[c] = pd.to_datetime(out[c], errors="coerce")
    out["first_in"] = out[["first_stock", "first_upd"]].min(axis=1)
    out = out.dropna(subset=["nm"])
    out["nm"] = out["nm"].astype("int64")
    return out[["nm", "first_in", "last_upd"]]


def _margin(a: date, b: date, tag: str) -> pd.DataFrame:
    from .menu import margin
    m = margin(a, b, "nm")
    if m.empty:
        return pd.DataFrame(columns=["nm"])
    m = m.rename(columns={"nm_id": "nm"})
    m["nm"] = pd.to_numeric(m["nm"], errors="coerce")
    m = m.dropna(subset=["nm"])
    m["nm"] = m["nm"].astype("int64")
    g = m.groupby("nm", as_index=False)[["qty", "rev", "cogs", "res"]].sum()
    g[f"m_{tag}"] = g["res"] / g["rev"].where(g["rev"] != 0) * 100
    g[f"cogs_u_{tag}"] = g["cogs"] / g["qty"].where(g["qty"] != 0)
    return g[["nm", f"m_{tag}", f"cogs_u_{tag}"]]


def analyze(d1: date, d2: date, last: date | None = None,
            min_cut: float | None = None) -> tuple[pd.DataFrame, dict]:
    from .drop import _sales, _stock_now
    from .digest import loaded_date
    last = last or loaded_date() or (date.today() - timedelta(days=1))
    a0, b0 = d1 - timedelta(days=BEFORE_DAYS), d1 - timedelta(days=1)
    n1 = (last - d1).days + 1

    s0 = _sales(a0, b0)[["nm", "q", "s", "bs", "brand", "category", "sa", "title"]]
    s1 = _sales(d1, last)[["nm", "q", "s", "bs", "brand", "category", "sa", "title"]]
    sc = _sales(d1, d2)[["nm", "q", "s"]]
    df = s0.merge(s1, on="nm", how="outer", suffixes=("0", "1"))
    df = df.merge(sc.rename(columns={"q": "qc", "s": "sc"}), on="nm", how="left")
    for c in ("brand", "category", "sa", "title"):
        df[c] = df[f"{c}1"].fillna(df[f"{c}0"])
    for c in ("q0", "s0", "bs0", "q1", "s1", "bs1", "qc", "sc"):
        df[c] = pd.to_numeric(df.get(c), errors="coerce").fillna(0.0)

    df["p0"] = df["s0"] / df["q0"].where(df["q0"] > 0)            # цена до СПП, до
    # цена в дни переоценки; если в эти дни продаж нет — за весь период «после»
    df["pc"] = (df["sc"] / df["qc"].where(df["qc"] > 0)).fillna(
        df["s1"] / df["q1"].where(df["q1"] > 0))
    df["p1"] = df["s1"] / df["q1"].where(df["q1"] > 0)
    df["bp0"] = df["bs0"] / df["q0"].where(df["q0"] > 0)
    df["bp1"] = df["bs1"] / df["q1"].where(df["q1"] > 0)
    df["cut"] = (df["pc"] / df["p0"] - 1) * 100
    df["u0"] = df["q0"] / BEFORE_DAYS
    df["u1"] = df["q1"] / n1

    df = df.merge(_stock_on(b0), on="nm", how="left").merge(_stock_now(last), on="nm", how="left")
    from .drop import _STOCK_DATE
    now = _stock_on(_STOCK_DATE.get("d"), "_n")[["nm", "to_n", "frm_n"]]
    df = df.merge(now, on="nm", how="left")
    for c in ("to0", "frm0", "to_n", "frm_n"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    df["stock0"] = df["stock0"].fillna(0)
    df["stock"] = df["stock"].fillna(0)
    # снимок остатков бывает неполным (0 при продажах после) — тогда остаток на
    # начало не меньше, чем «продано после + остаток сейчас»
    df["stock0"] = df[["stock0"]].assign(est=df["stock"] + df["q1"]).max(axis=1)
    # оборачиваемость остатка на начало переоценки
    df["days_old"] = df["stock0"] / df["u0"].where(df["u0"] > 0)
    df["days_new"] = df["stock0"] / df["u1"].where(df["u1"] > 0)
    # недополучено: продали по новой цене то, что раньше уходило по старой
    df["lost"] = (df["p0"] - df["p1"]) * df["q1"]
    df = df.merge(_margin(a0, b0, "0"), on="nm", how="left").merge(_margin(d1, last, "1"),
                                                                    on="nm", how="left")
    df = df.merge(_arrivals(), on="nm", how="left")
    # сколько дней товар в продаже на первый день переоценки
    df["age"] = (pd.Timestamp(d1) - df["first_in"]).dt.days

    cut_lim = -abs(min_cut) if min_cut is not None else MIN_CUT
    hit = df[(df["cut"] <= cut_lim) & (df["q1"] > 0) & df["p0"].notna()].copy()
    hit["fast"] = hit["days_new"] <= FAST_DAYS
    hit = hit.sort_values("lost", ascending=False)
    from .drop import _STOCK_DATE
    meta = {"d1": d1, "d2": d2, "last": last, "a0": a0, "b0": b0, "n1": n1,
            "stock_date": _STOCK_DATE.get("d")}
    return hit, meta


COLUMNS = [
    ("brand", "Бренд"), ("sa", "Артикул продавца"), ("nm", "Артикул WB"), ("title", "Наименование"),
    ("p0", "Цена до СПП ДО, ₽"), ("pc", "Цена до СПП в дни переоценки, ₽"), ("cut", "Изменение цены, %"),
    ("bp0", "Цена покупателя ДО, ₽"), ("bp1", "Цена покупателя ПОСЛЕ, ₽"),
    ("u0", "Продажи ДО, шт/день"), ("u1", "Продажи ПОСЛЕ, шт/день"),
    ("stock0", "Остаток перед переоценкой, шт"), ("stock", "Остаток сейчас, шт"),
    ("days_old", "Хватило бы при старой цене, дней"), ("days_new", "Уходит при новой цене, дней"),
    ("q1", "Продано после переоценки, шт"), ("lost", "Недополучено к старой цене, ₽ (оценка)"),
    ("cogs_u_0", "Себестоимость, ₽/шт"), ("m_0", "Фин. результат ДО, % к выручке"),
    ("m_1", "Фин. результат ПОСЛЕ, % к выручке"),
]


def to_excel(hit: pd.DataFrame, meta: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Переоценки"
    ws["A1"] = (f"Переоценки {meta['d1']:%d.%m}–{meta['d2']:%d.%m.%Y}: снижение цены от "
                f"{abs(MIN_CUT):.0f}% · «до» — {meta['a0']:%d.%m}–{meta['b0']:%d.%m}, «после» — "
                f"{meta['d1']:%d.%m}–{meta['last']:%d.%m} ({meta['n1']} дн.)")
    ws["A1"].font = Font(bold=True, size=12)
    head = Font(bold=True, color="FFFFFF")
    for j, (_, name) in enumerate(COLUMNS, 1):
        c = ws.cell(row=3, column=j, value=name)
        c.font, c.fill = head, PatternFill("solid", fgColor="2F6656")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    warn = PatternFill("solid", fgColor="FBE3D6")
    for i, r in enumerate(hit.itertuples(index=False), 4):
        for j, (key, _) in enumerate(COLUMNS, 1):
            v = getattr(r, key)
            v = None if (isinstance(v, float) and pd.isna(v)) else v
            c = ws.cell(row=i, column=j, value=v)
            if key in ("cut", "m_0", "m_1"):
                c.number_format = '0.0;-0.0'
            elif key in ("u0", "u1", "days_old", "days_new"):
                c.number_format = '0.0'
            elif isinstance(v, (int, float)) and key not in ("nm",):
                c.number_format = '#,##0;-#,##0'
            if key == "days_new" and r.fast:
                c.fill = warn
    widths = [14, 16, 13, 40] + [14] * (len(COLUMNS) - 4)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = "E4"
    n = len(hit) + 3
    ws.cell(row=n + 2, column=1, value=(
        "Цена — наша цена до СПП с НДС по продажам WB; «недополучено» — (цена ДО − цена ПОСЛЕ) × "
        "проданные после переоценки штуки, оценка. Розовым — остаток уходит за неделю и быстрее. "
        "Изменение цены может быть и из-за акций WB."))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def summary(hit: pd.DataFrame, meta: dict, top: int = 10) -> str:
    if hit.empty:
        return "Снижений цены от %.0f%% за эти дни не найдено." % abs(MIN_CUT)
    fast = hit[hit["fast"]]
    lines = [f"Переоценки {meta['d1']:%d.%m}–{meta['d2']:%d.%m}: подешевело {len(hit)} арт., "
             f"из них {len(fast)} распродаются за ≤{FAST_DAYS} дн.",
             f"Недополучено к старой цене (оценка): {hit['lost'].sum():,.0f} ₽".replace(",", " ")]
    for r in hit.head(top).itertuples(index=False):
        dn = "—" if pd.isna(r.days_new) else f"{r.days_new:.0f}"
        do = "—" if pd.isna(r.days_old) else f"{r.days_old:.0f}"
        lines.append(f"• {r.brand} {r.sa} ({r.nm}): {r.p0:,.0f} → {r.pc:,.0f} ₽ ({r.cut:+.0f}%), "
                     f"{r.u0:.1f} → {r.u1:.1f} шт/д, остаток {r.stock0:,.0f} уходит за {dn} дн. "
                     f"вместо {do}; недополучено {r.lost:,.0f} ₽".replace(",", " "))
    return "\n".join(lines)


def run(d1: date, d2: date, out: str | None = None) -> str:
    hit, meta = analyze(d1, d2)
    path = Path(out or f"reports/reprice_{d1:%Y%m%d}_{d2:%Y%m%d}.xlsx")
    to_excel(hit, meta, path)
    return summary(hit, meta) + f"\n\nФайл: {path.resolve()}"
