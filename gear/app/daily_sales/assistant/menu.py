# gear/app/daily_sales/assistant/menu.py
"""Меню бота продаж: готовые отчёты по кнопкам, без ИИ.

Кнопки внизу экрана → выбор периода (быстрые кнопки или календарь «начало — конец»)
→ карточка с цифрами, график и Excel. Все цифры — SQL по тем же витринам, что
дашборд продаж и утренняя рассылка; продажи — «как на сайте WB», с НДС.

callback_data: «mn:<операция>:<отчёт>:…», даты — ГГГГММДД.
    mn:o:<r>                 — выбор периода
    mn:p:<r>:<пресет>        — быстрый период
    mn:c:<r>:<ГГГГММ>:<нач>  — листание календаря (нач = 0, пока не выбрано)
    mn:d:<r>:<дата>          — выбрано начало
    mn:e:<r>:<нач>:<кон>     — выбран конец → отчёт
    mn:f:<r>:<файл>:<нач>:<кон> — дополнительный Excel
"""
from __future__ import annotations

import calendar
import html
import re
from datetime import date, timedelta

import pandas as pd

# ---------------------------------------------------------------- кнопки
REPORTS = {
    "s": "📊 Продажи",
    "k": "📦 Остатки",
    "b": "🏷 Бренды",
    "c": "🗂 Категории",
    "w": "💸 Расходы WB",
    "a": "📣 Реклама",
    "d": "💳 Скидка СПП",
    "r": "📑 Коммерческий обзор",
    "n": "📅 Недели",
    "m": "💰 Маржинальность",
    "l": "🔻 Продаём в убыток",
    "z": "🔎 Почему упали продажи",
    "g": "🚨 Сигналы",
}
SINGLE_DAY = {"r", "g"}                 # отчёт на дату, а не за период
BTN_AI = "🤖 Спросить ИИ"
BTN_HELP = "ℹ️ Как пользоваться"
BY_TEXT = {v: k for k, v in REPORTS.items()}

PRESETS = (("y", "Вчера"), ("7", "7 дней"), ("w", "Прошлая неделя"),
           ("m", "Этот месяц"), ("pm", "Прошлый месяц"), ("30", "30 дней"))
MONTHS = ("Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
          "Сентябрь", "Октябрь", "Ноябрь", "Декабрь")
WEEK = ("Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс")
DATA_START = date(2024, 1, 29)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
CHART_MAX_DAYS = 62
TOP_N = 10
MOVERS_N = 6

QTY, RET, RP, RA, PAY = ("Итого, шт", "Возвраты, шт", "До СПП: итого, ₽",
                         "После СПП: итого, ₽", "К перечислению, ₽")
SOLD = "Продажи, шт"


APP_BTN = "📱 Приложение"


def app_url() -> str:
    """Адрес Mini App (https), из .env: TELEGRAM_WEBAPP_URL_SALES."""
    import os
    return (os.getenv("TELEGRAM_WEBAPP_URL_SALES") or "").strip()


def app_row() -> list:
    url = app_url()
    return [[{"text": APP_BTN, "web_app": {"url": url}}]] if url.startswith("https://") else []


def keyboard() -> dict:
    """Постоянное меню бота продаж."""
    r = REPORTS
    return {"keyboard": [[{"text": r["s"]}, {"text": r["k"]}],
                         [{"text": r["b"]}, {"text": r["c"]}],
                         [{"text": r["w"]}, {"text": r["a"]}],
                         [{"text": r["n"]}, {"text": r["m"]}],
                         [{"text": r["l"]}, {"text": r["d"]}],
                         [{"text": r["z"]}, {"text": r["g"]}],
                         [{"text": r["r"]}, {"text": BTN_AI}],
                         [{"text": BTN_HELP}]] + app_row(),
            "resize_keyboard": True, "is_persistent": True,
            "input_field_placeholder": "Кнопка или вопрос своими словами…"}


# ---------------------------------------------------------------- формат
def esc(t) -> str:
    return html.escape(str(t if t is not None else ""), quote=False)


def num(v, d=0) -> str:
    s = f"{float(v or 0):,.{d}f}".replace(",", " ").replace(".", ",")
    return s.replace("-", "−")


def rub(v) -> str:
    v = float(v or 0)
    if abs(v) >= 1e6:
        return num(v / 1e6, 2) + " млн ₽"
    return num(v) + " ₽"


def chg(cur, prev, pp=False) -> str:
    """▲ 8,2% / ▼ 3,1% к прошлому периоду; pp — разница в процентных пунктах."""
    cur, prev = float(cur or 0), float(prev or 0)
    if pp:
        d = cur - prev
        if abs(d) < 0.05:
            return "≈"
        return ("▲ " if d > 0 else "▼ ") + num(abs(d), 1) + " п.п."
    if not prev:
        return ""
    d = (cur - prev) / abs(prev) * 100
    if abs(d) < 0.05:
        return "≈ 0%"
    return ("▲ " if d > 0 else "▼ ") + num(abs(d), 1) + "%"


def span(a: date, b: date) -> str:
    if a == b:
        return f"{a:%d.%m.%Y}"
    if a.year == b.year:
        return f"{a:%d.%m}–{b:%d.%m.%Y}"
    return f"{a:%d.%m.%Y}–{b:%d.%m.%Y}"


def ndays(a: date, b: date) -> int:
    return (b - a).days + 1


def ymd(d: date) -> str:
    return f"{d:%Y%m%d}"


def from_ymd(s: str) -> date:
    return date(int(s[:4]), int(s[4:6]), int(s[6:8]))


# ---------------------------------------------------------------- периоды
def loaded() -> date:
    from .digest import loaded_date
    return loaded_date() or date.today() - timedelta(days=1)


def month_end(d: date) -> date:
    return (d.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)


def preset(code: str, hi: date) -> tuple[date, date]:
    if code.startswith("w") and code[1:].isdigit():          # N календарных недель по hi
        mon = hi - timedelta(days=hi.weekday())
        return mon - timedelta(weeks=int(code[1:]) - 1), hi
    if code == "y":
        return hi, hi
    if code.isdigit():                                       # N последних дней
        return hi - timedelta(days=int(code) - 1), hi
    if code == "30":
        return hi - timedelta(days=29), hi
    if code == "w":
        b = hi if hi.weekday() == 6 else hi - timedelta(days=hi.weekday() + 1)
        return b - timedelta(days=6), b
    if code == "m":
        return hi.replace(day=1), hi
    pb = hi.replace(day=1) - timedelta(days=1)          # прошлый месяц
    return pb.replace(day=1), pb


def prev_range(a: date, b: date) -> tuple[date, date]:
    """Прошлый период для сравнения: месяц — прошлый месяц (те же числа), иначе — столько же дней раньше."""
    if a.day == 1 and a.month == b.month and a.year == b.year:
        pa = (a - timedelta(days=1)).replace(day=1)
        if b == month_end(b):
            return pa, month_end(pa)
        return pa, min(pa.replace(day=1) + timedelta(days=b.day - 1), month_end(pa))
    n = ndays(a, b)
    return a - timedelta(days=n), a - timedelta(days=1)


WEEK_PRESETS = (("w2", "Эта и прошлая"), ("w4", "4 недели"), ("w8", "8 недель"), ("w13", "13 недель"))


DROP_PRESETS = (("7", "7 дней"), ("14", "14 дней"), ("28", "28 дней"),
                ("w", "Прошлая неделя"), ("m", "Этот месяц"), ("pm", "Прошлый месяц"))


def period_kb(rep: str) -> dict:
    if rep == "z":
        btn = [{"text": t, "callback_data": f"mn:p:z:{c}"} for c, t in DROP_PRESETS]
        return {"inline_keyboard": [btn[:3], btn[3:], [{"text": "📅 Свой период", "callback_data": "mn:c:z::0"}]]}
    if rep == "n":
        btn = [{"text": t, "callback_data": f"mn:p:n:{c}"} for c, t in WEEK_PRESETS]
        return {"inline_keyboard": [btn[:2], btn[2:],
                                    [{"text": "📅 Выбрать недели", "callback_data": "mn:c:n::0"}]]}
    if rep in SINGLE_DAY:
        return {"inline_keyboard": [[{"text": "Вчера", "callback_data": f"mn:p:{rep}:y"},
                                     {"text": "📅 Другая дата", "callback_data": f"mn:c:{rep}::0"}]]}
    rows, row = [], []
    for code, label in PRESETS:
        row.append({"text": label, "callback_data": f"mn:p:{rep}:{code}"})
        if len(row) == 3:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([{"text": "📅 Свой период", "callback_data": f"mn:c:{rep}::0"}])
    return {"inline_keyboard": rows}


def calendar_kb(rep: str, year: int, month: int, start: date | None, lo: date, hi: date) -> dict:
    """Месяц кнопками. Без start — выбираем начало, со start — конец (дни раньше начала серые)."""
    s = ymd(start) if start else "0"
    rows = [[{"text": f"{MONTHS[month - 1]} {year}", "callback_data": "mn:n"}],
            [{"text": w, "callback_data": "mn:n"} for w in WEEK]]
    for week in calendar.Calendar(firstweekday=0).monthdayscalendar(year, month):
        line = []
        for day in week:
            if not day:
                line.append({"text": " ", "callback_data": "mn:n"})
                continue
            d = date(year, month, day)
            if d < lo or d > hi or (start and d < start):
                line.append({"text": "·", "callback_data": "mn:n"})
            elif rep in SINGLE_DAY:
                line.append({"text": str(day), "callback_data": f"mn:e:{rep}:{ymd(d)}:{ymd(d)}"})
            elif start:
                mark = f"[{day}]" if d == start else str(day)
                line.append({"text": mark, "callback_data": f"mn:e:{rep}:{s}:{ymd(d)}"})
            else:
                line.append({"text": str(day), "callback_data": f"mn:d:{rep}:{ymd(d)}"})
        rows.append(line)
    first = date(year, month, 1)
    prev_m = first - timedelta(days=1)
    next_m = month_end(first) + timedelta(days=1)
    nav = [{"text": "‹", "callback_data": f"mn:c:{rep}:{prev_m:%Y%m}:{s}"} if prev_m >= lo
           else {"text": " ", "callback_data": "mn:n"},
           {"text": "Закрыть", "callback_data": f"mn:o:{rep}"},
           {"text": "›", "callback_data": f"mn:c:{rep}:{next_m:%Y%m}:{s}"} if next_m <= hi
           else {"text": " ", "callback_data": "mn:n"}]
    rows.append(nav)
    if start:
        rows.append([{"text": f"Только {start:%d.%m}", "callback_data": f"mn:e:{rep}:{s}:{s}"}])
    return {"inline_keyboard": rows}


def ask_text(rep: str) -> str:
    if rep in SINGLE_DAY:
        return (f"<b>{REPORTS[rep]}</b>\nНа какую дату? Нажмите кнопку или напишите дату: "
                "<code>15.10</code>")
    return (f"<b>{REPORTS[rep]}</b>\nЗа какой период? Нажмите кнопку или напишите даты: "
            "<code>01.10–15.10</code>, <code>01.09.2026-30.09.2026</code>, <code>05.10</code>")


_D = r"(\d{1,2})[./](\d{1,2})(?:[./](\d{2,4}))?"
_RANGE = re.compile(rf"^\s*(?:с\s*)?{_D}\s*(?:[-–—]|по|до|\s)\s*{_D}\s*$", re.I)
_ONE = re.compile(rf"^\s*{_D}\s*$")


def _mkdate(d, m, y, hi: date) -> date | None:
    year = hi.year if not y else (int(y) + 2000 if len(y) == 2 else int(y))
    try:
        out = date(year, int(m), int(d))
    except ValueError:
        return None
    if not y and out > hi:                      # «20.12» в январе — прошлый год
        out = out.replace(year=out.year - 1)
    return out


def parse_range(text: str, hi: date) -> tuple[date, date] | None:
    """«01.10–15.10», «1.10 по 15.10.26», «с 01.09.2026 до 30.09.2026», «05.10» → (начало, конец)."""
    t = (text or "").strip()
    m = _RANGE.match(t)
    if m:
        a, b = _mkdate(*m.groups()[:3], hi), _mkdate(*m.groups()[3:], hi)
        if m.group(6) is None and m.group(3) and b and a and b.year != a.year and b < a:
            b = b.replace(year=a.year)          # год указан только у начала
        if a and b:
            a, b = sorted((a, b))
            return max(a, DATA_START), min(b, hi)
        return None
    m = _ONE.match(t)
    if m:
        d = _mkdate(*m.groups(), hi)
        return (d, d) if d and DATA_START <= d <= hi else None
    return None


# ---------------------------------------------------------------- данные
def sales(a: date, b: date, level: str) -> pd.DataFrame:
    from ..wb_sales_export import fetch
    return fetch(a, b, level, [])


def _photo(fn):
    try:
        return fn()
    except Exception as e:                       # без картинки отчёт всё равно уходит
        print(f"[menu] график: {type(e).__name__}: {e}", flush=True)
        return None


HL_FILL = "E7F1ED"                                   # светло-зелёная заливка ключевых колонок


def _highlight(ws, mp, cols, n):
    """Колонки остатка — светло-зелёной заливкой (данные, без шапки)."""
    from openpyxl.styles import PatternFill
    fill = PatternFill("solid", start_color=HL_FILL, end_color=HL_FILL)
    for i, c in enumerate(cols):
        if str(c).startswith("Остаток") or str(c).startswith("Хватит остатка"):
            for r in range(8, 8 + n):
                ws.cell(row=r, column=mp.COL_CODE + i).fill = fill


def _book(title: str, subtitle: str, sheets: list[tuple[str, list, list]], notes=None) -> bytes:
    """Excel в стиле мэн пака: листы (имя, колонки, строки) + описание."""
    import io
    from openpyxl import Workbook
    from gear.management.commands import mp
    from .excel import _write_table

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        name, cols, rows = sheet[:3]
        opts = sheet[3] if len(sheet) > 3 else {}
        ws = wb.create_sheet(name[:31])
        _write_table(ws, mp, f"{title} — {name.lower()}", subtitle, None, cols, rows,
                     totals=opts.get("totals", True))
        _highlight(ws, mp, cols, len(rows))
        if "Наименование" in cols:                          # закрепляем до наименования включительно
            ws.freeze_panes = ws.cell(row=8, column=mp.COL_CODE + cols.index("Наименование") + 1)
    if notes:
        from openpyxl.styles import Font
        ws = wb.create_sheet("Описание")
        mp.sheet_setup(ws, landscape=False)
        mp.title_band(ws, "Описание", subtitle, 4, show_back=False)
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 90
        for i, (k, v) in enumerate(notes):
            mp.write(ws, 7 + i, 2, k, font=Font(name=mp.FONT, size=9, bold=True, color=mp.NAVY_3),
                     align="left", border=mp.B_BOTTOM)
            mp.write(ws, 7 + i, 3, v, font=Font(name=mp.FONT, size=9, color=mp.TEXT),
                     align="left", wrap=True, border=mp.B_BOTTOM)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table(title: str, cols: list, rows: list, kinds: list | None = None) -> dict:
    """Таблица для приложения: kinds — формат колонок (text|int|rub|pct|signed)."""
    return {"title": title, "cols": cols, "rows": rows, "kinds": kinds or ["text"] * len(cols)}


# Результат отчёта: {"text", "photo", "files": [(имя, байты, mime)], "kb"}
def _files_kb(rep, a, b, items) -> dict:
    s, e = ymd(a), ymd(b)
    btn = [{"text": "⬇️ " + label, "callback_data": f"mn:f:{rep}:{code}:{s}:{e}"} for code, label in items]
    rows = [btn[i:i + 2] for i in range(0, len(btn), 2)]
    rows.append([{"text": "🔁 Другой период", "callback_data": f"mn:o:{rep}"}])
    return {"inline_keyboard": rows}


# ---------------------------------------------------------------- 📊 продажи
SALES_FILES = (("day", "Excel: по дням"), ("brand", "Excel: по брендам"),
               ("category", "Excel: по категориям"), ("nm", "Excel: по артикулам"))
LEVEL_NAMES = {"day": "по дням", "brand": "по брендам", "category": "по категориям",
               "nm": "по артикулам", "brand_day": "бренды по дням"}


def _tot(df, col):
    return float(df[col].sum()) if col in df.columns and not df.empty else 0.0


RP_S, RP_R = "До СПП: продажи, ₽", "До СПП: возвраты, ₽"


def _why(cur, old) -> str:
    """За счёт чего изменились продажи до СПП: количество, средняя цена, возвраты.

    Методика коммерческого обзора 2.0: количество = Δшт × прошлая цена,
    цена = Δцены × текущие шт, возвраты = −Δвозвратов в ₽. Сумма = изменение итога.
    """
    q, oq = _tot(cur, SOLD), _tot(old, SOLD)
    s, os_ = _tot(cur, RP_S), _tot(old, RP_S)
    if not q or not oq:
        return ""
    p, op = s / q, os_ / oq
    eff = [("количество", (q - oq) * op), ("средняя цена", (p - op) * q),
           ("возвраты", -(_tot(cur, RP_R) - _tot(old, RP_R)))]
    total = sum(v for _, v in eff)
    if not total:
        return ""
    sign = lambda v: ("+" if v >= 0 else "−") + rub(abs(v))
    main = max(eff, key=lambda e: abs(e[1]))
    text = (f"<b>Почему {'выросли' if total > 0 else 'упали'} продажи ({sign(total)})</b>\n"
            + "\n".join(f"• {name}: {sign(v)}" for name, v in eff)
            + f"\nГлавное — <b>{main[0]}</b>. ")
    if main[0] == "возвраты":
        text += "Смотреть надо на возвраты: размеры, фото, описание — цена тут ни при чём."
    elif main[0] == "количество" and main[1] < 0:
        text += "Упал спрос: проверьте наличие лидеров, позиции в выдаче и рекламу."
    elif main[0] == "количество":
        text += "Рост за счёт количества — здоровый рост. Проверьте, хватит ли запаса."
    elif main[1] < 0:
        text += ("Средняя цена ниже: скидки или больше дешёвых позиций в продажах. "
                 "Подробно — в «💳 Скидка СПП».")
    else:
        text += "Цена выросла без потери количества — хороший знак."
    return text


def _sales_tables(a: date, b: date) -> list:
    out = []
    try:
        for level, ttl, key in (("brand", "Бренды", "Бренд"), ("category", "Категории", "Категория")):
            m, *_ = _compare(level, a, b)
            tot = m[RP].sum() or 1
            out.append(table(ttl, [key, "До СПП", "Шт", "Доля", "К прошлому"],
                             [[r[key], r[RP], r[QTY], r[RP] / tot * 100,
                               (r["d_rp"] / r[RP + "_old"] * 100) if r[RP + "_old"] else None]
                              for _, r in m.iterrows() if r[RP] or r[RP + "_old"]],
                             ["text", "rub", "int", "pct", "signed_pct"]))
    except Exception as e:
        print(f"[menu] таблицы продаж: {type(e).__name__}: {e}", flush=True)
    return out


def report_sales(a: date, b: date) -> dict | None:
    from . import digest_charts as ch
    pa, pb = prev_range(a, b)
    df = sales(pa, b, "day")
    if df.empty:
        return None
    dcol = df.columns[0]
    df[dcol] = pd.to_datetime(df[dcol]).dt.date
    cur, old = df[df[dcol] >= a], df[df[dcol] <= pb]
    if cur.empty:
        return None
    rp, ra, pay, q, r, sold = (_tot(cur, c) for c in (RP, RA, PAY, QTY, RET, SOLD))
    o_rp, o_ra, o_pay, o_q, o_r, o_sold = (_tot(old, c) for c in (RP, RA, PAY, QTY, RET, SOLD))
    ret_pct = r / sold * 100 if sold else 0
    o_ret_pct = o_r / o_sold * 100 if o_sold else 0
    spp = (1 - ra / rp) * 100 if rp else 0
    o_spp = (1 - o_ra / o_rp) * 100 if o_rp else 0
    n = ndays(a, b)
    why = _why(cur, old)
    lines = [f"<b>📊 Продажи · {span(a, b)}</b>",
             f"<i>{n} дн. · сравнение с {span(pa, pb)}</i>", "",
             f"До СПП: <b>{rub(rp)}</b>  {chg(rp, o_rp)}",
             f"После СПП: <b>{rub(ra)}</b>  {chg(ra, o_ra)}",
             f"К перечислению: <b>{rub(pay)}</b>  {chg(pay, o_pay)}",
             f"Продано: <b>{num(q)} шт</b>  {chg(q, o_q)}",
             f"Возвраты: <b>{num(r)} шт</b> · {num(ret_pct, 1)}%  {chg(ret_pct, o_ret_pct, pp=True)}",
             f"Средняя цена до СПП: <b>{rub(rp / q if q else 0)}</b>  "
             f"{chg(rp / q if q else 0, o_rp / o_q if o_q else 0)}",
             f"СПП: <b>{num(spp, 1)}%</b>  {chg(spp, o_spp, pp=True)}"]
    if why:
        lines += ["", why]
    if n > 1:
        best = cur.loc[cur[RP].idxmax()]
        lines += ["", f"Лучший день: <b>{best[dcol]:%d.%m}</b> — {rub(best[RP])}",
                  f"В среднем за день: {rub(rp / n)} и {num(q / n)} шт"]
    lines += ["", "<i>Как на сайте WB: с НДС, по дате операции в отчёте реализации.</i>"]

    photo = None
    if n > 1:
        show = df if ndays(pa, b) <= CHART_MAX_DAYS else cur
        days = list(show[dcol])
        vals = [float(v) / 1e6 for v in show[RP]]
        n_cur = len(cur) if show is df else len(show)
        photo = _photo(lambda: ch.sales_period(
            days, vals, n_cur, None, f"Продажи до СПП — {rub(rp)}",
            f"млн ₽ по дням · {span(a, b)}" + (" · бледные — прошлый период" if show is df else "")))
    show = df if ndays(pa, b) <= CHART_MAX_DAYS else cur
    chart = {"type": "bar", "title": "Продажи до СПП по дням, ₽",
             "labels": [f"{d:%d.%m}" for d in show[dcol]], "values": [round(float(v)) for v in show[RP]],
             "n_cur": len(cur) if show is df else len(show)}
    return {"text": "\n".join(lines), "photo": photo, "auto": "day", "chart": chart,
            "files": SALES_FILES, "tables": _sales_tables(a, b), "kb": _files_kb("s", a, b, SALES_FILES)}


def file_sales(level: str, a: date, b: date):
    from ..wb_sales_export import to_excel
    if level not in LEVEL_NAMES:
        return None
    df = sales(a, b, level)
    name = f"Продажи WB {LEVEL_NAMES[level]} {span(a, b)}".replace("/", ".")
    data, fname = to_excel(df, name, a, b, level, [])
    return fname, data, XLSX


# ---------------------------------------------------------------- 🏷 бренды / 🗂 категории
GROUP = {"b": ("brand", "Бренд", "брендов", "🏷 Бренды"),
         "c": ("category", "Категория", "категорий", "🗂 Категории")}


def _compare(level, a, b):
    pa, pb = prev_range(a, b)
    key = {"brand": "Бренд", "category": "Категория"}[level]
    cur, old = sales(a, b, level), sales(pa, pb, level)
    cols = [key, QTY, RET, SOLD, RP, RA]
    cur = cur[cols] if not cur.empty else pd.DataFrame(columns=cols)
    old = old[cols] if not old.empty else pd.DataFrame(columns=cols)
    m = cur.merge(old, on=key, how="outer", suffixes=("", "_old")).fillna(0)
    m["d_rp"] = m[RP] - m[RP + "_old"]
    total = m[RP].sum()
    m["share"] = m[RP] / total * 100 if total else 0
    return m.sort_values(RP, ascending=False), pa, pb, key


def report_group(rep: str, a: date, b: date) -> dict | None:
    from . import digest_charts as ch
    level, key, many, title = GROUP[rep]
    m, pa, pb, key = _compare(level, a, b)
    if m.empty or not m[RP].sum():
        return None
    lines = [f"<b>{title} · {span(a, b)}</b>",
             f"<i>сравнение с {span(pa, pb)} · продажи до СПП с НДС</i>", "",
             f"<b>Топ-{min(TOP_N, len(m))} по продажам</b>"]
    for i, r in enumerate(m.head(TOP_N).itertuples(index=False), 1):
        d = dict(zip(m.columns, r))
        lines.append(f"{i}. {esc(d[key])} — <b>{rub(d[RP])}</b> · {num(d['share'], 1)}%  "
                     f"{chg(d[RP], d[RP + '_old'])}")
    up = m[m["d_rp"] > 0].sort_values("d_rp", ascending=False).head(MOVERS_N)
    down = m[m["d_rp"] < 0].sort_values("d_rp").head(MOVERS_N)
    if not up.empty:
        lines += ["", "<b>▲ Растут сильнее всех</b>"] + [
            f"• {esc(x)} — +{rub(v)}" for x, v in zip(up[key], up["d_rp"])]
    if not down.empty:
        lines += ["", "<b>▼ Падают сильнее всех</b>"] + [
            f"• {esc(x)} — −{rub(abs(v))}" for x, v in zip(down[key], down["d_rp"])]
    lines.append(f"\nВсего {many} с продажами: {len(m[m[RP] > 0])}. Полный список — в Excel.")
    movers = [(str(x), float(v)) for x, v in
              list(zip(up[key], up["d_rp"]))[:4] + list(zip(down[key], down["d_rp"]))[:4]]
    photo = _photo(lambda: ch.brands(
        movers, f"{title[2:]}: кто вырос и упал", f"{span(a, b)} к {span(pa, pb)}, до СПП")) \
        if movers else None
    files = (("cmp", "Excel: сравнение"), ("days", "Excel: по дням"))
    chart = {"type": "hbar", "title": "Изменение к прошлому периоду, ₽",
             "labels": [x for x, _ in movers], "values": [round(v) for _, v in movers]} if movers else None
    return {"text": "\n".join(lines), "photo": photo, "auto": "cmp", "chart": chart,
            "files": files, "kb": _files_kb(rep, a, b, files)}


def file_group(rep: str, code: str, a: date, b: date):
    level, key, many, title = GROUP[rep]
    if code == "days":
        if level == "brand":
            return file_sales("brand_day", a, b)
        df = sales(a, b, "nm_day")
        if df.empty:
            return None
        df["Дата"] = pd.to_datetime(df["Дата"]).dt.date
        g = df.groupby(["Дата", "Категория"], as_index=False)[[SOLD, RET, QTY, RP, RA]].sum()
        cols = ["Дата", "Категория", SOLD, RET, QTY, RP, RA]
        rows = g[cols].values.tolist()
        data = _book("Категории", f"{span(a, b)} · как на сайте WB, с НДС",
                     [("По дням", cols, rows)])
        return f"Категории по дням {span(a, b)}.xlsx", data, XLSX
    m, pa, pb, key = _compare(level, a, b)
    cols = [key, "Шт", "Шт прошлый", "Изменение шт", "До СПП, ₽", "До СПП прошлый, ₽",
            "Изменение, ₽", "Изменение, %", "Доля, %", "После СПП, ₽", "Возвраты, %"]
    rows = []
    for _, r in m.iterrows():
        rows.append([r[key], r[QTY], r[QTY + "_old"], r[QTY] - r[QTY + "_old"], r[RP],
                     r[RP + "_old"], r["d_rp"],
                     (r["d_rp"] / r[RP + "_old"] * 100) if r[RP + "_old"] else None,
                     r["share"], r[RA], (r[RET] / r[SOLD] * 100) if r[SOLD] else None])
    sub = f"{span(a, b)} к {span(pa, pb)} · как на сайте WB, с НДС"
    data = _book(f"{title[2:]}", sub, [("Сравнение", cols, rows)], notes=[
        ("Период", span(a, b)), ("Сравнение", span(pa, pb)),
        ("До СПП", "Наша цена до скидки постоянного покупателя, с НДС, продажи минус возвраты."),
        ("После СПП", "«Вайлдберриз реализовал Товар (Пр)» — цена покупателя."),
        ("Источник", "Отчёт реализации WB, дата операции.")])
    return f"{title[2:]} {span(a, b)}.xlsx", data, XLSX


# ---------------------------------------------------------------- 📦 остатки
STOCK_SQL = """
WITH wb AS (
    SELECT nm_id, SUM(COALESCE(quantity, 0)) AS wb_qty,
           SUM(COALESCE(in_way_to_client, 0)) + SUM(COALESCE(in_way_from_client, 0)) AS way
    FROM stocks.unpacked_stocks WHERE date_from::DATE = $d::DATE GROUP BY nm_id
),
fbs AS (
    SELECT nm_id, SUM(COALESCE(quantity, 0)) AS fbs_qty
    FROM stocks.unpacked_fbs_stocks WHERE date_from::DATE = $d::DATE GROUP BY nm_id
)
SELECT COALESCE(wb.nm_id, fbs.nm_id) AS nm_id,
       COALESCE(wb.wb_qty, 0) + COALESCE(wb.way, 0) + COALESCE(fbs.fbs_qty, 0) AS qty,
       COALESCE(NULLIF(TRIM(UPPER(p.brand)), ''), 'НЕ УКАЗАН') AS brand,
       COALESCE(NULLIF(TRIM(p.subject_name), ''), 'Не указана') AS category,
       COALESCE(p.sa_name, '') AS sa, COALESCE(NULLIF(TRIM(p.title), ''), '') AS title
FROM wb FULL OUTER JOIN fbs ON wb.nm_id = fbs.nm_id
LEFT JOIN inventories.wb_product p ON p.card_id = COALESCE(wb.nm_id, fbs.nm_id)
"""


def _stock_dates(a, b):
    from conns import get_duckdb_conn_with_opt
    from .stocks import _last_date
    with get_duckdb_conn_with_opt(ro=True) as con:
        db = _last_date(con, b)
        da = _last_date(con, a - timedelta(days=1)) or _last_date(con, a)
    return da, db


def _stock_nm(d):
    from conns import get_duckdb_conn_with_opt
    with get_duckdb_conn_with_opt(ro=True) as con:
        return con.execute(STOCK_SQL, {"d": d}).df()


def _stock_table(a, b, by):
    """Остатки на начало и конец по группе + продажи за период и дни запаса."""
    da, db = _stock_dates(a, b)
    s_end = _stock_nm(db)
    s_beg = _stock_nm(da) if da else s_end.iloc[0:0]
    sold = sales(a, b, "nm")
    sold = sold.rename(columns={"Артикул WB": "nm_id"})[["nm_id", QTY]] if not sold.empty \
        else pd.DataFrame(columns=["nm_id", QTY])
    attrs = pd.concat([s_end, s_beg])[["nm_id", "brand", "category", "sa", "title"]] \
        .drop_duplicates("nm_id")
    m = attrs.merge(s_beg[["nm_id", "qty"]].rename(columns={"qty": "beg"}), on="nm_id", how="left") \
        .merge(s_end[["nm_id", "qty"]].rename(columns={"qty": "end"}), on="nm_id", how="left") \
        .merge(sold, on="nm_id", how="left").fillna({"beg": 0, "end": 0, QTY: 0})
    keys = {"brand": ["brand"], "category": ["category"],
            "nm": ["brand", "category", "sa", "nm_id", "title"]}[by]
    g = m.groupby(keys, as_index=False)[["beg", "end", QTY]].sum()
    g = g[(g["beg"] != 0) | (g["end"] != 0) | (g[QTY] != 0)]
    days = ndays(a, b)
    g["chg"] = g["end"] - g["beg"]
    g["cover"] = g.apply(lambda r: r["end"] / (r[QTY] / days) if r[QTY] > 0 else None, axis=1)
    return g.sort_values("end", ascending=False), da, db


def report_stocks(a: date, b: date) -> dict | None:
    from . import digest_charts as ch
    from .stocks import series, totals
    end = totals(b)
    if not end:
        return None
    beg = totals(a - timedelta(days=1)) or totals(a)
    places = (("На складах WB", "wb"), ("Наш склад FBS", "fbs"),
              ("Едет к покупателям", "to_client"), ("Возвращается", "from_client"))
    d = end["total"] - (beg["total"] if beg else 0)
    lines = [f"<b>📦 Остатки на {end['date']:%d.%m.%Y}</b>",
             f"<i>изменение с {beg['date']:%d.%m.%Y}</i>" if beg else "", "",
             f"Всего: <b>{num(end['total'])} шт</b>"
             + (f"  {'▲' if d > 0 else '▼'} {num(abs(d))} шт ({chg(end['total'], beg['total'])[2:]})"
                if beg and d else "")]
    for name, k in places:
        dd = end[k] - beg[k] if beg else 0
        lines.append(f"{name}: <b>{num(end[k])}</b>" + (f"  ({'+' if dd > 0 else '−'}{num(abs(dd))})"
                                                       if dd else ""))
    q = _tot(sales(a, b, "day"), QTY)
    n = ndays(a, b)
    if q > 0:
        lines += ["", f"Продано за период: <b>{num(q)} шт</b> (≈{num(q / n)} в день)",
                  f"Хватит примерно на <b>{num(end['total'] / (q / n))} дн.</b>"]
    tables = []
    try:
        g, _, _ = _stock_table(a, b, "brand")
        tot_end = g["end"].sum() or 1
        top = g.sort_values("end", ascending=False).head(10)
        lines += ["", "<b>Как остатки распределены по брендам</b>"] + [
            f"{i}. {esc(x)} — <b>{num(e)} шт</b> · {num(e / tot_end * 100, 1)}%"
            + (f" ({'+' if c > 0 else '−'}{num(abs(c))})" if c else "")
            for i, (x, e, c) in enumerate(zip(top["brand"], top["end"], top["chg"]), 1)]
        if len(g) > 10:
            lines.append(f"…ещё {len(g) - 10} — в Excel и в приложении.")
        for by, ttl in (("brand", "Бренды"), ("category", "Категории")):
            t = g if by == "brand" else _stock_table(a, b, "category")[0]
            t = t.sort_values("end", ascending=False)
            te = t["end"].sum() or 1
            tables.append(table(ttl, [ttl[:-1] if by == "brand" else "Категория", "Остаток, шт",
                                      "Изм., шт", "Доля", "Запас, дн."],
                                [[r[by], r["end"], r["chg"], r["end"] / te * 100,
                                  r["cover"] if r["cover"] == r["cover"] else None]
                                 for _, r in t.iterrows()],
                                ["text", "int", "signed", "pct", "int"]))
        up = g[g["chg"] > 0].sort_values("chg", ascending=False).head(5)
        down = g[g["chg"] < 0].sort_values("chg").head(5)
        if not up.empty:
            lines += ["", "<b>▲ Прибавилось больше всего</b>"] + [
                f"• {esc(x)} — +{num(v)} шт" for x, v in zip(up["brand"], up["chg"])]
        if not down.empty:
            lines += ["", "<b>▼ Убыло больше всего</b>"] + [
                f"• {esc(x)} — −{num(abs(v))} шт" for x, v in zip(down["brand"], down["chg"])]
        short = g[(g["cover"].notna()) & (g["cover"] < 14) & (g[QTY] > 0)].sort_values("cover").head(5)
        if not short.empty:
            lines += ["", "<b>⚠️ Кончается (запас меньше 2 недель)</b>"] + [
                f"• {esc(x)} — {num(e)} шт, ≈{num(c)} дн." for x, e, c in
                zip(short["brand"], short["end"], short["cover"])]
    except Exception as e:
        print(f"[menu] остатки по брендам: {type(e).__name__}: {e}", flush=True)
    lines += ["", "<i>Все места вместе: склады WB, в пути и наш склад FBS.</i>"]
    ser = series(b, min(max(n, 8), CHART_MAX_DAYS))
    photo = None
    if len(ser) > 1:
        parts = [(name, end[k], (end[k] - beg[k]) if beg else None) for name, k in places]
        ref = next((i for i, r in enumerate(ser) if beg and r["date"] == beg["date"]), None)
        photo = _photo(lambda: ch.stocks(
            [r["date"] for r in ser], [r["total"] for r in ser], parts,
            f"Остатки на {end['date']:%d.%m} — {num(end['total'])} шт",
            "склады WB, в пути, наш склад FBS", ref=ref or 0, note="за период"))
    files = (("brand", "Excel: по брендам"), ("category", "Excel: по категориям"),
             ("nm", "Excel: по артикулам"))
    chart = {"type": "line", "title": "Остатки всего, шт", "labels": [f"{r['date']:%d.%m}" for r in ser],
             "values": [r["total"] for r in ser]} if len(ser) > 1 else None
    return {"text": "\n".join(x for x in lines if x is not None), "photo": photo, "auto": "brand",
            "chart": chart, "files": files, "tables": tables, "kb": _files_kb("k", a, b, files)}


def file_stocks(by: str, a: date, b: date):
    g, da, db = _stock_table(a, b, by)
    head = {"brand": ["Бренд"], "category": ["Категория"],
            "nm": ["Бренд", "Категория", "Артикул продавца", "Артикул WB", "Наименование"]}[by]
    src = {"brand": ["brand"], "category": ["category"],
           "nm": ["brand", "category", "sa", "nm_id", "title"]}[by]
    cols = head + [f"На {da:%d.%m}, шт" if da else "На начало, шт", f"На {db:%d.%m}, шт",
                   "Изменение, шт", "Продано за период, шт", "Запас, дней"]
    rows = [[r[c] for c in src] + [r["beg"], r["end"], r["chg"], r[QTY],
                                   round(r["cover"]) if r["cover"] is not None and r["cover"] == r["cover"] else None]
            for _, r in g.iterrows()]
    name = {"brand": "по брендам", "category": "по категориям", "nm": "по артикулам"}[by]
    data = _book("Остатки", f"{span(a, b)} · склады WB + в пути + FBS", [(f"Остатки {name}", cols, rows)],
                 notes=[("Остатки", "Склады WB + в пути к клиенту + в пути от клиента + наш склад FBS."),
                        ("Продано", "Продажи минус возвраты за период, как на сайте WB."),
                        ("Запас, дней", "Остаток на конец / средние продажи в день за период.")])
    return f"Остатки {name} {span(a, b)}.xlsx", data, XLSX


# ---------------------------------------------------------------- 💸 расходы WB
def _costs_chart(rows, title, sub) -> bytes:
    from . import digest_charts as ch
    plt = ch._plt()
    rows = sorted(rows, key=lambda r: -r[1])
    fig, ax = plt.subplots(figsize=(10.8, 1.8 + 0.9 * len(rows)), dpi=100)
    mx = max(v for _, v, _ in rows) or 1
    for i, (name, v, old) in enumerate(rows):
        y = len(rows) - 1 - i
        ax.barh(y, v, height=0.58, color=ch.GREEN)
        ax.text(-mx * 0.02, y, name, va="center", ha="right", fontsize=15, color=ch.INK)
        lab = rub(v) + (f"   {chg(v, old)}" if old else "")
        ax.text(v + mx * 0.02, y, lab, va="center", ha="left", fontsize=14.5, fontweight="bold",
                color=ch.INK)
    ax.set_xlim(-mx * 0.45, mx * 1.7)
    ax.set_ylim(-0.6, len(rows) - 0.4)
    ax.axis("off")
    h = fig.get_figheight()
    ch._head(fig, title, sub, y=1 - 0.42 / h)
    fig.subplots_adjust(left=0.03, right=0.97, top=1 - 1.3 / h, bottom=0.03)
    return ch._png(plt, fig)


def report_costs(a: date, b: date) -> dict | None:
    from .digest_period import wb_costs_range
    pa, pb = prev_range(a, b)
    rows = [r for r in wb_costs_range(a, b, pa, pb) if r[1] or r[2]]
    if not rows:
        return None
    rp = _tot(sales(a, b, "day"), RP)
    o_rp = _tot(sales(pa, pb, "day"), RP)
    tot, o_tot = sum(r[1] for r in rows), sum(r[2] for r in rows)
    lines = [f"<b>💸 Расходы WB · {span(a, b)}</b>",
             f"<i>сравнение с {span(pa, pb)} · суммы с НДС, как в отчёте WB</i>", "",
             f"Всего: <b>{rub(tot)}</b>  {chg(tot, o_tot)}"]
    if rp:
        lines.append(f"Доля от продаж до СПП: <b>{num(tot / rp * 100, 1)}%</b>  "
                     f"{chg(tot / rp * 100, o_tot / o_rp * 100 if o_rp else 0, pp=True)}")
    lines.append("")
    for name, v, o in sorted(rows, key=lambda r: -r[1]):
        share = f" · {num(v / rp * 100, 1)}% от продаж" if rp else ""
        lines.append(f"{name}: <b>{rub(v)}</b>{share}  {chg(v, o)}")
    worst = max(rows, key=lambda r: r[1] - r[2])
    if worst[1] - worst[2] > 0 and worst[2]:
        lines += ["", f"Сильнее всего выросла статья «{worst[0]}»: +{rub(worst[1] - worst[2])}."]
    lines += ["", "<i>Комиссия WB сюда не входит. Подробно по дням и статьям — в Excel; "
                  "какие карточки и склады тянут логистику и штрафы — в PDF кнопкой ниже.</i>"]
    photo = _photo(lambda: _costs_chart(rows, f"Расходы WB — {rub(tot)}",
                                        f"{span(a, b)} · изменение к {span(pa, pb)}"))
    files = (("xl", "Excel: расходы WB"), ("lpdf", "PDF: логистика и штрафы"),
             ("lxl", "Excel: логистика и штрафы"))
    srt = sorted(rows, key=lambda r: -r[1])
    chart = {"type": "hbar", "title": "Расходы WB по статьям, ₽",
             "labels": [r[0] for r in srt], "values": [round(r[1]) for r in srt]}
    return {"text": "\n".join(lines), "photo": photo, "auto": "xl", "chart": chart,
            "files": files, "kb": _files_kb("w", a, b, files)}


def file_costs(code: str, a: date, b: date):
    from .. import wb_top_cards_report as tc
    s, e = a.isoformat(), b.isoformat()
    if code == "lpdf":
        return f"Логистика и штрафы {span(a, b)}.pdf", tc.build_top_cards_pdf_bytes(s, e), PDF
    if code == "lxl":
        return f"Логистика и штрафы {span(a, b)}.xlsx", tc.build_top_cards_excel(s, e), XLSX
    from ..wb_expenses_excel import build_wb_expenses_excel
    return f"Расходы WB {span(a, b)}.xlsx", build_wb_expenses_excel(a.isoformat(), b.isoformat()), XLSX


# ---------------------------------------------------------------- 📣 реклама
_ads_cache: dict = {}


def _ads_bundle(a: date, b: date) -> dict:
    """Те же данные, что у PDF-записки и Excel: считаем один раз на текст и оба файла."""
    import time
    from .. import ad_campaigns_report as ad
    key = (a, b)
    hit = _ads_cache.get(key)
    if hit and time.time() - hit[0] < 600:
        return hit[1]
    bundle = ad._fetch_report_bundle(a.isoformat(), b.isoformat())
    if len(_ads_cache) > 4:
        _ads_cache.clear()
    _ads_cache[key] = (time.time(), bundle)
    return bundle


def report_ads(a: date, b: date) -> dict | None:
    from .. import ad_campaigns_report as ad
    bnd = _ads_bundle(a, b)
    t, o = bnd.get("totals") or {}, bnd.get("prev_totals") or {}
    if not t.get("spend_rub") and not t.get("views"):
        return None
    pa, pb = date.fromisoformat(str(bnd["prev_start"])[:10]), date.fromisoformat(str(bnd["prev_end"])[:10])
    g = lambda d, k: d.get(k) or 0
    dd = lambda d: d.get("drr_direct") if d.get("drr_direct") is not None else d.get("drr")  # noqa: E731
    rating = ad._drr_rating(dd(t))
    lines = [f"<b>📣 Реклама WB · {span(a, b)}</b>",
             f"<i>сравнение с {span(pa, pb)}</i>", "",
             f"Расход: <b>{rub(g(t, 'spend_rub'))}</b>  {chg(g(t, 'spend_rub'), g(o, 'spend_rub'))}",
             f"Выручка с рекламы: <b>{rub(g(t, 'revenue_rub'))}</b>  "
             f"{chg(g(t, 'revenue_rub'), g(o, 'revenue_rub'))}",
             f"ДРР по рекламируемым товарам: <b>{num(dd(t) or 0, 1)}%</b>  "
             f"{chg(dd(t) or 0, dd(o) or 0, pp=True)}" + (f" · {rating.lower()}" if rating else ""),
             f"ДРР с учётом связанных заказов: {num(g(t, 'drr'), 1)}%",
             f"Заказы: <b>{num(g(t, 'orders'))}</b>  {chg(g(t, 'orders'), g(o, 'orders'))}",
             f"Стоимость заказа: <b>{rub(g(t, 'cpo'))}</b>  {chg(g(t, 'cpo'), g(o, 'cpo'))}",
             f"Показы: {num(g(t, 'views'))} · клики: {num(g(t, 'clicks'))} · CTR {num(g(t, 'ctr'), 2)}%"]
    c = bnd.get("campaigns")
    if c is not None and not c.empty:
        c = c.copy()
        if "drr_direct" in c.columns:                 # как в записке: по рекламируемым товарам
            c["drr"] = c["drr_direct"]
        top = c.sort_values("spend_rub", ascending=False).head(5)
        lines += ["", "<b>Больше всего тратят</b>"] + [
            f"• {esc(str(n)[:40])} — {rub(s_)} · ДРР "
            + (f"{num(d, 1)}%" if d == d and d is not None else "—")
            for n, s_, d in zip(top["name"], top["spend_rub"], top["drr"])]
        bad = c[c["quality"] == ad.QUALITY_BAD].sort_values("spend_rub", ascending=False).head(5)
        if not bad.empty:
            lines += ["", f"<b>⚠️ Неэффективные (ДРР выше {num(ad.DRR_HIGH_THRESHOLD)}%)</b>"] + [
                f"• {esc(str(n)[:40])} — {rub(s_)}, ДРР "
                + (f"{num(d, 1)}%" if d == d and d is not None else "нет заказов")
                for n, s_, d in zip(bad["name"], bad["spend_rub"], bad["drr"])]
    bud = bnd.get("budget") or {}
    if bud.get("net") is not None:
        left = f"Доступно на счёте рекламы: <b>{rub(bud['net'])}</b>"
        if bud.get("budgets_total"):
            left += f" + в бюджетах кампаний {rub(bud['budgets_total'])}"
        lines += ["", left]
        if bud.get("runway_days"):
            lines.append(f"При расходе последних {ad.RUNWAY_LOOKBACK_DAYS} дней "
                         f"(≈{rub(bud['avg_daily_spend'])} в день) хватит примерно на "
                         f"<b>{num(bud['runway_days'])} дн.</b>")
    lines += ["", "<i>По статистике кампаний WB — те же цифры, что в аналитической записке. "
                  "Записка (PDF) и детализация (Excel) — следом.</i>"]
    daily = bnd.get("daily")
    chart = None
    try:
        if daily is not None and len(daily) > 1:
            dd = pd.to_datetime(daily[daily.columns[0]])
            chart = {"type": "bar", "title": "Расход на рекламу по дням, ₽",
                     "labels": [f"{x:%d.%m}" for x in dd],
                     "values": [round(float(v or 0)) for v in daily["spend_rub"]], "n_cur": len(daily)}
    except Exception:
        chart = None
    return {"text": "\n".join(lines), "photo": None, "auto": ["pdf", "xl"], "chart": chart,
            "files": (("pdf", "PDF: записка"), ("xl", "Excel: кампании"), ("per", "Excel: по периодам")),
            "kb": _files_kb("a", a, b, (("per", "Excel: по периодам"),))}


def file_ads(code: str, a: date, b: date):
    from .. import ad_campaigns_report as ad
    if code == "pdf":
        return f"Реклама WB — записка {span(a, b)}.pdf", ad.build_ad_campaigns_pdf(_ads_bundle(a, b)), PDF
    if code == "per":
        return (f"Реклама WB по периодам {span(a, b)}.xlsx",
                ad.build_ad_campaigns_periods_excel(a.isoformat(), b.isoformat()), XLSX)
    return f"Реклама WB {span(a, b)}.xlsx", ad._assemble_excel(_ads_bundle(a, b)), XLSX


# ---------------------------------------------------------------- 💳 скидка СПП
DISC = "СПП, %"
AVG_RP, AVG_RA = "Средняя цена до СПП, ₽", "Средняя цена после СПП, ₽"


def _spp(df) -> float:
    rp = _tot(df, RP)
    return (1 - _tot(df, RA) / rp) * 100 if rp else 0.0


def _spp_chart(days, spp, avg_rp, avg_ra, title, sub) -> bytes:
    from . import digest_charts as ch
    plt = ch._plt()
    fig = plt.figure(figsize=(10.8, 9.0), dpi=100)
    ax = fig.add_axes([0.08, 0.52, 0.86, 0.30])
    xs = range(len(days))
    ax.plot(xs, spp, color=ch.WARM, lw=2.6)
    ax.scatter([xs[-1]], [spp[-1]], s=90, color=ch.WARM, zorder=3, edgecolors="white", linewidths=2)
    ax.text(xs[-1], spp[-1] + 0.6, num(spp[-1], 1) + "%", ha="center", fontsize=16,
            fontweight="bold", color=ch.INK)
    ax2 = fig.add_axes([0.08, 0.08, 0.86, 0.30])
    ax2.plot(xs, avg_rp, color=ch.GREEN, lw=2.4, label="до СПП")
    ax2.plot(xs, avg_ra, color=ch.MID, lw=2.4, label="после СПП (платит покупатель)")
    ax2.legend(frameon=False, fontsize=12.5, loc="lower right", bbox_to_anchor=(1.0, 1.02), ncol=2)
    step = 1 if len(days) <= 10 else 3 if len(days) <= 31 else 7
    for a_ in (ax, ax2):
        a_.set_xticks(list(xs))
        a_.set_xticklabels([f"{d:%d.%m}" if (len(days) - 1 - i) % step == 0 else ""
                            for i, d in enumerate(days)], color=ch.MUTED, fontsize=12)
        a_.yaxis.set_major_locator(plt.MaxNLocator(4))
        a_.tick_params(axis="y", labelcolor=ch.MUTED, labelsize=12)
        a_.yaxis.grid(True, color=ch.GRID, lw=1)
        a_.set_axisbelow(True)
        ch._clean(a_)
    ax.yaxis.set_major_formatter(lambda v, _: f"{v:g}%".replace(".", ","))
    ax2.yaxis.set_major_formatter(lambda v, _: num(v))
    ch._head(fig, title, sub, y=0.945)
    fig.text(0.06, 0.845, "СПП по дням, %", fontsize=13, fontweight="bold", color=ch.INK)
    fig.text(0.06, 0.41, "Средняя цена за штуку, ₽", fontsize=13, fontweight="bold", color=ch.INK)
    return ch._png(plt, fig)


def report_disc(a: date, b: date) -> dict | None:
    pa, pb = prev_range(a, b)
    df = sales(pa, b, "day")
    if df.empty:
        return None
    dcol = df.columns[0]
    df[dcol] = pd.to_datetime(df[dcol]).dt.date
    cur, old = df[df[dcol] >= a], df[df[dcol] <= pb]
    if cur.empty:
        return None
    q, oq = _tot(cur, QTY), _tot(old, QTY)
    rp, ra, orp, ora = _tot(cur, RP), _tot(cur, RA), _tot(old, RP), _tot(old, RA)
    spp, ospp = _spp(cur), _spp(old)
    lines = [f"<b>💳 Скидка СПП · {span(a, b)}</b>",
             f"<i>сравнение с {span(pa, pb)}</i>", "",
             f"СПП: <b>{num(spp, 1)}%</b>  {chg(spp, ospp, pp=True)}",
             f"Скидка WB в рублях: <b>{rub(rp - ra)}</b>  {chg(rp - ra, orp - ora)}",
             f"Средняя цена до СПП: <b>{rub(rp / q if q else 0)}</b>  "
             f"{chg(rp / q if q else 0, orp / oq if oq else 0)}",
             f"Платит покупатель (после СПП): <b>{rub(ra / q if q else 0)}</b>  "
             f"{chg(ra / q if q else 0, ora / oq if oq else 0)}"]
    if len(cur) > 1:
        cur = cur.assign(_spp=[(1 - r_a / r_p) * 100 if r_p else None
                               for r_p, r_a in zip(cur[RP], cur[RA])])
        hi_d, lo_d = cur.loc[cur["_spp"].idxmax()], cur.loc[cur["_spp"].idxmin()]
        lines += ["", f"Максимум: <b>{num(hi_d['_spp'], 1)}%</b> — {hi_d[dcol]:%d.%m}",
                  f"Минимум: <b>{num(lo_d['_spp'], 1)}%</b> — {lo_d[dcol]:%d.%m}"]
    try:
        g, og = sales(a, b, "brand"), sales(pa, pb, "brand")
        m = g[["Бренд", RP, RA]].merge(og[["Бренд", RP, RA]], on="Бренд", how="left",
                                       suffixes=("", "_o")).fillna(0)
        m = m[m[RP] >= max(rp * 0.01, 1)]
        m["s"] = (1 - m[RA] / m[RP]) * 100
        m["so"] = [(1 - x / y) * 100 if y else None for x, y in zip(m[RA + "_o"], m[RP + "_o"])]
        m["d"] = m["s"] - m["so"]
        up = m[m["d"].notna()].sort_values("d", ascending=False).head(5)
        dn = m[m["d"].notna()].sort_values("d").head(5)
        if not up.empty and up["d"].iloc[0] > 0.5:
            lines += ["", "<b>▲ Скидка выросла сильнее всего</b>"] + [
                f"• {esc(x)} — {num(s_, 1)}% ({chg(s_, so, pp=True)})"
                for x, s_, so, d_ in zip(up["Бренд"], up["s"], up["so"], up["d"]) if d_ > 0.5]
        if not dn.empty and dn["d"].iloc[0] < -0.5:
            lines += ["", "<b>▼ Скидка снизилась сильнее всего</b>"] + [
                f"• {esc(x)} — {num(s_, 1)}% ({chg(s_, so, pp=True)})"
                for x, s_, so, d_ in zip(dn["Бренд"], dn["s"], dn["so"], dn["d"]) if d_ < -0.5]
    except Exception as e:
        print(f"[menu] СПП по брендам: {type(e).__name__}: {e}", flush=True)
    lines += ["", "<i>СПП — скидка постоянного покупателя за счёт WB: 1 − после СПП / до СПП.</i>"]
    photo = None
    show = df if ndays(pa, b) <= CHART_MAX_DAYS else cur
    if len(show) > 1:
        vals = [((1 - r_a / r_p) * 100 if r_p else 0, r_p / qq if qq else 0, r_a / qq if qq else 0)
                for r_p, r_a, qq in zip(show[RP], show[RA], show[QTY])]
        photo = _photo(lambda: _spp_chart(list(show[dcol]), [v[0] for v in vals],
                                          [v[1] for v in vals], [v[2] for v in vals],
                                          f"СПП — {num(spp, 1)}%",
                                          f"{span(show[dcol].iloc[0], b)} · до и после СПП"))
    files = (("day", "Excel: по дням"), ("brand", "Excel: по брендам"), ("category", "Excel: по категориям"))
    chart = {"type": "line", "title": "СПП по дням, %", "labels": [f"{d:%d.%m}" for d in show[dcol]],
             "values": [round((1 - r_a / r_p) * 100, 1) if r_p else None
                        for r_p, r_a in zip(show[RP], show[RA])]} if len(show) > 1 else None
    return {"text": "\n".join(lines), "photo": photo, "auto": "day", "chart": chart,
            "files": files, "kb": _files_kb("d", a, b, files)}


def file_disc(code: str, a: date, b: date):
    if code not in ("day", "brand", "category"):
        return None
    df = sales(a, b, code)
    if df.empty:
        return None
    key = {"day": "Дата", "brand": "Бренд", "category": "Категория"}[code]
    if code == "day":
        df[key] = pd.to_datetime(df[key]).dt.date
    cols = [key, QTY, RP, RA, AVG_RP, AVG_RA, DISC]
    df = df[cols].copy()
    df.insert(4, "Скидка WB, ₽", df[RP] - df[RA])
    if code != "day":
        df = df.sort_values(RP, ascending=False)
    cols = list(df.columns)
    name = {"day": "по дням", "brand": "по брендам", "category": "по категориям"}[code]
    data = _book("Скидка СПП", f"{span(a, b)} · как на сайте WB, с НДС",
                 [(f"СПП {name}", cols, df.values.tolist())],
                 notes=[("СПП, %", "1 − После СПП / До СПП, по чистым суммам."),
                        ("Скидка WB, ₽", "До СПП минус после СПП: сколько скидки дал WB."),
                        ("Средняя цена", "Сумма / проданные штуки (продажи минус возвраты).")])
    return f"Скидка СПП {name} {span(a, b)}.xlsx", data, XLSX


# ---------------------------------------------------------------- 📑 коммерческий обзор 2.0
def report_review(a: date, b: date) -> dict | None:
    lines = [f"<b>📑 Коммерческий обзор 2.0 на {b:%d.%m.%Y}</b>", "",
             "Тот же PDF, что в дашборде: продажи и их причины, бренды и категории, "
             "остатки, FBS, расходы WB. Собираю — это может занять до минуты."]
    return {"text": "\n".join(lines), "photo": None, "auto": "pdf", "files": (("pdf", "PDF: обзор"),),
            "kb": {"inline_keyboard": [
        [{"text": "🔁 Другая дата", "callback_data": "mn:o:r"}]]}}


def file_review(b: date):
    from ..commercial_review.data import build_review_data
    from ..commercial_review.report import build_commercial_review_pdf
    payload, fbs = build_review_data(b)
    return f"Коммерческий обзор {b:%d.%m.%Y}.pdf", build_commercial_review_pdf(payload, fbs=fbs), PDF


# ---------------------------------------------------------------- 💰 маржинальность / 🔻 убыток
MANAGEMENT_NOTE = ("<i>Управленческий учёт, как в дашборде продаж: выручка без НДС, себестоимость "
                   "по FIFO, продажи скорректированы на возвраты (возврат сторнирует исходную "
                   "продажу по её себестоимости), комиссия WB без НДС; расходы WB (логистика, "
                   "хранение, штрафы, продвижение) распределены на проданные штуки по неделям.</i>")
MARGIN_SQL = """
WITH commissions AS (
    SELECT rrd_id,
           COALESCE(SUM(val / (100 + vat_rate) * 100)
                    FILTER (WHERE field = 'comission' AND oper = 'dt'), 0)
           - COALESCE(SUM(val / (100 + vat_rate) * 100)
                      FILTER (WHERE field = 'comission' AND oper = 'cr'), 0) AS net_comission
    FROM sales.sales_long GROUP BY rrd_id
),
wk AS (
    SELECT YEARWEEK(date_from::DATE) AS yw,
           SUM(CASE WHEN cr_rev > 0 THEN 1 WHEN cr_rev < 0 THEN -1 ELSE 0 END) AS n
    FROM base WHERE cr_rev <> 0 GROUP BY 1
),
cw AS (
    SELECT yw, SUM(dt / (100 + vat_rate) * 100) - SUM(cr / (100 + vat_rate) * 100) AS costs
    FROM wb_costs GROUP BY yw
),
cps AS (
    SELECT wk.yw, CASE WHEN COALESCE(wk.n, 0) = 0 THEN NULL ELSE ABS(cw.costs) / wk.n END AS r
    FROM wk LEFT JOIN cw ON cw.yw = wk.yw
)
SELECT {dims},
       SUM(CASE WHEN t.cr_rev > 0 THEN 1 WHEN t.cr_rev < 0 THEN -1 ELSE 0 END) AS qty,
       SUM(t.cr_rev / (100 + t.vat_rate) * 100) / 100 AS rev,
       SUM(t.adjusted_cogs_man) / 100 AS cogs,
       SUM(c.net_comission) / 100 AS comm,
       SUM((CASE WHEN t.cr_rev > 0 THEN 1 WHEN t.cr_rev < 0 THEN -1 ELSE 0 END)
           * COALESCE(cps.r, 0)) / 100 AS wb
FROM base t
LEFT JOIN commissions c ON c.rrd_id = t.rrd_id
LEFT JOIN cps ON cps.yw = YEARWEEK(t.date_from::DATE)
WHERE t.cr_rev <> 0 AND t.date_from::DATE BETWEEN ? AND ?
GROUP BY ALL
"""
MDIMS = {
    "brand": ("COALESCE(NULLIF(TRIM(UPPER(t.brand)), ''), 'НЕ УКАЗАН') AS name", ["Бренд"]),
    "category": ("COALESCE(NULLIF(TRIM(t.subject_name), ''), 'Не указана') AS name", ["Категория"]),
    "nm": ("COALESCE(NULLIF(TRIM(UPPER(t.brand)), ''), 'НЕ УКАЗАН') AS brand, "
           "COALESCE(NULLIF(TRIM(t.subject_name), ''), 'Не указана') AS category, "
           "t.usk AS nm_id, COALESCE(NULLIF(TRIM(t.title), ''), '') AS name",
           ["Бренд", "Категория", "Артикул WB", "Наименование"]),
}
_mcache: dict = {}


def margin(a: date, b: date, by: str) -> pd.DataFrame:
    """Маржа по методике дашборда продаж (DashboardData: base + wb_costs)."""
    import time
    key = (a, b, by)
    hit = _mcache.get(key)
    if hit and time.time() - hit[0] < 1800:
        return hit[1].copy()
    from ...data.base import DashboardData
    with DashboardData() as d:
        df = d.con.execute(MARGIN_SQL.replace("{dims}", MDIMS[by][0]), [a, b]).df()
    for c in ("qty", "rev", "cogs", "comm", "wb"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    df["md1"] = df["rev"] - df["cogs"] + df["comm"]
    df["res"] = df["md1"] - df["wb"]
    if len(_mcache) > 12:
        _mcache.clear()
    _mcache[key] = (time.time(), df)
    return df.copy()


def _pct(v, base):
    return v / base * 100 if base else 0.0


def report_margin(a: date, b: date) -> dict | None:
    pa, pb = prev_range(a, b)
    g, og = margin(a, b, "brand"), margin(pa, pb, "brand")
    if g.empty or not g["rev"].sum():
        return None
    T = {k: g[k].sum() for k in ("rev", "cogs", "comm", "md1", "wb", "res", "qty")}
    O = {k: og[k].sum() for k in ("rev", "md1", "res")} if not og.empty else {"rev": 0, "md1": 0, "res": 0}
    m1, om1 = _pct(T["md1"], T["rev"]), _pct(O["md1"], O["rev"])
    m2, om2 = _pct(T["res"], T["rev"]), _pct(O["res"], O["rev"])
    lines = [f"<b>💰 Маржинальность · {span(a, b)}</b>",
             f"<i>сравнение с {span(pa, pb)} · управленческий учёт</i>", "",
             f"Выручка без НДС: <b>{rub(T['rev'])}</b>  {chg(T['rev'], O['rev'])}",
             f"Себестоимость: <b>{rub(T['cogs'])}</b> · {num(_pct(T['cogs'], T['rev']), 1)}%",
             f"Комиссия WB: <b>{rub(-T['comm'])}</b> · {num(_pct(-T['comm'], T['rev']), 1)}%",
             f"Маржинальность: <b>{num(m1, 1)}%</b> · {rub(T['md1'])}  {chg(m1, om1, pp=True)}",
             f"Расходы WB: <b>{rub(T['wb'])}</b> · {num(_pct(T['wb'], T['rev']), 1)}%",
             f"Финансовый результат: <b>{num(m2, 1)}%</b> · {rub(T['res'])}  {chg(m2, om2, pp=True)}"]
    g = g[g["rev"] > max(T["rev"] * 0.005, 1)].copy()
    g["m2"] = g["res"] / g["rev"] * 100
    best = g.sort_values("m2", ascending=False).head(5)
    worst = g.sort_values("m2").head(5)
    lines += ["", "<b>▲ Самые доходные бренды</b>"] + [
        f"• {esc(n)} — {num(p, 1)}% · {rub(r)}" for n, p, r in zip(best["name"], best["m2"], best["res"])]
    lines += ["", "<b>▼ Самые слабые бренды</b>"] + [
        f"• {esc(n)} — {num(p, 1)}% · {rub(r)}" for n, p, r in zip(worst["name"], worst["m2"], worst["res"])]
    lines += ["", "Маржинальность = выручка − себестоимость − комиссия; финансовый результат — "
                  "ещё минус расходы WB. % — от выручки без НДС.", "", MANAGEMENT_NOTE]
    files = (("brand", "Excel: по брендам"), ("category", "Excel: по категориям"),
             ("nm", "Excel: по артикулам"))
    tables = []
    for by, ttl in (("brand", "Бренды"), ("category", "Категории")):
        t = margin(a, b, by).sort_values("rev", ascending=False)
        tables.append(table(ttl, ["Бренд" if by == "brand" else "Категория", "Выручка", "Маржа",
                                  "Фин. результат", "Фин. рез., ₽"],
                            [[r["name"], r["rev"], _pct(r["md1"], r["rev"]) if r["rev"] else None,
                              _pct(r["res"], r["rev"]) if r["rev"] else None, r["res"]]
                             for _, r in t.iterrows()],
                            ["text", "rub", "pct", "pct", "rub"]))
    return {"text": "\n".join(lines), "photo": None, "auto": "brand", "files": files,
            "tables": tables, "kb": _files_kb("m", a, b, files)}


def _margin_rows(df, by):
    df = df.sort_values("res")
    keys = {"brand": ["name"], "category": ["name"], "nm": ["brand", "category", "nm_id", "name"]}[by]
    cols = MDIMS[by][1] + ["Продано, шт", "Выручка без НДС, ₽", "Себестоимость, ₽",
                           "Комиссия WB, ₽", "Маржинальность, ₽", "Маржинальность, %",
                           "Расходы WB, ₽", "Финансовый результат, ₽", "Финансовый результат, %",
                           "Доля выручки, %"]
    tot = df["rev"].sum()
    rows = [[r[k] for k in keys] + [r["qty"], r["rev"], -r["cogs"], r["comm"], r["md1"],
                                    _pct(r["md1"], r["rev"]) if r["rev"] else None, -r["wb"],
                                    r["res"], _pct(r["res"], r["rev"]) if r["rev"] else None,
                                    _pct(r["rev"], tot)]
            for _, r in df.iterrows()]
    return cols, rows


MARGIN_NOTES = [
    ("Учёт", "Управленческий, по методике дашборда продаж."),
    ("Выручка", "Наша цена до СПП без НДС, продажи минус возвраты."),
    ("Себестоимость", "Управленческая, FIFO. Возврат сторнирует исходную продажу по её "
                      "себестоимости (корректировка продаж на возвраты)."),
    ("Комиссия WB", "Без НДС."),
    ("Расходы WB", "Логистика, хранение, приёмка, штрафы, продвижение, лояльность, коррекции — "
                   "без НДС, распределены на проданные штуки по неделям, как в дашборде."),
    ("Финансовый результат", "Маржинальность минус расходы WB. Меньше нуля — продажи в убыток."),
]


def file_margin(by: str, a: date, b: date, only_loss: bool = False):
    if by not in MDIMS:
        return None
    df = margin(a, b, by)
    if only_loss:
        df = df[df["res"] < 0]
    if df.empty:
        return None
    cols, rows = _margin_rows(df, by)
    name = {"brand": "по брендам", "category": "по категориям", "nm": "по артикулам"}[by]
    title = "В убытке" if only_loss else "Маржинальность"
    data = _book(title, f"{span(a, b)} · управленческий учёт", [(f"{title} {name}", cols, rows)],
                 notes=MARGIN_NOTES)
    return f"{title} {name} {span(a, b)}.xlsx", data, XLSX


def report_loss(a: date, b: date) -> dict | None:
    nm = margin(a, b, "nm")
    if nm.empty:
        return None
    loss = nm[nm["res"] < 0].sort_values("res")
    total_rev = nm["rev"].sum()
    lines = [f"<b>🔻 Продаём в убыток · {span(a, b)}</b>", "<i>управленческий учёт</i>", ""]
    if loss.empty:
        lines += ["Артикулов с отрицательным финансовым результатом нет. 👍", "", MANAGEMENT_NOTE]
        return {"text": "\n".join(lines), "photo": None, "auto": None,
                "kb": _files_kb("l", a, b, ())}
    lines += [f"Артикулов в убытке: <b>{num(len(loss))}</b> из {num(len(nm))}",
              f"Их выручка: <b>{rub(loss['rev'].sum())}</b> · {num(_pct(loss['rev'].sum(), total_rev), 1)}% "
              f"от всей", f"Убыток по ним: <b>{rub(loss['res'].sum())}</b>", "",
              "<b>Самые убыточные артикулы</b>"]
    for _, r in loss.head(8).iterrows():
        lines.append(f"• {esc(r['brand'])} · {esc(str(r['name'])[:38])} ({r['nm_id']}) — "
                     f"{rub(r['res'])}, {num(r['qty'])} шт")
    br = loss.groupby("brand", as_index=False)[["res", "rev"]].sum().sort_values("res").head(5)
    lines += ["", "<b>Бренды, где больше всего убыточных продаж</b>"] + [
        f"• {esc(n)} — {rub(v)}" for n, v in zip(br["brand"], br["res"])]
    lines += ["", "Убыток = выручка без НДС − себестоимость − комиссия − расходы WB < 0.", "",
              MANAGEMENT_NOTE]
    return {"text": "\n".join(lines), "photo": None, "auto": "nm",
            "files": (("nm", "Excel: убыточные артикулы"), ("brand", "Excel: убыток по брендам")),
            "kb": _files_kb("l", a, b, (("brand", "Excel: убыток по брендам"),))}


# ---------------------------------------------------------------- 📅 недели
def _weeks(a: date, b: date) -> list[tuple[date, date, bool]]:
    """Календарные недели пн–вс от недели с a до недели с b: (пн, последний день, полная?)."""
    mon = a - timedelta(days=a.weekday())
    out = []
    while mon <= b:
        end = min(mon + timedelta(days=6), b)
        out.append((mon, end, end == mon + timedelta(days=6)))
        mon += timedelta(days=7)
    return out


def _week_stats(df, dcol, x, y):
    part = df[(df[dcol] >= x) & (df[dcol] <= y)]
    q, sold, ret = _tot(part, QTY), _tot(part, SOLD), _tot(part, RET)
    rp, ra = _tot(part, RP), _tot(part, RA)
    return {"rp": rp, "ra": ra, "pay": _tot(part, PAY), "q": q,
            "price": rp / q if q else 0, "ret": ret / sold * 100 if sold else 0,
            "spp": (1 - ra / rp) * 100 if rp else 0, "rps": _tot(part, RP_S), "rpr": _tot(part, RP_R),
            "sold": sold, "days": len(part)}


def report_weeks(a: date, b: date) -> dict | None:
    weeks = _weeks(a, b)
    if not weeks:
        return None
    first = weeks[0][0] - timedelta(days=7)                       # для сравнения первой недели
    df = sales(first, b, "day")
    if df.empty:
        return None
    dcol = df.columns[0]
    df[dcol] = pd.to_datetime(df[dcol]).dt.date
    rows = []
    for mon, end, full in weeks:
        n = (end - mon).days
        cur = _week_stats(df, dcol, mon, end)
        prev = _week_stats(df, dcol, mon - timedelta(days=7), mon - timedelta(days=7) + timedelta(days=n))
        rows.append((mon, end, full, cur, prev))
    last_mon, last_end, last_full, cur, prev = rows[-1]
    lbl = lambda m, e, f: f"{m:%d.%m}–{(m + timedelta(days=6)):%d.%m}" + ("" if f else f" (по {e:%d.%m})")
    lines = [f"<b>📅 Недели · {span(rows[0][0], b)}</b>",
             f"<i>неделя {lbl(last_mon, last_end, last_full)} к прошлой · календарные недели пн–вс"
             + ("" if last_full else ", неполная — к тем же дням") + "</i>", "",
             f"До СПП: <b>{rub(cur['rp'])}</b>  {chg(cur['rp'], prev['rp'])}",
             f"Продано: <b>{num(cur['q'])} шт</b>  {chg(cur['q'], prev['q'])}",
             f"Средняя цена: <b>{rub(cur['price'])}</b>  {chg(cur['price'], prev['price'])}",
             f"Возвраты: <b>{num(cur['ret'], 1)}%</b>  {chg(cur['ret'], prev['ret'], pp=True)}",
             f"СПП: <b>{num(cur['spp'], 1)}%</b>  {chg(cur['spp'], prev['spp'], pp=True)}"]
    cw = df[(df[dcol] >= last_mon) & (df[dcol] <= last_end)]
    pw = df[(df[dcol] >= last_mon - timedelta(days=7)) &
            (df[dcol] <= last_mon - timedelta(days=7) + (last_end - last_mon))]
    why = _why(cw, pw)
    if why:
        lines += ["", why]
    if len(rows) > 1:
        lines += ["", "<b>По неделям (до СПП · шт · к прошлой)</b>"]
        for mon, end, full, c, p in rows:
            lines.append(f"• {lbl(mon, end, full)} — {rub(c['rp'])} · {num(c['q'])} шт  {chg(c['rp'], p['rp'])}")
        full_rows = [r for r in rows if r[2]]
        if len(full_rows) > 1:
            best = max(full_rows, key=lambda r: r[3]["rp"])
            worst = min(full_rows, key=lambda r: r[3]["rp"])
            avg = sum(r[3]["rp"] for r in full_rows) / len(full_rows)
            lines += ["", f"Лучшая неделя: <b>{lbl(best[0], best[1], True)}</b> — {rub(best[3]['rp'])}",
                      f"Самая слабая: <b>{lbl(worst[0], worst[1], True)}</b> — {rub(worst[3]['rp'])}",
                      f"В среднем за полную неделю: {rub(avg)}"]
            ups = sum(1 for r in full_rows[-3:] if r[3]["rp"] > r[4]["rp"])
            if len(full_rows) >= 3 and ups in (0, 3):
                lines.append("Три недели подряд " + ("рост" if ups == 3 else "падение")
                             + " — это уже тренд, а не случайность.")
    lines += ["", "<b>Цена и скидка по неделям</b>"] + [
        f"• {m:%d.%m} — до СПП {rub(c['price'])} · покупатель {rub(c['ra'] / c['q'] if c['q'] else 0)}"
        f" · СПП {num(c['spp'], 1)}%" for m, e, f, c, p in rows]
    if len(rows) > 1:
        f0, f1 = rows[0][3], rows[-1][3]
        dp = f1["price"] - f0["price"]
        ds = f1["spp"] - f0["spp"]
        parts = []
        if abs(dp) >= 1:
            parts.append(f"цена до СПП {'выросла' if dp > 0 else 'снизилась'} на {rub(abs(dp))}")
        if abs(ds) >= 0.1:
            parts.append(f"СПП {'выросла' if ds > 0 else 'снизилась'} на {num(abs(ds), 1)} п.п.")
        lines.append(f"За {len(rows)} нед.: " + ", ".join(parts) + "." if parts
                     else "Цена и скидка за эти недели почти не менялись.")
    lines += ["", "<i>Как на сайте WB: с НДС, по дате операции в отчёте реализации.</i>"]
    labels = [f"{m:%d.%m}" + ("" if f else "*") for m, e, f, *_ in rows]
    chart = {"type": "bar", "title": "Продажи до СПП по неделям, ₽", "labels": labels,
             "values": [round(r[3]["rp"]) for r in rows], "n_cur": len(rows), "label_all": True}
    charts = [chart,
              {"type": "line", "title": "Средняя цена за штуку, ₽", "labels": labels, "label_all": True,
               "series": [{"name": "до СПП", "values": [round(r[3]["price"]) for r in rows]},
                          {"name": "платит покупатель", "values": [
                              round(r[3]["ra"] / r[3]["q"]) if r[3]["q"] else None for r in rows]}]},
              {"type": "line", "title": "СПП по неделям, %", "labels": labels, "label_all": True,
               "values": [round(r[3]["spp"], 1) for r in rows]}]
    wk_table = table("Недели", ["Неделя", "До СПП", "Шт", "К прошлой", "Цена до СПП", "Цена покупателя",
                                "СПП", "Возвраты"],
                     [[lbl(m, e, f), c["rp"], c["q"], (c["rp"] / p["rp"] - 1) * 100 if p["rp"] else None,
                       c["price"], c["ra"] / c["q"] if c["q"] else None, c["spp"], c["ret"]]
                      for m, e, f, c, p in rows],
                     ["text", "rub", "int", "signed_pct", "rub", "rub", "pct", "pct"])
    photo = None
    if len(rows) > 1:
        from . import digest_charts as ch
        photo = _photo(lambda: ch.trend([f"{m:%d.%m}" + ("" if f else "*") for m, _, f, *_ in rows],
                                        [r[3]["rp"] / 1e6 for r in rows], [None] * len(rows),
                                        "Продажи по неделям, млн ₽",
                                        f"{span(rows[0][0], b)} · до СПП"
                                        + ("" if rows[-1][2] else " · * неделя ещё идёт")))
    photos = []
    if len(rows) > 1:                                    # цена и скидка — второй картинкой
        from . import digest_charts as ch
        pic = _photo(lambda: ch.price_spp(
            labels, [r[3]["price"] for r in rows],
            [r[3]["ra"] / r[3]["q"] if r[3]["q"] else None for r in rows], [r[3]["spp"] for r in rows],
            "Цена и скидка по неделям", f"{span(rows[0][0], b)} · с НДС"
            + ("" if rows[-1][2] else " · * неделя ещё идёт")))
        if pic:
            photos.append(pic)
    files = (("weeks", "Excel: недели"), ("brands", "Excel: бренды по неделям"),
             ("cats", "Excel: категории по неделям"))
    return {"text": "\n".join(lines), "photo": photo, "photos": photos, "auto": "weeks", "chart": chart,
            "charts": charts,
            "files": files, "tables": [wk_table], "kb": _files_kb("n", a, b, files)}


def file_weeks(code: str, a: date, b: date):
    weeks = _weeks(a, b)
    if not weeks:
        return None
    lab = lambda m: f"{m:%d.%m}–{m + timedelta(days=6):%d.%m}"
    if code == "weeks":
        df = sales(weeks[0][0] - timedelta(days=7), b, "day")
        if df.empty:
            return None
        dcol = df.columns[0]
        df[dcol] = pd.to_datetime(df[dcol]).dt.date
        cols = ["Неделя", "Дней", "Продано, шт", "До СПП, ₽", "Изменение к прошлой, %", "После СПП, ₽",
                "К перечислению, ₽", "Средняя цена до СПП, ₽", "Возвраты, %", "СПП, %"]
        rows = []
        for mon, end, full in weeks:
            c = _week_stats(df, dcol, mon, end)
            p = _week_stats(df, dcol, mon - timedelta(days=7), mon - timedelta(days=7) + (end - mon))
            rows.append([lab(mon) + ("" if full else f" (по {end:%d.%m})"), c["days"], c["q"], c["rp"],
                         (c["rp"] / p["rp"] - 1) * 100 if p["rp"] else None, c["ra"], c["pay"],
                         c["price"], c["ret"], c["spp"]])
        data = _book("Недели", f"{span(weeks[0][0], b)} · календарные недели пн–вс, с НДС",
                     [("Недели", cols, rows)],
                     notes=[("Неделя", "Понедельник–воскресенье. Неполная неделя сравнивается с теми "
                                       "же днями прошлой недели."),
                            ("Источник", "Отчёт реализации WB, дата операции, как на сайте WB.")])
        return f"Недели {span(weeks[0][0], b)}.xlsx", data, XLSX
    level, key = {"brands": ("brand_day", "Бренд"), "cats": ("nm_day", "Категория")}.get(code, (None, None))
    if not level:
        return None
    df = sales(weeks[0][0], b, level)
    if df.empty:
        return None
    df["Дата"] = pd.to_datetime(df["Дата"]).dt.date
    df["_w"] = [d - timedelta(days=d.weekday()) for d in df["Дата"]]
    piv = df.pivot_table(index=key, columns="_w", values=RP, aggfunc="sum", fill_value=0)
    piv = piv[sorted(piv.columns)]
    piv["_t"] = piv.sum(axis=1)
    piv = piv.sort_values("_t", ascending=False)
    wk = [c for c in piv.columns if c != "_t"]
    cols = [key] + [f"{lab(w)}, ₽" for w in wk] + ["Итого, ₽"]
    last_full = weeks[-1][2]
    if len(wk) > 1 and last_full:
        cols.append("Последняя к предыдущей, %")
    rows = []
    for name, r in piv.iterrows():
        row = [name] + [float(r[w]) for w in wk] + [float(r["_t"])]
        if len(wk) > 1 and last_full:
            row.append((r[wk[-1]] / r[wk[-2]] - 1) * 100 if r[wk[-2]] else None)
        rows.append(row)
    what = "Бренды" if code == "brands" else "Категории"
    data = _book(what, f"{span(weeks[0][0], b)} · продажи до СПП с НДС"
                 + ("" if last_full else f" · последняя неделя по {b:%d.%m}"),
                 [("По неделям", cols, rows)])
    return f"{what} по неделям {span(weeks[0][0], b)}.xlsx", data, XLSX


# ---------------------------------------------------------------- диспетчер
def build(rep: str, a: date, b: date) -> dict | None:
    if rep == "s":
        return report_sales(a, b)
    if rep in GROUP:
        return report_group(rep, a, b)
    if rep == "k":
        return report_stocks(a, b)
    if rep == "w":
        return report_costs(a, b)
    if rep == "a":
        return report_ads(a, b)
    if rep == "d":
        return report_disc(a, b)
    if rep == "r":
        return report_review(a, b)
    if rep == "n":
        return report_weeks(a, b)
    if rep == "m":
        return report_margin(a, b)
    if rep == "l":
        return report_loss(a, b)
    if rep == "z":
        from . import drop
        return drop.report(a, b)
    if rep == "g":
        from . import drop
        return drop.report_signals(a, b)
    return None


def file(rep: str, code: str, a: date, b: date):
    """(имя, байты, mime) или None."""
    if rep in ("z", "g"):
        from . import drop
        return drop.file_signals(b) if rep == "g" else drop.file(code, a, b)
    if rep == "s":
        return file_sales(code, a, b)
    if rep in GROUP:
        return file_group(rep, code, a, b)
    if rep == "k":
        return file_stocks(code, a, b)
    if rep == "w":
        return file_costs(code, a, b)
    if rep == "a":
        return file_ads(code, a, b)
    if rep == "d":
        return file_disc(code, a, b)
    if rep == "r":
        return file_review(b)
    if rep == "n":
        return file_weeks(code, a, b)
    if rep == "m":
        return file_margin(code, a, b)
    if rep == "l":
        return file_margin(code, a, b, only_loss=True)
    return None


HELP = ("<b>Как пользоваться</b>\n"
        "• Кнопки внизу — готовые отчёты: нажмите, выберите период — придут цифры, "
        "график и Excel. Дополнительные файлы — кнопками под отчётом.\n"
        "• Свой период — «📅 Свой период»: нажмите день начала, потом день конца. "
        "Или просто напишите даты после кнопки отчёта: 01.10–15.10.\n"
        "• Всё сравнивается с прошлым таким же периодом (месяц — с прошлым месяцем).\n"
        "• «🔎 Почему упали продажи» — из чего сложилось изменение: остатки, цена, СПП, реклама, "
        "возвраты, спрос; по каждому артикулу — в Excel.\n"
        "• «🚨 Сигналы» — артикулы, которые за 7 дней заметно просели, и вероятная причина.\n"
        "• Продажи — как на сайте WB: с НДС, до СПП, после СПП и к перечислению.\n"
        "• «📅 Недели» — календарные недели пн–вс: эта к прошлой или несколько недель подряд.\n"
        "• «💳 Скидка СПП» — как менялась скидка WB и цена для покупателя.\n"
        "• «📑 Коммерческий обзор» — PDF коммерческого обзора 2.0 на выбранную дату.\n"
        "• «💰 Маржинальность» и «🔻 Продаём в убыток» — по управленческому учёту, как в "
        "дашборде: себестоимость FIFO, продажи скорректированы на возвраты.\n"
        "• Нужного отчёта нет на кнопках — просто напишите вопрос своими словами, ответит ИИ.\n\n"
        "Важные цифры сверяйте с дашбордом. Если что-то не так — напишите Дарье.")


AI_TEXT = ("🤖 <b>Спросите своими словами</b>\n"
           "Если нужного раздела нет на кнопках — просто напишите вопрос обычным "
           "сообщением, я посчитаю и при необходимости пришлю Excel.")
