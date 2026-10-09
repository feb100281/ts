# gear/app/manpack/excel.py
"""Выгрузки дашборда в Excel — оформление мэн пака."""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook


def tables_xlsx(tables) -> bytes:
    """tables: [(лист, заголовок, подзаголовок, колонки, строки[, опции])] → xlsx."""
    from gear.management.commands import mp
    from ..daily_sales.assistant.excel import _write_table

    wb = Workbook()
    first = True
    stamp = datetime.now().strftime("%d.%m.%Y %H:%M")
    for t in tables:
        sheet, title, subtitle, cols, rows = t[:5]
        opts = t[5] if len(t) > 5 else {}
        row_formats = {i: getattr(mp, f) for i, f in (opts.get("row_formats") or {}).items()}
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet[:31]
        _write_table(ws, mp, title, subtitle, f"Сформировано {stamp} · строк: {len(rows)}",
                     list(cols), [list(r) for r in rows],
                     row_formats=row_formats or None, totals=opts.get("totals", True))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def summary_xlsx(s: dict, me: date, as_of: date) -> bytes:
    """Лист «Выводы» мэн пака за выбранный месяц."""
    from gear.management.commands import mp
    wb = Workbook()
    ws = wb.active
    ws.title = "Выводы"
    mp.build_summary(ws, s, me, as_of=as_of)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def management_pack(report_date: date):
    """Полный мэн пак (как выгрузка из меню «Экспорт»). → (bytes, имя файла)."""
    from gear.management.commands import mp
    path, _stats = mp.build_management_pack(report_date)
    p = Path(path)
    return p.read_bytes(), p.name
